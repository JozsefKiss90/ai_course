"""Segment 5 — Guided setup block (30 min, hands-on, facilitated).

Three slides per lesson_1.md Segment 5: the setup checklist that stays
on screen-share for the whole block (four numbered items in large type
plus the '✅ 1–4' meeting-chat convention), the troubleshooting slide
with the three typical failure points and their fixes, and the catch-up
path (Practice Week 1 drop-in, the offered 10-minute slot, and the note
that Obsidian and Perfect Wiki are deliberately not installed today).
"""

from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches

from deck import template
from deck.template import (
    CONTENT_WIDTH, HIGHLIGHT, PALE, PALE_WARM, PRIMARY, TEXT, WHITE,
)


# --- Slide 1: the setup checklist --------------------------------------------

# The four items from lesson_1.md Segment 5. This slide stays on
# screen-share for the whole 30-minute block, so the rows are taller and
# the type larger than the deck's other numbered lists.
_CHECKLIST_ITEMS = [
    ("Google account + Gemini",
     "Sign in to Google; open gemini.google.com and get an answer to a "
     "test prompt."),
    ("Copilot Chat — institutional account",
     "Open Copilot Chat in Teams or microsoft365.com with your "
     "institutional account (needed from Lesson 2)."),
    ("Perplexity free account",
     "Create a free account at perplexity.ai (needed in Lesson 3 — "
     "cheap to do now)."),
    ("First experimental prompt in each",
     "Send one prompt in each of the three tools — any work question "
     "(nothing you couldn't publish)."),
]

_ROW_HEIGHT = 0.92
_ROW_GAP = 0.08


def _checklist_row(slide, number, top, heading, detail):
    """Numbered checklist row: navy number chip + pale item card."""
    chip = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, template.CONTENT_LEFT, Inches(top),
        Inches(0.7), Inches(_ROW_HEIGHT),
    )
    template.solid(chip, PRIMARY)
    tf = chip.text_frame
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.paragraphs[0].alignment = PP_ALIGN.CENTER
    run = tf.paragraphs[0].add_run()
    run.text = str(number)
    template.style(run, size=26, color=WHITE, font=template.HEADING_FONT,
                   bold=True)

    card = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(template.CONTENT_LEFT.inches + 0.9), Inches(top),
        Inches(CONTENT_WIDTH.inches - 0.9), Inches(_ROW_HEIGHT),
    )
    template.solid(card, PALE)
    tf = card.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.3)
    tf.margin_right = Inches(0.3)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    para = tf.paragraphs[0]
    run = para.add_run()
    run.text = heading
    template.style(run, size=19, color=PRIMARY, font=template.HEADING_FONT,
                   bold=True)
    run = para.add_run()
    run.text = f"  —  {detail}"
    template.style(run, size=15, color=TEXT)


def _checklist_slide(prs):
    slide = template.add_content_slide(prs, "Setup checklist")

    top = 1.45
    for number, (heading, detail) in enumerate(_CHECKLIST_ITEMS, start=1):
        _checklist_row(slide, number, top, heading, detail)
        top += _ROW_HEIGHT + _ROW_GAP

    template.add_band(
        slide, Inches(top + 0.1), Inches(1.0),
        "Done with an item? Post its number in the meeting chat: ✅ 1 … ✅ 4",
        size=18, fill=HIGHLIGHT,
        sub="Posted ✅ 1–4? You're fully set up — help a neighbor or "
            "explore your new tools.",
    )
    template.set_notes(slide, (
        "This slide stays on screen-share for the whole 30-minute block — "
        "it is the only hands-on part of Lesson 1, and everyone leaves "
        "with working tools. Read the four items aloud once, show the "
        "chat convention (✅ plus the item number, ✅ 1–4 when done), "
        "then let people work. The chat marks make stragglers visible "
        "without anyone having to speak up — watch the chat, not the "
        "gallery, and go to whoever has posted nothing after ten "
        "minutes.\n\n"
        "Quick finishers become co-troubleshooters: pair anyone posting "
        "✅ 1–4 early with someone who is stuck. Switch to the "
        "Troubleshooting slide only when the same failure hits more "
        "than one person; otherwise handle issues one-to-one and keep "
        "the checklist up."
    ))


# --- Slide 2: troubleshooting ------------------------------------------------

# The three typical failure points from lesson_1.md Segment 5, each with
# its fix: symptom as the card heading, fix as the card body.
_FAILURE_POINTS = [
    ("Copilot shows the wrong account — or no Copilot at all",
     "Fix: you are signed in with a personal Microsoft account. Sign out "
     "and sign back in with the institutional account."),
    ("Google account blocked in your default browser",
     "Fix: institutional policy is blocking the sign-in. Create a "
     "separate browser profile and sign in to Google there."),
    ("Gemini shows region or consent screens",
     "Fix: click through the consent prompts — they are one-time. If a "
     "block remains, flag it in the chat and we'll look together."),
]


def _troubleshooting_slide(prs):
    slide = template.add_content_slide(prs, "Troubleshooting")

    top = 1.6
    for symptom, fix in _FAILURE_POINTS:
        template.add_card(
            slide, template.CONTENT_LEFT, Inches(top), CONTENT_WIDTH,
            Inches(1.35), PALE_WARM, symptom, HIGHLIGHT, fix,
        )
        top += 1.55

    template.add_label(
        slide, template.CONTENT_LEFT, Inches(top), CONTENT_WIDTH,
        "Something else? Say so in the chat — that's what this block is for.",
        italic=True, align=PP_ALIGN.CENTER,
    )
    template.set_notes(slide, (
        "Show this slide only when a failure hits more than one person — "
        "otherwise keep the checklist up and fix issues one-to-one. These "
        "are the three typical failure points; walk the affected person "
        "through the fix on the card while others keep working.\n\n"
        "The account mix-ups dominate: a personal Microsoft account "
        "where the institutional one is needed blocks Copilot Chat "
        "entirely, and institutional policy can block the Google "
        "sign-in in the default browser — the separate browser profile "
        "sidesteps that. Region/consent screens in Gemini usually just "
        "need clicking through."
    ))


# --- Slide 3: catch-up path --------------------------------------------------


def _catchup_slide(prs):
    slide = template.add_content_slide(prs, "Not finished? Catch-up path")

    template.add_card(
        slide, template.CONTENT_LEFT, Inches(1.7), Inches(5.96), Inches(1.9),
        PALE, "Practice Week 1 drop-in (Week 3)", PRIMARY,
        "The catch-up venue: an optional facilitated session with no new "
        "material — bring whatever refused to install or sign in.",
    )
    template.add_card(
        slide, Inches(template.CONTENT_LEFT.inches + 6.17), Inches(1.7),
        Inches(5.96), Inches(1.9), PALE, "…or a 10-minute slot this week",
        PRIMARY,
        "Don't wait two weeks stuck: say so now (chat or aloud) and we "
        "book a 10-minute call this week to finish together.",
    )

    template.add_band(
        slide, Inches(4.1), Inches(1.15),
        "Not installed today — on purpose: Obsidian and Perfect Wiki",
        size=18,
        sub="Tools arrive when they're needed: Perfect Wiki in Lesson 3, "
            "Obsidian in Lesson 4. Today's three are enough.",
    )
    template.set_notes(slide, (
        "Show this in the last minutes of the block. Anyone without "
        "✅ 1–4 in the chat: flag them now, by name — don't let an "
        "unfinished setup surface silently in Lesson 2 when the "
        "assignment needs Gemini. The default catch-up venue is Practice "
        "Week 1's drop-in in Week 3, but that is two weeks away, so "
        "actively offer the 10-minute slot during this week and book it "
        "before the call ends.\n\n"
        "If anyone asks about the other course tools: Obsidian and "
        "Perfect Wiki are deliberately not installed today — each "
        "arrives in the lesson that first uses it (Perfect Wiki in "
        "Lesson 3, Obsidian in Lesson 4). Close the session by pointing "
        "at the assignment slide once more: everything this week "
        "happens in Gemini, which everyone now has."
    ))


def add_slides(prs):
    _checklist_slide(prs)
    _troubleshooting_slide(prs)
    _catchup_slide(prs)
