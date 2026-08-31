"""Segment 2, Part B — The 7 Levels of AI, the course's map (~5 min).

Two slides, per lesson_1.md Segment 2 Part B: the 7 Levels framework as an
ascending stair (Level 1 at the bottom, "You are (probably) here"), each
level with its research-management translation and course coverage; and the
KM System teaser diagram closing on "by Week 8, you will have built this".
Terminology follows CONTEXT.md (Wiki Sandbox, Markdown Generator, Personal
Vault, Round-Trip).
"""

from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches

from deck import template
from deck.template import (
    ACCENT, HEADING_FONT, HIGHLIGHT, MUTED, PALE, PALE_WARM, PRIMARY, WHITE,
)

# (level, name, research-management translation, coverage, chip treatment)
# — compressed from the lesson_1.md Part B table so each row stays one line.
# Chip treatments: "here" = the Level 1 marker, "course" = covered lessons,
# "outlook" = beyond this course.
_LEVELS = [
    (1, "Question Asker",
     "AI as a fancier search engine — one-off questions about calls, rules",
     "◀ You are (probably) here", "here"),
    (2, "Prompt Crafter",
     "structured prompts: instruction–context–constraints, examples",
     "Today + Lesson 2", "course"),
    (3, "Power User",
     "persistent context — workspaces that already know your domain",
     "Lessons 2–3", "course"),
    (4, "Tool Orchestrator",
     "the right specialized tool per step: NotebookLM, Perplexity",
     "Lessons 3–4", "course"),
    (5, "Builder",
     "your own tools without code — custom agents, pipelines",
     "Lesson 5 · Markdown Generator", "course"),
    (6, "Systems Builder",
     "systems that run without you — agentic workflows",
     "Outlook only — final session", "outlook"),
    (7, "AI-Run Operation",
     "aspirational: an agent workforce running an operation",
     "Mentioned for perspective", "outlook"),
]

_VIBE_CODING_ASIDE = (
    "Aside: “vibe coding” — describing software in plain language and letting "
    "AI write it — is reshaping who can build tools, but it is not part of "
    "this course."
)

_ROW_HEIGHT = 0.62
_ROW_PITCH = 0.7
_ASIDE_GAP = 0.42  # extra vertical slot between the Level 5 and Level 4 rows
_STAIR_STEP = 0.22  # left indent added per level
_BAR_RIGHT = 9.95
_CHIP_LEFT = 10.1
_CHIP_WIDTH = 2.63


def _level_row(slide, level, name, translation, coverage, treatment, top):
    indent = (level - 1) * _STAIR_STEP

    badge = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(template.CONTENT_LEFT.inches + indent), Inches(top),
        Inches(0.55), Inches(_ROW_HEIGHT),
    )
    template.solid(badge, PRIMARY)
    tf = badge.text_frame
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.paragraphs[0].alignment = PP_ALIGN.CENTER
    run = tf.paragraphs[0].add_run()
    run.text = str(level)
    template.style(run, size=20, color=WHITE, font=HEADING_FONT, bold=True)

    bar_left = template.CONTENT_LEFT.inches + indent + 0.67
    bar = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(bar_left), Inches(top),
        Inches(_BAR_RIGHT - bar_left), Inches(_ROW_HEIGHT),
    )
    template.solid(bar, PALE)
    tf = bar.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.18)
    tf.margin_right = Inches(0.12)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    para = tf.paragraphs[0]
    run = para.add_run()
    run.text = name
    template.style(run, size=15, color=PRIMARY, font=HEADING_FONT, bold=True)
    run = para.add_run()
    run.text = f"  —  {translation}"
    template.style(run, size=12, color=MUTED)

    chip_fill, chip_color, chip_bold = {
        "here": (HIGHLIGHT, WHITE, True),
        "course": (ACCENT, WHITE, False),
        "outlook": (PALE, MUTED, False),
    }[treatment]
    chip = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(_CHIP_LEFT), Inches(top),
        Inches(_CHIP_WIDTH), Inches(_ROW_HEIGHT),
    )
    template.solid(chip, chip_fill)
    tf = chip.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.12)
    tf.margin_right = Inches(0.12)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.paragraphs[0].alignment = PP_ALIGN.CENTER
    run = tf.paragraphs[0].add_run()
    run.text = coverage
    template.style(run, size=11, color=chip_color, font=HEADING_FONT,
                   bold=chip_bold)


def _seven_levels_slide(prs):
    slide = template.add_content_slide(prs, "The 7 Levels of AI — the course's map")

    top = 1.55
    for level, name, translation, coverage, treatment in reversed(_LEVELS):
        _level_row(slide, level, name, translation, coverage, treatment, top)
        top += _ROW_PITCH
        if level == 5:  # the one-sentence vibe-coding aside sits at rows 4–5
            box = template.add_label(
                slide, template.CONTENT_LEFT, Inches(top - 0.04),
                Inches(12.13), _VIBE_CODING_ASIDE, size=11, italic=True,
            )
            box.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
            top += _ASIDE_GAP

    template.set_notes(slide, (
        "~4 min. Pace: one sentence per level, no dwelling — the levels get "
        "lived out in the coming weeks, not explained today. Walk the stair "
        "bottom-up from Level 1; the orange chip says it out loud: you are "
        "(probably) here, and that is exactly where this course expects you. "
        "The blue chips are the course arc — today's session already starts "
        "the climb to Level 2; Level 5 is where you build the Markdown "
        "Generator. The grey levels 6–7 are outlook only.\n\n"
        "Vibe coding: exactly the one sentence printed at the Level 4–5 rows, "
        "then move on — this is not an engineering course; the closest we "
        "come to building is Lesson 5's Agent Builder, which needs no code "
        "at all.\n\n"
        "Close of segment: this course takes you from Level 1 to Level 4–5. "
        "Then advance to the destination slide — the teaser."
    ))


def _km_teaser_slide(prs):
    slide = template.add_content_slide(
        prs,
        "The destination: your KM System",
        [
            "This course takes you from Level 1 to Level 4–5. The destination "
            "artifact is a two-way knowledge management system:",
        ],
    )

    cards = [
        ("Wiki Sandbox",
         "Your own Perfect Wiki — the single source of truth for concepts, "
         "processes, guides, and projects", PALE, PRIMARY),
        ("Markdown Generator",
         "A custom Copilot agent that converts wiki content into Markdown "
         "notes — you build it in Lesson 5", PALE_WARM, HIGHLIGHT),
        ("Personal Vault",
         "Your local Obsidian vault — Graph View, Canvas: your visual "
         "thinking space", PALE, PRIMARY),
    ]
    card_width = 3.53
    arrow_gap = 0.77
    x = template.CONTENT_LEFT.inches
    for i, (heading, body, fill, heading_color) in enumerate(cards):
        template.add_card(slide, Inches(x), Inches(2.55), Inches(card_width),
                          Inches(1.95), fill, heading, heading_color, body)
        x += card_width
        if i < len(cards) - 1:
            template.add_label(slide, Inches(x), Inches(3.25), Inches(arrow_gap),
                               "→", size=28, color=ACCENT, align=PP_ALIGN.CENTER)
            x += arrow_gap

    template.add_label(
        slide, template.CONTENT_LEFT, Inches(4.75), Inches(12.13),
        "⟲  …and manual export back to the wiki closes the Round-Trip — a "
        "person pulls and exports; nothing syncs on its own.",
        italic=True, align=PP_ALIGN.CENTER,
    )
    template.add_band(
        slide, Inches(5.5), Inches(0.95),
        "By Week 8, you will have built this.",
        fill=HIGHLIGHT,
        sub="Built around one research project of yours — real or fictional — "
            "one lesson at a time.",
    )

    template.set_notes(slide, (
        "~1 min, teaser only — show, don't teach. Name the three parts and "
        "leave them as names: Wiki Sandbox (your own Perfect Wiki, single "
        "source of truth), Markdown Generator (a custom Copilot agent — "
        "built in Lesson 5, once the basics are in place), Personal Vault "
        "(your local Obsidian vault, the visual thinking space). One breath "
        "on the return arrow: the Round-Trip is human-mediated — pulling "
        "and exporting are deliberate acts, never automated sync.\n\n"
        "Land the line and let it sit: by Week 8, you will have built this. "
        "Then straight into Segment 3 — prompt engineering, where the climb "
        "to Level 2 starts right now."
    ))


def add_slides(prs):
    _seven_levels_slide(prs)
    _km_teaser_slide(prs)
