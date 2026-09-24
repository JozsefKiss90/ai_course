"""Segment 5 — Assignment briefing + close (8 min, Presentation + Q&A).

Four slides, per lesson_3.md Segment 5: the assignment slide as the plan
writes it (six items, the outcome); the chain — six steps, each naming
its tool, lane, input and output, only step 1 required and only when the
material is sensitive; effort and where the work is shown; and the close
(Practice Week 2, the next Teaching Session).

The chain crosses the Daily Driver and the local tools, so its rows are
colour-coded per step; the rest of the segment is neutral.
"""

from pptx.util import Inches

from deck.lesson3 import template
from deck.lesson3.template import (
    DAILY_DRIVER, HEADING_FONT, INK, LOCAL, PANEL, SLATE, TEXT,
)

_LEFT = template.CONTENT_LEFT
_WIDTH = template.CONTENT_WIDTH

# --- Assignment slide --------------------------------------------------------

_ITEMS = [
    ("1", "Notebook",
     "at least 5 sources, anonymised or public; one Deep Research run or "
     "Perplexity search that adds a source you did not have, with its citation"),
    ("2", "Vault",
     "at least 5 concepts, 2 processes, 1 guide, 1 project note, each produced "
     "from a cited notebook answer, each with a filled source line and "
     "status: checked after you have compared the text with the citation"),
    ("3", "Links",
     "at least 10 [[links]], every note linked at least once; the project "
     "note links to every process. Bring a Graph View screenshot"),
    ("4", "Smart Connections",
     "one relationship the Connections view surfaced that you had not "
     "linked; add the link and write one line on why you had missed it"),
    ("5", "Copilot for Obsidian",
     "one Vault QA question about your project, the cited answer, and one "
     "line on whether the cited notes actually support it"),
    ("6", "AI-use log", "tool, date and purpose per step, as in Lesson 2"),
]


def _assignment_slide(prs):
    slide = template.add_content_slide(
        prs, "Assignment: your KM System, first two floors")
    template.add_label(
        slide, _LEFT, Inches(1.28), _WIDTH,
        "Populate your Personal Vault from your Anchor Project, through your "
        "notebook, and put the vault's own AI to work.",
        size=12.5, color=INK, bold=True,
    )
    top = 1.75
    heights = [0.62, 0.78, 0.62, 0.62, 0.62, 0.42]
    for (number, name, detail), height in zip(_ITEMS, heights):
        template.add_badge(slide, _LEFT, Inches(top + (height - 0.42) / 2),
                           Inches(0.42), number)
        card = template.add_open_card(slide, _LEFT + Inches(0.55), Inches(top),
                                      _WIDTH - Inches(0.55), Inches(height), SLATE)
        card.text_frame.margin_top = 0
        card.text_frame.margin_bottom = 0
        para = card.text_frame.paragraphs[0]
        run = para.add_run()
        run.text = f"{name}:  "
        template.style(run, size=11.5, color=INK, font=HEADING_FONT, bold=True)
        run = para.add_run()
        run.text = detail
        template.style(run, size=11, color=TEXT)
        top += height + 0.07
    template.add_band(
        slide, Inches(6.05), Inches(0.62),
        "Outcome: Level 4 reached — the right specialised tool per step, "
        "chained, with the outputs kept somewhere they outlive the chat",
        size=13,
    )
    template.set_notes(slide, (
        "~3 min. Read the six items as they stand; this slide text is pasted "
        "into the Teams chat after the session. Populate your Personal Vault "
        "from your Anchor Project, through your notebook, and put the "
        "vault's own AI to work.\n\n"
        "Outcome: Level 4 reached — the right specialised tool per step, "
        "chained, with the outputs kept somewhere they outlive the chat. "
        "Level 3 was the Lesson 2 assignment; today's recap map opened with "
        "that row."
    ))
    return slide


# --- How the chain reads -----------------------------------------------------

# The plan's chain table: (#, step, tool, lane, in, out, edge colour).
_CHAIN = [
    ("1", "Gather and anonymise sources", "No AI tool; Word or Excel, by hand",
     "None (manual)", "Your Anchor Project documents",
     "Files that may be uploaded → step 2", SLATE),
    ("2", "Fill the gap", "Deep Research in the notebook, or Perplexity",
     "Daily Driver", "One question your sources do not answer",
     "One cited source → step 3", DAILY_DRIVER.color),
    ("3", "Ground", "NotebookLM", "Daily Driver",
     "The files from 1 and the source from 2",
     "Cited answers per entity type → step 4", DAILY_DRIVER.color),
    ("4", "Keep", "Obsidian", "Local", "The cited answers",
     "Notes with source lines and links → steps 5, 6", LOCAL.color),
    ("5", "Find", "Smart Connections", "Local, offline", "The vault",
     "Missed relationships → back into step 4 as links", LOCAL.color),
    ("6", "Ask", "Copilot for Obsidian, Vault QA", "Daily Driver",
     "The vault and one question", "A cited answer, checked against the notes",
     DAILY_DRIVER.color),
]


def _chain_slide(prs):
    slide = template.add_content_slide(
        prs, "How the chain reads: every step names its input and its output")
    columns = [
        (Inches(0.45), "#"), (Inches(2.05), "Step"), (Inches(2.75), "Tool"),
        (Inches(1.45), "Lane"), (Inches(2.65), "In"), (Inches(2.78), "Out"),
    ]
    rows = [[n, step, tool, lane, inp, out] for n, step, tool, lane, inp, out, _ in _CHAIN]
    template.add_table(
        slide, Inches(1.32), columns, rows, Inches(0.58), size=10,
        row_edges=[edge for *_, edge in _CHAIN], first_bold=False,
    )
    template.add_detail_card(
        slide, _LEFT, Inches(5.65), Inches(5.95), Inches(1.0), SLATE,
        "Only step 1 is required in the strict sense",
        ["and only when your material contains sensitive data; every other "
         "step is the assignment itself"],
    )
    template.add_detail_card(
        slide, _LEFT + Inches(6.18), Inches(5.65), Inches(5.95), Inches(1.0), SLATE,
        "The gate that is not a step, from Lesson 2",
        ["check every output before it becomes the next step's input, and "
         "log it"],
    )
    template.set_notes(slide, (
        "~2 min. The same discipline as the Lesson 2 blueprint: every step "
        "names its input and its output, and the lane it runs in. Read the "
        "Out column: each output is the next step's input, and step 5's "
        "output goes back into step 4 as links.\n\n"
        "Only step 1 is required in the strict sense, and only when your "
        "material contains sensitive data; every other step is the "
        "assignment itself. The gate that is not a step, from Lesson 2, "
        "applies: check every output before it becomes the next step's "
        "input, and log it.\n\n"
        "Colour: amber steps are cloud calls (Daily Driver), green steps stay "
        "on the machine, slate is no AI tool at all."
    ))
    return slide


# --- Effort and where the work is shown --------------------------------------


def _effort_slide(prs):
    slide = template.add_content_slide(prs, "Effort, and where the work is shown")
    template.add_card(
        slide, _LEFT, Inches(1.45), Inches(3.9), Inches(1.6), PANEL,
        "3–4 hours across the two weeks", INK,
        "The vault notes are the bulk; the plugin steps are minutes once "
        "installed.",
    )
    template.add_list_card(
        slide, _LEFT + Inches(4.12), Inches(1.45), Inches(8.01), Inches(2.55), PANEL,
        "Done looks like",
        ["A vault whose Graph View has no isolated nodes",
         "Nine or more checked notes with source lines",
         "One added link you owe to Smart Connections",
         "One cited Vault QA answer you have judged",
         "An AI-use log"],
        size=12.5,
    )
    template.add_detail_card(
        slide, _LEFT, Inches(3.2), Inches(3.9), Inches(0.8), SLATE,
        "nine or more · no isolated nodes",
        ["the two numbers to remember"],
    )
    template.add_card(
        slide, _LEFT, Inches(4.25), Inches(5.95), Inches(2.0), PANEL,
        "Practice Week 2's drop-in (W6)", INK,
        "Troubleshooting: installation, plugin setup, linking, free-tier "
        "surprises.",
    )
    template.add_card(
        slide, _LEFT + Inches(6.18), Inches(4.25), Inches(5.95), Inches(2.0), PANEL,
        "The next Teaching Session's debrief (W7)", INK,
        "2–3 volunteers walk through notebook → vault → Vault QA.",
    )
    template.set_notes(slide, (
        "~1 min. Expect 3–4 hours across the two weeks: the vault notes are "
        "the bulk, the plugin steps are minutes once installed. Done looks "
        "like a vault whose Graph View has no isolated nodes, nine or more "
        "checked notes with source lines, one added link you owe to Smart "
        "Connections, one cited Vault QA answer you have judged, and a "
        "log.\n\n"
        "Show the work at Practice Week 2's drop-in (W6) for "
        "troubleshooting, and at the next Teaching Session's debrief (W7): "
        "2–3 volunteers walk through notebook → vault → Vault QA."
    ))
    return slide


# --- Close -------------------------------------------------------------------


def _close_slide(prs):
    slide = template.add_content_slide(
        prs, "Close: Practice Week 2 and the next Teaching Session")
    template.add_card(
        slide, _LEFT, Inches(1.5), Inches(5.95), Inches(2.6), PANEL,
        "Practice Week 2 (W6)", INK,
        "Optional facilitated drop-in, no new material. Bring installation "
        "problems, plugin setup, linking questions and free-tier surprises.",
    )
    template.add_card(
        slide, _LEFT + Inches(6.18), Inches(1.5), Inches(5.95), Inches(2.6), PANEL,
        "Next Teaching Session (W7)", INK,
        "Builds on the vault as source of truth. The detailed plan follows "
        "the redesign of the remaining lessons and is announced in Teams "
        "before W6.",
    )
    template.add_band(
        slide, Inches(4.4), Inches(1.3),
        "Two floors of your KM System are standing: a source-grounded "
        "notebook and a linked vault",
        size=17,
        sub="capture in the notebook, keep in the vault, query the vault — "
            "see you at the drop-in",
    )
    template.set_notes(slide, (
        "~1 min. Practice Week 2 (W6): optional facilitated drop-in, no new "
        "material; bring installation problems, plugin setup, linking "
        "questions and free-tier surprises. Next Teaching Session (W7): "
        "builds on the vault as source of truth; the detailed plan follows "
        "the redesign of the remaining lessons and is announced in Teams "
        "before W6.\n\n"
        "After the session: paste the assignment slide text and the "
        "hands-on checklist into the Teams chat."
    ))
    return slide


def add_slides(prs):
    _assignment_slide(prs)
    _chain_slide(prs)
    _effort_slide(prs)
    _close_slide(prs)
