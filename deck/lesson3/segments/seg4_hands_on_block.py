"""Segment 4 — Hands-on block: notebook + vault (25 min, Hands-on,
facilitated).

Three slides, per lesson_3.md Segment 4: the checklist the facilitator
screen-shares, split at the plan's own line — steps 1–5 are the minimum
working system, steps 6–8 the vault's own AI — each step with its "done
when"; and the troubleshooting slide with the typical failure points and
their fixes. Participants mark progress in the Teams chat ("✅ 1–4").

The block spans both lanes and the local tools, so the slides carry no
single lane; each step's edge bar says where that step runs (amber: a
cloud call; green: on the machine; slate: neither).
"""

from pptx.util import Inches

from deck.lesson3 import template
from deck.lesson3.template import (
    DAILY_DRIVER, INK, LOCAL, MUTED, SLATE, TEXT,
)

_LEFT = template.CONTENT_LEFT
_WIDTH = template.CONTENT_WIDTH

# The plan's checklist: (number, step, done when, edge colour).
_STEPS = [
    ("1", "Create a notebook for your Anchor Project; add at least three "
          "anonymised or public sources (the Lesson 2 blueprint's step 2 "
          "files qualify)",
     "Sources listed, notebook answers a test question with a citation",
     DAILY_DRIVER.color),
    ("2", "Ask one concept question and one process question in the shapes "
          "from Segment 2; open one citation each",
     "Two cited answers on screen", DAILY_DRIVER.color),
    ("3", "Install Obsidian; create a vault named after your Anchor Project; "
          "make the four folders and the template",
     "Vault open, folders visible", LOCAL.color),
    ("4", "Paste the two answers as two notes in the right folders; fill "
          "type, source, captured; leave status: draft",
     "Two notes with frontmatter", LOCAL.color),
    ("5", "Link them to each other with [[ ]]; create one more note by "
          "clicking a link to a title that does not exist yet",
     "Graph View shows three connected nodes", LOCAL.color),
    ("6", "Turn off Restricted mode; install and enable Smart Connections; "
          "open the Connections view on one note",
     "Connections view lists the other notes", LOCAL.color),
    ("7", "Get a Gemini API key at Google AI Studio; install Copilot for "
          "Obsidian; enter the key; refresh the index",
     "Settings saved without an error", DAILY_DRIVER.color),
    ("8", "In Copilot for Obsidian's Vault QA, ask a question your three "
          "notes can answer",
     "An answer citing at least one note", DAILY_DRIVER.color),
]

_STEP_W = 7.6
_DONE_W = 3.8
_ROW_H = 0.7
_ROW_PITCH = 0.8


def _checklist_rows(slide, steps, top):
    template.add_label(slide, _LEFT + Inches(0.65), Inches(top - 0.32),
                       Inches(_STEP_W), "STEP", size=10, color=MUTED, bold=True)
    template.add_label(slide, _LEFT + Inches(0.65 + _STEP_W + 0.1), Inches(top - 0.32),
                       Inches(_DONE_W), "DONE WHEN", size=10, color=MUTED, bold=True)
    for number, step, done, edge in steps:
        template.add_badge(slide, _LEFT, Inches(top + 0.1), Inches(0.5), number)
        card = template.add_open_card(slide, _LEFT + Inches(0.65), Inches(top),
                                      Inches(_STEP_W), Inches(_ROW_H), edge)
        run = card.text_frame.paragraphs[0].add_run()
        run.text = step
        template.style(run, size=11.5, color=TEXT)
        card = template.add_open_card(slide, _LEFT + Inches(0.65 + _STEP_W + 0.1),
                                      Inches(top), Inches(_DONE_W), Inches(_ROW_H), edge)
        run = card.text_frame.paragraphs[0].add_run()
        run.text = done
        template.style(run, size=11, color=INK, bold=True)
        top += _ROW_PITCH


def _checklist_a_slide(prs):
    slide = template.add_content_slide(
        prs, "Checklist, steps 1–5: the minimum working system")
    _checklist_rows(slide, _STEPS[:5], 1.75)
    template.add_band(
        slide, Inches(5.85), Inches(0.85),
        "Steps 1–5 are the minimum: anyone at step 5 by the 20-minute mark "
        "has a working system",
        size=14,
        sub="mark progress in the Teams chat (“✅ 1–4”) so stragglers are "
            "visible  ·  quick finishers help",
    )
    template.set_notes(slide, (
        "25 min in all. Screen-share this slide; everyone builds the "
        "smallest working version of their KM System on their own Anchor "
        "Project. Participants post progress in the Teams chat as “✅ 1–4” "
        "so stragglers are visible; quick finishers help.\n\n"
        "Steps 1–5 are the minimum; anyone at step 5 by the 20-minute mark "
        "has a working system. Steps 6–8 (next slide) finish in the "
        "assignment or at Practice Week 2 if time runs out.\n\n"
        "Participants who bring no anonymised material use the "
        "facilitator's sample documents — prepared, fictional. Keep the "
        "Lesson 2 “what must never enter Gemini” list on the second monitor; "
        "the vault inherits it. Say the free-tier chat cap (50 a day) once: "
        "far above what the block needs, but nobody should burn questions "
        "on chit-chat."
    ))
    return slide


def _checklist_b_slide(prs):
    slide = template.add_content_slide(
        prs, "Checklist, steps 6–8: the vault's own AI")
    _checklist_rows(slide, _STEPS[5:], 1.75)
    template.add_band(
        slide, Inches(4.3), Inches(0.85),
        "Steps 6–8 finish in the assignment or at Practice Week 2 if time runs out",
        size=14, fill=SLATE,
        sub="the setup order from Segment 3 is the recipe; the key goes to "
            "Google's free API tier, so Copilot for Obsidian is a Daily Driver call",
    )
    template.add_label(
        slide, _LEFT, Inches(5.4), _WIDTH,
        "Amber steps call a cloud model (Daily Driver); green steps stay on "
        "the machine.",
        size=11, color=MUTED, italic=True,
    )
    template.set_notes(slide, (
        "Steps 6–8 follow Segment 3's setup order. Smart Connections' first "
        "embedding run is a one-off and can take a while on a slow machine — "
        "wait for it. Copilot for Obsidian needs the Gemini API key from "
        "Google AI Studio on the personal Google account and a refresh of "
        "the vault index before Vault QA answers anything.\n\n"
        "If time runs out these three finish in the assignment or at "
        "Practice Week 2; do not hold the room for them."
    ))
    return slide


# The plan's typical failure points and their fixes.
_FAILURES = [
    ("Obsidian install blocked on a managed machine",
     "portable build, or a personal device"),
    ("Google account in the wrong browser profile", "same fix as Lesson 1"),
    ("API key pasted with a trailing space", "paste it again"),
    ("Copilot for Obsidian's index refresh not run, so Vault QA answers "
     "“no relevant notes”", "run “Refresh vault index”"),
    ("Smart Connections still embedding on a slow machine",
     "wait — it is a one-off"),
]


def _trouble_slide(prs):
    slide = template.add_content_slide(
        prs, "If it does not work — the usual reasons")
    template.add_table(
        slide, Inches(1.42),
        [(Inches(7.0), "Symptom"), (Inches(5.13), "Fix")],
        [[a, b] for a, b in _FAILURES], Inches(0.62), size=12, edge=SLATE,
    )
    template.add_detail_card(
        slide, _LEFT, Inches(5.6), Inches(5.95), Inches(1.0), SLATE,
        "No anonymised material of your own?",
        ["Use the facilitator's sample documents — prepared, fictional."],
    )
    template.add_detail_card(
        slide, _LEFT + Inches(6.18), Inches(5.6), Inches(5.95), Inches(1.0), SLATE,
        "On the second monitor",
        ["The Lesson 2 “what must never enter Gemini” list — the vault "
         "inherits it, and so does the notebook."],
    )
    template.set_notes(slide, (
        "Keep this slide handy rather than on screen; the checklist is the "
        "one to share. The five symptoms are the plan's typical failure "
        "points, each with its fix; all five come up at Practice Week 2 too. "
        "Anyone without anonymised material uses the sample documents."
    ))
    return slide


def add_slides(prs):
    _checklist_a_slide(prs)
    _checklist_b_slide(prs)
    _trouble_slide(prs)
