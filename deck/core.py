"""Shared deck mechanics — everything about producing a PowerPoint file that
does not depend on which lesson is being built.

A lesson hands the core a `Design` (fonts, theme colours, canvas, footer)
and gets back a 16:9 presentation whose theme part and slide master carry
that design. The per-lesson template module owns the design constants and
the slide builders on top of these helpers; nothing here knows a lesson's
palette or layout.

Contents: presentation setup at 16:9, writing fonts and colours into the
theme part, the master footer, the safe master shape-id helper (the
near-2^31 slide-layout-id gotcha), the flat/outlined shape styling
helpers, run styling, and the blank-slide, text-box and notes helpers.
"""

from dataclasses import dataclass, field

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.util import Emu, Inches, Pt

SLIDE_WIDTH = Emu(12192000)  # 13.333 in — 16:9 for Teams
SLIDE_HEIGHT = Emu(6858000)  # 7.5 in

# Footer geometry — the same strip on every lesson's master.
FOOTER_LEFT = Inches(0.6)
FOOTER_WIDTH = Inches(12.13)
FOOTER_RULE_TOP = Inches(7.08)
FOOTER_TEXT_TOP = Inches(7.13)
FOOTER_TEXT_SIZE = 10

# Ids at or above this are the stock template's near-2^31 slide-layout ids
# (p:sldLayoutId), not shapes; see master_next_id().
_STOCK_ID_FLOOR = 0x7F000000

_A = "http://schemas.openxmlformats.org/drawingml/2006/main"
_P = "http://schemas.openxmlformats.org/presentationml/2006/main"


@dataclass(frozen=True)
class Design:
    """What a lesson's template hands the core: the parts of a design that
    live in the theme part and on the slide master.

    scheme maps theme colour names (dk2, accent1, accent2, ...) to the
    RGBColor written into the theme's colour scheme; names absent from the
    stock scheme are ignored.
    """

    heading_font: str
    body_font: str
    canvas: RGBColor
    footer_text: str
    footer_rule: RGBColor
    footer_color: RGBColor
    scheme: dict = field(default_factory=dict)


# --- Presentation setup -----------------------------------------------------


def new_presentation(design):
    """A 16:9 presentation carrying the design's theme, canvas, and footer."""
    prs = Presentation()
    prs.slide_width = SLIDE_WIDTH
    prs.slide_height = SLIDE_HEIGHT
    _apply_theme(prs, design)
    master = prs.slide_master
    master.background.fill.solid()
    master.background.fill.fore_color.rgb = design.canvas
    _add_master_footer(prs, design)
    return prs


def _apply_theme(prs, design):
    """Write the design's fonts and colour scheme into the master's theme part."""
    theme_part = prs.slide_master.part.part_related_by(RT.THEME)
    root = etree.fromstring(theme_part.blob)
    for tag, font in (("majorFont", design.heading_font), ("minorFont", design.body_font)):
        latin = root.find(f".//{{{_A}}}{tag}/{{{_A}}}latin")
        latin.set("typeface", font)
    scheme = root.find(f".//{{{_A}}}clrScheme")
    for name, color in design.scheme.items():
        srgb = scheme.find(f"{{{_A}}}{name}/{{{_A}}}srgbClr")
        if srgb is not None:
            srgb.set("val", str(color))
    theme_part._blob = etree.tostring(
        root, xml_declaration=True, encoding="UTF-8", standalone=True
    )


# --- Shape styling ----------------------------------------------------------


def solid(shape, color):
    """Flat fill with no outline or shadow."""
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    shape.shadow.inherit = False


def outlined(shape, fill, border, weight=1.0):
    """Solid fill with a thin border, no shadow."""
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = border
    shape.line.width = Pt(weight)
    shape.shadow.inherit = False


def add_bar(slide, left, top, width, height, color):
    """A plain filled rectangle — accent bars, rules, edge caps."""
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    solid(bar, color)
    return bar


def style(run, size, color, font, bold=False, italic=False):
    run.font.name = font
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.font.bold = bold
    run.font.italic = italic


# --- Master shapes ----------------------------------------------------------


def master_next_id(master):
    """Next free shape id on the master.

    python-pptx's own `shapes._next_shape_id` scans every id attribute in
    the master part, which includes the stock template's slide-layout ids
    (p:sldLayoutId, numbered from 2^31). A shape given the next id after
    those exceeds the signed 32-bit range and PowerPoint reports the file
    as corrupt. This helper looks at shape ids (p:cNvPr) only."""
    ids = [
        int(el.get("id"))
        for el in master.shapes._spTree.iter(f"{{{_P}}}cNvPr")
    ]
    return max((i for i in ids if i < _STOCK_ID_FLOOR), default=1) + 1


def _master_add_shape(master, sp):
    """Append a shape element to the master's shape tree (python-pptx's
    MasterShapes collection has no add_shape/add_textbox of its own)."""
    from pptx.shapes.autoshape import Shape

    master.shapes._spTree.append(sp)
    return Shape(sp, master.shapes)


def _add_master_footer(prs, design):
    """Footer rule + text on the slide master, inherited by every slide."""
    from pptx.oxml.shapes.autoshape import CT_Shape

    master = prs.slide_master

    sp = CT_Shape.new_autoshape_sp(
        master_next_id(master), "Footer Rule", "rect",
        FOOTER_LEFT, FOOTER_RULE_TOP, FOOTER_WIDTH, Pt(1.2),
    )
    line = _master_add_shape(master, sp)
    solid(line, design.footer_rule)

    sp = CT_Shape.new_textbox_sp(
        master_next_id(master), "Footer Text",
        FOOTER_LEFT, FOOTER_TEXT_TOP, FOOTER_WIDTH, Inches(0.3),
    )
    box = _master_add_shape(master, sp)
    tf = box.text_frame
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    run = tf.paragraphs[0].add_run()
    run.text = design.footer_text
    style(run, size=FOOTER_TEXT_SIZE, color=design.footer_color, font=design.body_font)


# --- Slide-level helpers ----------------------------------------------------


def blank_slide(prs):
    return prs.slides.add_slide(prs.slide_master.slide_layouts[6])


def add_text(slide, left, top, width, height, align=PP_ALIGN.LEFT, anchor=None):
    """A word-wrapping text box; returns its text frame."""
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    if anchor is not None:
        tf.vertical_anchor = anchor
    tf.paragraphs[0].alignment = align
    return tf


def set_notes(slide, notes):
    slide.notes_slide.notes_text_frame.text = notes
