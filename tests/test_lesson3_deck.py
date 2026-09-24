"""Tests for the Lesson 3 deck (`.scratch/lesson-3-deck/`).

Every assertion runs against a saved-and-reopened .pptx, so it checks what
PowerPoint will load. lesson_3.md is the oracle: the Timing Overview for
segment names, durations and modes, and its tables for the recap map, the
free-tier limits, the Studio outputs, the question shapes, the plugin
comparison, the hands-on checklist and the assignment chain.
"""

import os
import re
import tempfile
from pathlib import Path

import pytest
from pptx import Presentation

from deck.lesson3.build import build_presentation


def _slide_text(slide):
    """All visible text on a slide, joined."""
    return "\n".join(
        shape.text_frame.text for shape in slide.shapes if shape.has_text_frame
    )


def _slide_corpus(slide):
    """Everything a slide says: its visible text and its speaker notes."""
    text = _slide_text(slide)
    if slide.has_notes_slide:
        text += "\n" + slide.notes_slide.notes_text_frame.text
    return text


@pytest.fixture(scope="module")
def reopened():
    """Build the deck, save it, and reopen it the way PowerPoint would."""
    prs = build_presentation()
    fd, path = tempfile.mkstemp(suffix=".pptx")
    os.close(fd)
    try:
        prs.save(path)
        yield Presentation(path)
    finally:
        os.unlink(path)


# --- Title slide ------------------------------------------------------------


def test_title_slide_names_the_lesson_course_week_and_length(reopened):
    text = _slide_text(reopened.slides[0])
    assert "Lesson 3 — Knowledge Capture: NotebookLM and the Personal Vault" in text
    assert "ai for research management" in text.lower()
    assert "Week 4" in text
    assert "90 minutes" in text


# --- Section dividers ---------------------------------------------------------

# The Timing Overview table from lesson_3.md — transcribed, so the copy is
# independent of the build.
TIMING_OVERVIEW = [
    ("Recap + assignment debrief", "15 min", "Participant demos, facilitated"),
    ("The KM System redrawn + NotebookLM in depth", "20 min",
     "Presentation + live facilitator demo"),
    ("Obsidian: the Personal Vault and its two plugins", "20 min",
     "Live facilitator demo"),
    ("Hands-on block: notebook + vault", "25 min", "Hands-on, facilitated"),
    ("Assignment briefing + close", "8 min", "Presentation + Q&A"),
]
BUFFER_MINUTES = 2

SEG1_NAME, SEG2_NAME, SEG3_NAME, SEG4_NAME, SEG5_NAME = [
    name for name, _, _ in TIMING_OVERVIEW
]


def test_each_segment_has_a_divider_with_duration_and_mode(reopened):
    all_slides = [_slide_text(s) for s in reopened.slides]
    for number, (name, duration, mode) in enumerate(TIMING_OVERVIEW, start=1):
        dividers = [t for t in all_slides if name in t and f"{duration} · {mode}" in t]
        assert dividers, f"no divider found for segment {number} ({name})"
        assert any(f"Segment {number} of 5" in t for t in dividers)


def test_dividers_appear_in_session_order_after_the_title(reopened):
    all_slides = [_slide_text(s) for s in reopened.slides]
    positions = [
        next(i for i, t in enumerate(all_slides) if name in t)
        for name, _, _ in TIMING_OVERVIEW
    ]
    assert positions == sorted(positions)
    assert positions[0] > 0


# --- Segments as slide ranges ------------------------------------------------

SEGMENT_NAMES = [name for name, _, _ in TIMING_OVERVIEW]


def _segment_ranges(reopened):
    """Map divider name -> the slides between that divider and the next."""
    all_slides = [_slide_text(s) for s in reopened.slides]
    starts = {
        name: next(i for i, t in enumerate(all_slides) if name in t)
        for name in SEGMENT_NAMES
    }
    ordered = sorted(starts.items(), key=lambda kv: kv[1])
    ranges = {}
    for (name, start), nxt in zip(ordered, [s for _, s in ordered[1:]] + [len(all_slides)]):
        ranges[name] = list(reopened.slides)[start + 1 : nxt]
    return ranges


def _dividers(reopened):
    """The section dividers: the slide immediately before each segment's
    first content slide."""
    all_slides = list(reopened.slides)
    return [
        all_slides[all_slides.index(slides[0]) - 1]
        for slides in _segment_ranges(reopened).values()
    ]


def _segment_corpus(reopened, name):
    """One segment's divider and slides, each as slide text plus notes."""
    slides = _segment_ranges(reopened)[name]
    all_slides = list(reopened.slides)
    divider = all_slides[all_slides.index(slides[0]) - 1]
    return [_slide_corpus(s) for s in [divider, *slides]]


def _slide_title(slide):
    """A content slide's title — the first text box add_content_slide adds."""
    for shape in slide.shapes:
        if shape.has_text_frame:
            return shape.text_frame.text
    return ""


def _segment_slide(reopened, segment, title):
    """The slide inside one segment whose title is exactly this."""
    matches = [
        s for s in _segment_ranges(reopened)[segment]
        if _slide_title(s) == title
    ]
    assert len(matches) == 1, f"{title!r}: {len(matches)} slides matched"
    return matches[0]


def _shape_containing(slide, token):
    """The single shape whose text contains `token` (case-sensitive)."""
    matches = [
        sh for sh in slide.shapes
        if sh.has_text_frame and token in sh.text_frame.text
    ]
    assert len(matches) == 1, f"{token!r}: {len(matches)} shapes matched"
    return matches[0]


def _assert_titles_in_order(reopened, segment, titles):
    slides = _segment_ranges(reopened)[segment]
    found = [_slide_title(s) for s in slides]
    positions = []
    for title in titles:
        assert title in found, f"{title!r} is not in {segment!r}: {found}"
        positions.append(found.index(title))
    assert positions == sorted(positions), f"{segment}: slides out of plan order"


# --- The plan as oracle -------------------------------------------------------

PLAN_PATH = Path(__file__).resolve().parents[1] / "lesson_3.md"
CONTEXT_PATH = PLAN_PATH.parent / "CONTEXT.md"


def _plan_section(heading):
    """The body of one section of lesson_3.md, up to the next heading of
    the same or a higher level."""
    text = PLAN_PATH.read_text(encoding="utf-8")
    start = text.index(heading)
    level = len(heading) - len(heading.lstrip("#"))
    rest = text[start + len(heading):]
    stop = re.search(r"\n#{1,%d} " % level, rest)
    return rest[: stop.start()] if stop else rest


def _plan_table(section_heading):
    """The rows of the first Markdown table in a plan section, header
    row excluded, as lists of cell strings."""
    rows = []
    for line in _plan_section(section_heading).splitlines():
        line = line.strip()
        if not line.startswith("|") or set(line) <= set("|-: "):
            continue
        rows.append([c.strip() for c in line.strip("|").split("|")])
    return rows[1:]


def _segment_refs(text):
    """The segment numbers a piece of text points forward to — "Segment 2",
    "Segments 2–3", "Segments 3, 5"."""
    refs = set()
    for chunk in re.findall(r"Segments?\s+(\d+(?:\s*(?:·|,|and|–|-)\s*\d+)*)", text):
        numbers = [int(d) for d in re.findall(r"\d+", chunk)]
        if "–" in chunk or "-" in chunk:
            refs.update(range(numbers[0], numbers[-1] + 1))
        else:
            refs.update(numbers)
    return refs


def _strip_md(cell):
    return re.sub(r"[*`]", "", cell)


# --- Master, theme and the Lesson 3 look -------------------------------------


def _theme_xml(reopened):
    from pptx.opc.constants import RELATIONSHIP_TYPE as RT

    return reopened.slide_masters[0].part.part_related_by(RT.THEME).blob.decode("utf-8")


def test_master_defines_the_lesson_3_footer_and_theme_fonts(reopened):
    from deck.lesson3 import template

    master = reopened.slide_masters[0]
    master_text = "\n".join(
        s.text_frame.text for s in master.shapes if s.has_text_frame
    )
    assert "Lesson 3 · Week 4" in master_text
    assert "Lesson 1" not in master_text and "Lesson 2" not in master_text
    theme_xml = _theme_xml(reopened)
    assert f'typeface="{template.HEADING_FONT}"' in theme_xml
    assert f'typeface="{template.BODY_FONT}"' in theme_xml


def test_look_is_neither_lesson_1s_nor_lesson_2s(reopened):
    """Different fonts from both earlier decks, none of their theme
    colours, and a canvas that is neither Lesson 1's navy nor Lesson 2's
    warm paper."""
    from deck import template as lesson1
    from deck.lesson2 import template as lesson2
    from deck.lesson3 import template as lesson3

    theme_xml = _theme_xml(reopened)
    assert "Segoe" not in theme_xml
    assert "Cambria" not in theme_xml
    for color in (lesson1.PRIMARY, lesson1.ACCENT, lesson1.HIGHLIGHT,
                  lesson2.SENSITIVE_LANE.color, lesson2.DAILY_DRIVER.color):
        assert str(color) not in theme_xml
    canvas = reopened.slide_masters[0].background.fill.fore_color.rgb
    assert canvas not in (lesson1.CANVAS, lesson2.CANVAS)
    assert canvas == lesson3.CANVAS
    assert {lesson3.HEADING_FONT, lesson3.BODY_FONT}.isdisjoint(
        {lesson2.HEADING_FONT, lesson2.BODY_FONT}
    )


def test_fonts_and_sizes_are_explicit_everywhere(reopened):
    from deck.lesson3 import template

    for i, slide in enumerate(reopened.slides):
        for shape in slide.shapes:
            if not shape.has_text_frame:
                continue
            for para in shape.text_frame.paragraphs:
                for run in para.runs:
                    assert run.font.name in (template.HEADING_FONT, template.BODY_FONT), (
                        f"slide {i + 1}: font {run.font.name!r} in {run.text!r}"
                    )
                    assert run.font.size is not None, (
                        f"slide {i + 1}: inherited size in {run.text!r}"
                    )


# --- Builders keep the Lesson 2 call shapes ----------------------------------


def test_builders_accept_the_lesson_2_call_shapes_and_survive_a_roundtrip():
    from pptx.util import Inches

    from deck.lesson3 import template

    prs = template.new_presentation()
    slide = template.add_content_slide(
        prs, "Sample", ["one", (1, "two")], notes="n", lane=template.DAILY_DRIVER,
    )
    template.add_demo_cue(slide, "open NotebookLM", Inches(1.6), lane=template.DAILY_DRIVER)
    template.add_band(slide, Inches(2.4), Inches(0.8), "Band", sub="sub")
    template.add_card(slide, Inches(0.6), Inches(3.5), Inches(4.0), Inches(1.5),
                      template.DAILY_TINT, "Card", template.INK, "body")
    template.add_card(slide, Inches(4.8), Inches(3.5), Inches(4.0), Inches(1.5),
                      template.LOCAL.color, "Label block", template.WHITE)
    template.add_list_card(slide, Inches(9.0), Inches(3.5), Inches(3.5), Inches(1.5),
                           template.LOCAL_TINT, "List", ["a", "b"])
    template.add_detail_card(slide, Inches(0.6), Inches(5.2), Inches(4.0), Inches(1.2),
                             template.LOCAL.color, "Lead", ["detail"])
    template.add_label(slide, Inches(4.8), Inches(5.3), Inches(4.0), "caption")
    template.add_placeholder_slide(prs, "Seg", "Planned", "gist", "NN")

    fd, path = tempfile.mkstemp(suffix=".pptx")
    os.close(fd)
    try:
        prs.save(path)
        reopened = Presentation(path)
    finally:
        os.unlink(path)
    text = _slide_text(reopened.slides[0])
    for expected in ("Sample", "one", "two", "LIVE DEMO — open NotebookLM", "Band",
                     "sub", "Card", "body", "Label block", "List", "a", "Lead",
                     "detail", "caption", "DAILY DRIVER"):
        assert expected in text, expected
    assert reopened.slides[0].notes_slide.notes_text_frame.text == "n"
    assert "[PLACEHOLDER]" in _slide_text(reopened.slides[1])


# --- Segment 1: Recap + assignment debrief (ticket 02) -----------------------

RECAP_TITLE = "Lessons 1–2 → Lesson 3: the recap map"
DEBRIEF_TITLE = "Assignment debrief — the next 12 minutes"
CAPTURE_TITLE = "Common problems — captured live"
SEG1_TITLES = [RECAP_TITLE, DEBRIEF_TITLE, CAPTURE_TITLE]

# One (concept token, link token) pair per row of the plan's recap table,
# in the plan's order. Each token is asserted to appear in its plan cell
# first, so the pairs cannot drift from the plan.
RECAP_TOKENS = [
    ("7 Levels", "Level 4 is today's target"),
    ("KM System teaser", "single source of truth"),
    ("More context is not better context", "source-grounded notebook"),
    ("Hallucination", "verification step"),
    ("two-lane rule", "call a cloud model"),
    ("Gems", "organises sources"),
    ("Deep Research in Gemini", "straight into the sources"),
    ("Gemini Canvas", "different thing"),
    ("Microsoft Copilot Chat (Basic)", "unrelated community plugin"),
    ("NotebookLM preview", "walkthrough"),
    ("Lesson 2 blueprint", "notebook output is vault input"),
]


def _recap_rows_from_plan():
    return [(c, l) for c, l in _plan_table("### Part A: The recap map")]


def test_seg1_slides_present_in_order(reopened):
    _assert_titles_in_order(reopened, SEG1_NAME, SEG1_TITLES)


def test_recap_map_pairs_every_plan_concept_with_its_lesson_3_link(reopened):
    """Every row of the plan's Part A table is on the slide as a concept to
    the left of the Lesson 3 topic it leads into, keeping the plan's
    forward references, in the plan's order."""
    slide = _segment_slide(reopened, SEG1_NAME, RECAP_TITLE)
    plan_rows = _recap_rows_from_plan()
    assert len(plan_rows) == len(RECAP_TOKENS)

    tops = []
    for (concept_cell, link_cell), (concept, link) in zip(plan_rows, RECAP_TOKENS):
        assert concept in concept_cell, f"{concept!r} not in the plan's concept cell"
        assert link in link_cell, f"{link!r} not in the plan's leads-into cell"
        concept_shape = _shape_containing(slide, concept)
        link_shape = _shape_containing(slide, link)
        assert concept_shape.top == link_shape.top, f"{concept!r}: link not alongside"
        assert concept_shape.left < link_shape.left, f"{concept!r}: link not to its right"
        planned = _segment_refs(link_cell)
        assert planned <= _segment_refs(link_shape.text_frame.text), (
            f"{concept!r}: slide points to "
            f"{_segment_refs(link_shape.text_frame.text)}, plan to {planned}"
        )
        tops.append(concept_shape.top)
    assert tops == sorted(tops), "recap rows are not in the plan's order"


DEFINITION_MARKERS = [
    " is a ", " is an ", " is the ", " are the ", " means ", " refers to ",
    " stands for ", "in other words", "that is, ", "which is ",
]


def test_recap_map_restates_no_earlier_definition(reopened):
    slide = _segment_slide(reopened, SEG1_NAME, RECAP_TITLE)
    for shape in slide.shapes:
        if not shape.has_text_frame:
            continue
        text = shape.text_frame.text
        if RECAP_TITLE in text:
            continue
        lowered = text.lower()
        for marker in DEFINITION_MARKERS:
            assert marker not in lowered, f"{marker!r} in {text!r}"
        assert not re.search(r"\.\s+[A-Z“]", text), f"explanation prose in {text!r}"


def test_debrief_slide_states_the_format_the_three_items_and_the_questions(reopened):
    text = _slide_text(_segment_slide(reopened, SEG1_NAME, DEBRIEF_TITLE))
    assert "2–3 volunteers" in text
    assert "4 minutes" in text
    assert "12 minutes" in text
    assert "Practice Week 2" in text
    lowered = text.lower()
    positions = [lowered.index(item) for item in ("blueprint", "presentation", "ai-use log")]
    assert positions == sorted(positions), "the three items are out of plan order"
    for question in ("which steps did you skip", "learning outcome", "rejected at the check"):
        assert question in lowered, f"question {question!r} missing"


def test_capture_slide_tags_the_three_items_and_routes_the_problems(reopened):
    slide = _segment_slide(reopened, SEG1_NAME, CAPTURE_TITLE)
    text = _slide_text(slide)
    for tag in ("Blueprint", "Presentation", "Log"):
        assert tag in text, f"capture lane {tag!r} missing"
    assert _segment_refs(text) >= {2, 3}, "the routing loses a segment"
    assert "free-tier" in text.lower()
    assert "Practice Week 2" in text
    blank = [
        sh for sh in slide.shapes
        if sh.has_text_frame and not sh.text_frame.text.strip()
    ]
    assert len(blank) >= 3, "no blank capture lanes to write in"


def test_seg1_notes_carry_the_time_keeping_script_and_the_routing_check(reopened):
    notes = " ".join(
        _segment_slide(reopened, SEG1_NAME, t).notes_slide.notes_text_frame.text
        for t in SEG1_TITLES
    ).lower()
    assert "3-minute" in notes
    assert "cut at 4" in notes
    assert "practice week 2" in notes
    assert "backup" in notes
    assert "anonymise-before-upload" in notes


# --- Segment 2: The KM System redrawn + NotebookLM in depth (ticket 03) ------

SEG2_KM_TITLE = "The KM System, redrawn"
SEG2_FLOORS_TITLE = "The four-floor map"
SEG2_WHAT_TITLE = "NotebookLM (Gemini Notebook): what it is"
SEG2_LIMITS_TITLE = "Free tier at the time of writing (September 2026)"
SEG2_GROUNDING_TITLE = "Grounding versus exploration: when to use which"
SEG2_STUDIO_TITLE = "Studio outputs: one notebook, many formats"
SEG2_SHAPES_TITLE = "From notebook to note: the four question shapes"
SEG2_PERPLEXITY_TITLE = "Perplexity, in passing"
SEG2_DEMO_TITLE = "Demo script — six steps, fictional material"
SEG2_TITLES = [
    SEG2_KM_TITLE, SEG2_FLOORS_TITLE, SEG2_WHAT_TITLE, SEG2_LIMITS_TITLE,
    SEG2_GROUNDING_TITLE, SEG2_STUDIO_TITLE, SEG2_SHAPES_TITLE,
    SEG2_PERPLEXITY_TITLE, SEG2_DEMO_TITLE,
]

ENTITY_TYPES = ["concept", "process", "guide", "project"]


def _fills(slide):
    fills = set()
    for shape in slide.shapes:
        try:
            if shape.fill.type is not None and shape.fill.fore_color.rgb:
                fills.add(str(shape.fill.fore_color.rgb))
        except (AttributeError, TypeError):
            continue
    return fills


def _seg2_slide(reopened, title):
    return _segment_slide(reopened, SEG2_NAME, title)


def test_seg2_slides_present_in_plan_order(reopened):
    _assert_titles_in_order(reopened, SEG2_NAME, SEG2_TITLES)


def test_km_system_slide_draws_notebook_to_vault_with_a_person_between(reopened):
    """Sources → NotebookLM (the distiller) → Obsidian (the source of truth);
    the arrow into the vault is a person, there is no automated sync, and
    the four entity types survive as the vault's four folders."""
    slide = _seg2_slide(reopened, SEG2_KM_TITLE)
    text = _slide_text(slide)
    notebook = _shape_containing(slide, "the distiller")
    vault = _shape_containing(slide, "source of truth")
    assert notebook.top < vault.top, "the notebook does not sit above the vault"
    for token in ("NotebookLM", "Obsidian", "read, check", "source line",
                  "No automated sync", "Smart Connections", "Copilot for Obsidian"):
        assert token in text, f"{token!r} missing from the KM System slide"
    for entity in ENTITY_TYPES:
        assert entity in text.lower(), f"entity type {entity!r} missing"
    assert "paid" in _slide_corpus(slide).lower(), "why the wiki is gone is unsaid"


def test_four_floor_slide_puts_obsidian_on_the_fourth_floor(reopened):
    plan = _plan_section("### The KM System, redrawn")
    text = _slide_text(_seg2_slide(reopened, SEG2_FLOORS_TITLE))
    for token in ("grounding", "frontier-model", "specialist", "tools"):
        assert token in plan
        assert token in text.lower(), f"floor {token!r} missing"
    assert "Obsidian" in text
    assert "Lesson 2" in text, "floors 2 and 3 are not credited to Lesson 2"
    assert "the edge is never the tool" in text.lower()


def test_the_rename_is_said_once_and_notebooklm_is_the_name_afterwards(reopened):
    slides_text = [_slide_text(s) for s in reopened.slides]
    assert sum(t.count("Gemini Notebook") for t in slides_text) == 1
    assert sum(t.count("NotebookLM") for t in slides_text) > 10


def test_what_it_is_slide_states_grounding_citations_inputs_and_the_lane(reopened):
    plan = _plan_section("### NotebookLM: what it is")
    text = _slide_text(_seg2_slide(reopened, SEG2_WHAT_TITLE))
    assert "only" in text and "citation" in text.lower()
    for source in ("PDF", "Google Docs", "Markdown", "web page", "YouTube",
                   "audio", "images", "CSV"):
        assert source in plan
        assert source.lower() in text.lower(), f"source type {source!r} missing"
    assert "500,000 words or 200 MB" in text
    assert "Daily Driver" in text
    assert "never enter Gemini" in text


def test_free_tier_slide_matches_the_plan_table_and_is_dated(reopened):
    rows = _plan_table("### NotebookLM: what it is")
    assert len(rows) == 6
    slide = _seg2_slide(reopened, SEG2_LIMITS_TITLE)
    text = _slide_text(slide)
    for label, value in rows:
        assert _strip_md(label) in text, f"limit {label!r} missing"
        number = re.search(r"\d+", value)
        if number:
            assert number.group() in text, f"{label!r}: {number.group()} missing"
    assert "daily generation cap" in text
    assert "September 2026" in text
    assert "re-verif" in _slide_corpus(slide).lower()


def test_grounding_slide_gives_the_three_situations_and_the_one_liner(reopened):
    text = _slide_text(_seg2_slide(reopened, SEG2_GROUNDING_TITLE))
    for token in ("drowning in material", "Precision matters", "move across formats"):
        assert token in text, f"situation {token!r} missing"
    assert "stays in Gemini" in text
    assert "a notebook organises evidence; a Gem organises behaviour" in text
    assert "re-verif" in text.lower(), "the notebook-to-Gem flag is missing"


def test_studio_outputs_slide_matches_the_plan_table_in_order(reopened):
    rows = _plan_table("### Studio outputs")
    assert len(rows) == 7
    text = _slide_text(_seg2_slide(reopened, SEG2_STUDIO_TITLE))
    positions = []
    for output, _ in rows:
        name = re.match(r"\*\*(.+?)\*\*", output).group(1)
        assert name in text, f"Studio output {name!r} missing"
        positions.append(text.index(name))
    assert positions == sorted(positions)
    assert "10 a month" in text
    assert "speaker" in text.lower(), "the learning-science point is unattributed"
    assert "diffuse" in text.lower()
    assert "MIT" not in text, "the preliminary study is on the slide"


def test_question_shapes_slide_gives_one_shape_per_entity_type(reopened):
    rows = _plan_table("### From notebook to note")
    assert [_strip_md(r[0]).lower() for r in rows] == ENTITY_TYPES
    text = _slide_text(_seg2_slide(reopened, SEG2_SHAPES_TITLE))
    openers = ("Define", "List the steps", "Write a how-to", "Summarise the project")
    for (entity, shape), opener in zip(rows, openers):
        assert opener in shape
        assert opener in text, f"{entity}: question shape missing"
    assert "source line" in text
    assert "An answer without a citation does not become a note" in text


def test_perplexity_slide_gives_both_benefits_and_the_routing_line(reopened):
    text = _slide_text(_seg2_slide(reopened, SEG2_PERPLEXITY_TITLE))
    assert "research management" in text.lower()
    assert "knowledge management" in text.lower()
    assert "find → ground → keep" in text
    assert "Spaces" in text and "paid" in text.lower()
    assert "Daily Driver" in text
    assert "not on the km system's path" in text.lower()


def test_seg2_demo_script_has_the_six_steps_in_order_and_is_cued(reopened):
    slide = _seg2_slide(reopened, SEG2_DEMO_TITLE)
    text = _slide_text(slide)
    steps = ["Create a notebook", "concept question", "briefing", "Mind map",
             "Audio Overview", "Deep Research and Perplexity"]
    positions = [text.index(step) for step in steps]
    assert positions == sorted(positions)
    assert "fictional" in text
    assert "50-source" in text
    assert "LIVE DEMO" in text
    assert "live" in slide.notes_slide.notes_text_frame.text.lower()


def test_every_seg2_slide_carries_the_daily_driver_colour(reopened):
    from deck.lesson3 import template

    amber = str(template.DAILY_DRIVER.color)
    slides = _segment_ranges(reopened)[SEG2_NAME]
    divider = list(reopened.slides)[list(reopened.slides).index(slides[0]) - 1]
    for slide in [divider, *slides]:
        assert amber in _fills(slide), f"{_slide_title(slide)!r}: no Daily Driver colour"


# --- Segment 3: Obsidian, the Personal Vault and its two plugins (ticket 04) --

SEG3_WHAT_TITLE = "What Obsidian is"
SEG3_VAULT_TITLE = "The vault: four folders and one note template"
SEG3_LINKS_TITLE = "Links, backlinks, tags — and the linking discipline"
SEG3_GRAPH_TITLE = "Graph View and Canvas"
SEG3_PLUGINS_TITLE = "Two community plugins: the vault's own AI"
SEG3_WHY_TITLE = "Why two plugins"
SEG3_SETUP_TITLE = "Setup order for the hands-on block"
SEG3_DEMO_TITLE = "Demo script — six steps, one prepared vault"
SEG3_TITLES = [
    SEG3_WHAT_TITLE, SEG3_VAULT_TITLE, SEG3_LINKS_TITLE, SEG3_GRAPH_TITLE,
    SEG3_PLUGINS_TITLE, SEG3_WHY_TITLE, SEG3_SETUP_TITLE, SEG3_DEMO_TITLE,
]


def _seg3_slide(reopened, title):
    return _segment_slide(reopened, SEG3_NAME, title)


def test_seg3_slides_present_in_plan_order(reopened):
    _assert_titles_in_order(reopened, SEG3_NAME, SEG3_TITLES)


def test_what_obsidian_is_slide_states_local_free_fourth_floor_and_the_install_flag(reopened):
    slide = _seg3_slide(reopened, SEG3_WHAT_TITLE)
    text = _slide_text(slide)
    for token in ("plain Markdown files", "No account", "Free for personal use",
                  "Sync", "Publish", "not used", "fourth floor", "thread",
                  "administrator", "portable build"):
        assert token in text, f"{token!r} missing"
    assert "Nothing in Obsidian requires an account" in _slide_corpus(slide)


def test_vault_slide_shows_the_four_folders_and_the_template_fields(reopened):
    plan = _plan_section("### The vault: structure, notes, links, tags")
    text = _slide_text(_seg3_slide(reopened, SEG3_VAULT_TITLE))
    for folder in ("concepts/", "processes/", "guides/", "projects/", "_templates/"):
        assert folder in plan and folder in text, f"folder {folder!r} missing"
    for field in ("type:", "source:", "captured:", "status:"):
        assert field in text, f"frontmatter field {field!r} missing"
    assert "draft" in text and "checked" in text
    assert "## Related" in text
    assert "human gate" in text


def test_links_slide_states_the_mechanics_and_the_one_rule(reopened):
    text = _slide_text(_seg3_slide(reopened, SEG3_LINKS_TITLE))
    for token in ("[[Note title]]", "Backlinks", "Tags", "#eligible-cost",
                  "a note with no links is a file, not knowledge"):
        assert token in text, f"{token!r} missing"


def test_graph_slide_resolves_the_canvas_name_clash(reopened):
    text = _slide_text(_seg3_slide(reopened, SEG3_GRAPH_TITLE))
    for token in ("Graph View", "local graph", "Canvas", "Gemini Canvas", "whiteboard"):
        assert token in text, f"{token!r} missing"


PLUGIN_ROW_LABELS = [
    "Author", "What it does", "How it works", "Cost", "Which model", "Lane",
    "Job in the KM System", "Name clash",
]


def _shape_with_text(slide, text):
    matches = [
        sh for sh in slide.shapes
        if sh.has_text_frame and sh.text_frame.text.strip() == text
    ]
    assert len(matches) == 1, f"{text!r}: {len(matches)} shapes"
    return matches[0]


def test_plugin_slide_carries_every_row_of_the_plan_table_and_colours_the_columns(reopened):
    from deck.lesson3 import template

    rows = _plan_table("### Two community plugins")
    assert [_strip_md(r[0]) for r in rows] == PLUGIN_ROW_LABELS
    slide = _seg3_slide(reopened, SEG3_PLUGINS_TITLE)
    text = _slide_text(slide)
    for label in PLUGIN_ROW_LABELS:
        assert label in text, f"row {label!r} missing"
    for token in ("Restricted mode", "Brian Petro", "logancyang", "Lookup",
                  "Vault QA", "no API key", "Google AI Studio",
                  "Not Microsoft Copilot", "never the Sensitive Lane",
                  "not used", "is ever in the vault"):
        assert token in text, f"{token!r} missing"
    sc = _shape_with_text(slide, "Smart Connections")
    cfo = _shape_with_text(slide, "Copilot for Obsidian")
    assert str(sc.fill.fore_color.rgb) == str(template.LOCAL.color)
    assert str(cfo.fill.fore_color.rgb) == str(template.DAILY_DRIVER.color)
    assert sc.left < cfo.left


def test_why_two_plugins_slide_closes_the_loop(reopened):
    text = _slide_text(_seg3_slide(reopened, SEG3_WHY_TITLE))
    assert "different questions" in text
    assert "thirty notes" in text
    assert "capture in the notebook, keep in the vault, query the vault" in text


def test_setup_order_slide_has_the_five_steps_in_order(reopened):
    text = _slide_text(_seg3_slide(reopened, SEG3_SETUP_TITLE))
    steps = ["turn off Restricted mode", "Install and enable Smart Connections",
             "Get a Gemini API key", "Install and enable Copilot for Obsidian",
             "switch to Vault QA"]
    positions = [text.index(step) for step in steps]
    assert positions == sorted(positions)
    assert "Refresh vault index" in text


def test_seg3_demo_script_has_the_six_steps_in_order_and_is_cued(reopened):
    slide = _seg3_slide(reopened, SEG3_DEMO_TITLE)
    text = _slide_text(slide)
    steps = ["Create a vault", "Paste the Segment 2 briefing", "two concept notes",
             "Graph View", "Smart Connections on the prepared vault", "Vault QA"]
    positions = [text.index(step) for step in steps]
    assert positions == sorted(positions)
    assert "~25 linked notes" in text
    assert "Which deadlines apply to the periodic report" in text
    assert "LIVE DEMO" in text
    assert "live" in slide.notes_slide.notes_text_frame.text.lower()


def test_every_seg3_slide_carries_the_local_colour(reopened):
    from deck.lesson3 import template

    green, amber = str(template.LOCAL.color), str(template.DAILY_DRIVER.color)
    slides = _segment_ranges(reopened)[SEG3_NAME]
    divider = list(reopened.slides)[list(reopened.slides).index(slides[0]) - 1]
    for slide in [divider, *slides]:
        assert green in _fills(slide), f"{_slide_title(slide)!r}: no local colour"
    for title in (SEG3_PLUGINS_TITLE, SEG3_DEMO_TITLE):
        assert amber in _fills(_seg3_slide(reopened, title)), (
            f"{title!r}: the Copilot for Obsidian cloud call is not amber"
        )


# --- Segment 4: Hands-on block (ticket 05) ------------------------------------

SEG4_CHECKLIST_A_TITLE = "Checklist, steps 1–5: the minimum working system"
SEG4_CHECKLIST_B_TITLE = "Checklist, steps 6–8: the vault's own AI"
SEG4_TROUBLE_TITLE = "If it does not work — the usual reasons"
SEG4_TITLES = [SEG4_CHECKLIST_A_TITLE, SEG4_CHECKLIST_B_TITLE, SEG4_TROUBLE_TITLE]


def _seg4_slide(reopened, title):
    return _segment_slide(reopened, SEG4_NAME, title)


def test_seg4_slides_present_in_plan_order(reopened):
    _assert_titles_in_order(reopened, SEG4_NAME, SEG4_TITLES)


# One token from each step's "Done when" cell in the plan's checklist,
# asserted against the plan first.
CHECKLIST_DONE_TOKENS = [
    "answers a test question with a citation", "Two cited answers on screen",
    "Vault open, folders visible", "Two notes with frontmatter",
    "three connected nodes", "Connections view lists the other notes",
    "Settings saved without an error", "citing at least one note",
]


def test_checklist_slides_carry_every_step_with_its_done_when_in_order(reopened):
    rows = _plan_table("## Segment 4")
    assert len(rows) == 8
    text = "\n".join(
        _slide_text(_seg4_slide(reopened, t))
        for t in (SEG4_CHECKLIST_A_TITLE, SEG4_CHECKLIST_B_TITLE)
    )
    positions = []
    for (number, step, done), token in zip(rows, CHECKLIST_DONE_TOKENS):
        assert token in done, f"step {number}: {token!r} is not the plan's done-when"
        assert token in text, f"step {number}: done-when {token!r} missing"
        positions.append(text.index(token))
    assert positions == sorted(positions)
    for token in ("three anonymised or public sources", "one concept question",
                  "one process question", "four folders", "status: draft",
                  "[[ ]]", "Restricted mode", "Smart Connections",
                  "Google AI Studio", "Copilot for Obsidian", "Vault QA"):
        assert token in text, f"checklist drops {token!r}"


def test_checklist_marks_the_minimum_and_the_progress_convention(reopened):
    text = _slide_text(_seg4_slide(reopened, SEG4_CHECKLIST_A_TITLE))
    assert "minimum" in text.lower()
    assert "20-minute mark" in text
    assert "✅" in text
    text_b = _slide_text(_seg4_slide(reopened, SEG4_CHECKLIST_B_TITLE))
    assert "Practice Week 2" in text_b


FAILURE_POINTS = [
    "managed machine", "browser profile", "trailing space", "index refresh",
    "still embedding",
]


def test_trouble_slide_lists_the_plan_s_failure_points_and_fixes(reopened):
    plan = _plan_section("## Segment 4")
    slide = _seg4_slide(reopened, SEG4_TROUBLE_TITLE)
    text = _slide_text(slide)
    for token in FAILURE_POINTS:
        assert token in plan, f"{token!r} is not a failure point the plan names"
        assert token in text, f"failure point {token!r} missing"
    corpus = _slide_corpus(slide)
    assert "sample documents" in corpus
    assert "never enter Gemini" in corpus


def test_the_chat_cap_is_said_once_in_the_hands_on_block(reopened):
    """The plan's Segment 4 note: say the 50-a-day cap once, so nobody burns
    questions on chit-chat — and only once, in the block's own notes."""
    assert "50 a day" in _slide_corpus(_seg4_slide(reopened, SEG4_CHECKLIST_A_TITLE))
    saying_it = [i + 1 for i, s in enumerate(reopened.slides) if "50 a day" in _slide_corpus(s)]
    assert len(saying_it) == 1, f"the chat cap is said on slides {saying_it}"


# --- Segment 5: Assignment briefing + close (ticket 06) -----------------------

SEG5_ASSIGNMENT_TITLE = "Assignment: your KM System, first two floors"
SEG5_CHAIN_TITLE = "How the chain reads: every step names its input and its output"
SEG5_EFFORT_TITLE = "Effort, and where the work is shown"
SEG5_CLOSE_TITLE = "Close: Practice Week 2 and the next Teaching Session"
SEG5_TITLES = [SEG5_ASSIGNMENT_TITLE, SEG5_CHAIN_TITLE, SEG5_EFFORT_TITLE, SEG5_CLOSE_TITLE]


def _seg5_slide(reopened, title):
    return _segment_slide(reopened, SEG5_NAME, title)


def test_seg5_slides_present_in_plan_order(reopened):
    _assert_titles_in_order(reopened, SEG5_NAME, SEG5_TITLES)


ASSIGNMENT_ITEMS = ["Notebook", "Vault", "Links", "Smart Connections",
                    "Copilot for Obsidian", "AI-use log"]


def test_assignment_slide_carries_the_six_items_and_the_outcome(reopened):
    plan = _plan_section("### Assignment slide")
    text = _slide_text(_seg5_slide(reopened, SEG5_ASSIGNMENT_TITLE))
    positions = []
    for i, item in enumerate(ASSIGNMENT_ITEMS, start=1):
        assert f"{i}. **{item}" in plan, f"{item!r} is not item {i} of the plan"
        assert f"{item}:" in text, f"assignment item {item!r} missing"
        positions.append(text.index(f"{item}:"))
    assert positions == sorted(positions)
    for token in ("at least 5 sources", "5 concepts, 2 processes, 1 guide, 1 project note",
                  "status: checked", "at least 10", "Graph View screenshot",
                  "one line on why you had missed it", "Vault QA",
                  "tool, date and purpose per step"):
        assert token in text, f"assignment drops {token!r}"
    assert "Level 4 reached" in text


def _chain_steps_from_plan():
    rows = _plan_table("### How the chain reads")
    return [dict(zip(("#", "Step", "Tool", "Lane", "In", "Out"), r)) for r in rows]


def test_chain_slide_carries_every_step_with_tool_lane_input_and_output(reopened):
    steps = _chain_steps_from_plan()
    assert [s["Step"] for s in steps] == [
        "Gather and anonymise sources", "Fill the gap", "Ground", "Keep", "Find", "Ask",
    ]
    text = _slide_text(_seg5_slide(reopened, SEG5_CHAIN_TITLE))
    positions = []
    for step in steps:
        assert step["Step"] in text, f"step {step['#']} missing"
        positions.append(text.index(step["Step"]))
        for field in ("Tool", "Lane"):
            assert step[field].split(",")[0].split(" in ")[0] in text, (
                f"step {step['#']}: {field} {step[field]!r} missing"
            )
    assert positions == sorted(positions)
    for token in ("No AI tool", "None (manual)", "Local, offline", "Daily Driver",
                  "Only step 1 is required", "sensitive data",
                  "check every output before it becomes the next step's input"):
        assert token in text, f"chain slide drops {token!r}"
    assert text.count("→ step") + text.count("→ steps") >= 5, "outputs do not say where they go"


def test_effort_slide_states_the_effort_what_done_looks_like_and_both_venues(reopened):
    plan = _plan_section("### Effort and where the work is shown")
    text = _slide_text(_seg5_slide(reopened, SEG5_EFFORT_TITLE))
    for token in ("3–4 hours", "no isolated nodes", "nine or more", "W6", "W7",
                  "2–3 volunteers", "notebook → vault → Vault QA"):
        assert token in plan, f"{token!r} is not in the plan's effort section"
        assert token in text, f"the effort slide drops {token!r}"
    assert "Practice Week 2" in text


def test_close_names_both_venues_in_context_terms(reopened):
    context = CONTEXT_PATH.read_text(encoding="utf-8")
    text = _slide_text(_seg5_slide(reopened, SEG5_CLOSE_TITLE))
    for term in ("Practice Week", "Teaching Session", "KM System"):
        assert f"**{term}**" in context, f"{term!r} is not a CONTEXT.md term"
        assert term in text, f"the close does not use {term!r}"
    lowered = text.lower()
    for token in ("w6", "w7", "drop-in", "no new material", "optional",
                  "source of truth", "announced in teams"):
        assert token in lowered, f"the close drops {token!r}"


def test_seg5_chain_colours_lanes_and_local_steps_apart(reopened):
    from deck.lesson3 import template

    fills = _fills(_seg5_slide(reopened, SEG5_CHAIN_TITLE))
    assert str(template.DAILY_DRIVER.color) in fills
    assert str(template.LOCAL.color) in fills


# --- Integration pass: flow, timing, colour, terminology (ticket 07) ---------


def test_no_placeholder_slides_remain_anywhere(reopened):
    for i, slide in enumerate(reopened.slides):
        assert "[PLACEHOLDER]" not in _slide_text(slide), f"slide {i + 1}"


# Slides per segment, reviewed against the Timing Overview durations and
# pinned. Growing a segment past its pin means changing lesson_3.md first.
EXPECTED_SEGMENT_SLIDE_COUNTS = {
    SEG1_NAME: 3, SEG2_NAME: 9, SEG3_NAME: 8, SEG4_NAME: 3, SEG5_NAME: 4,
}


def test_total_slide_count(reopened):
    expected = 1 + len(TIMING_OVERVIEW) + sum(EXPECTED_SEGMENT_SLIDE_COUNTS.values())
    assert len(reopened.slides) == expected


def test_segment_slide_counts_fit_their_timings(reopened):
    minutes = {name: int(duration.split()[0]) for name, duration, _ in TIMING_OVERVIEW}
    for name, slides in _segment_ranges(reopened).items():
        expected = EXPECTED_SEGMENT_SLIDE_COUNTS[name]
        assert len(slides) == expected, f"{name}: {len(slides)} slides, expected {expected}"
        assert len(slides) <= minutes[name], f"{name}: over one slide per minute"


def test_the_session_fits_ninety_minutes_with_the_buffer(reopened):
    """The durations the dividers show plus the plan's 2-minute buffer add
    up to the length the title slide promises."""
    from deck.lesson3 import build

    shown = [
        int(re.search(r"(\d+) min", _slide_text(divider)).group(1))
        for divider in _dividers(reopened)
    ]
    assert len(shown) == len(TIMING_OVERVIEW)
    promised = int(re.search(r"(\d+) minutes", _slide_text(reopened.slides[0])).group(1))
    assert build.BUFFER_MINUTES == BUFFER_MINUTES
    assert sum(shown) + BUFFER_MINUTES == promised, f"the dividers add up to {sum(shown)}"
    assert promised == 90


def test_every_content_slide_has_speaker_notes(reopened):
    dividers = _dividers(reopened)
    for i, slide in enumerate(reopened.slides):
        if i == 0 or slide in dividers:
            continue
        assert slide.has_notes_slide, f"slide {i + 1} has no notes slide"
        assert slide.notes_slide.notes_text_frame.text.strip(), (
            f"slide {i + 1} has empty speaker notes"
        )


def test_every_live_demo_cue_sits_in_a_demo_segment_and_is_mirrored_in_notes(reopened):
    """The two demo segments (2 and 3) are the only ones that open a live
    tool; each cue is repeated in that slide's notes."""
    ranges = _segment_ranges(reopened)
    demo_slides = ranges[SEG2_NAME] + ranges[SEG3_NAME]
    cued = [(i, s) for i, s in enumerate(reopened.slides) if "LIVE DEMO" in _slide_text(s)]
    # Segment 2: steps 1–2, 3–5 and the script; Segment 3: steps 1–2, 3, 4
    # and the script. Step 6 of Segment 2 is slide-only by the plan.
    assert len(cued) == 7, f"{len(cued)} live-demo cues; the plan's demo steps need 7"
    for i, slide in cued:
        assert slide in demo_slides, f"slide {i + 1}: a live demo outside the demo segments"
        notes = slide.notes_slide.notes_text_frame.text
        assert re.search(r"\blive\b", notes, re.IGNORECASE), (
            f"slide {i + 1}: the demo is cued on the slide but not in its notes"
        )


def test_no_youtube_link_anywhere_in_the_deck(reopened):
    """The talk consulted while drafting is never linked; the word stays
    allowed because YouTube videos with captions are a NotebookLM source."""
    for i, slide in enumerate(reopened.slides):
        lowered = _slide_corpus(slide).lower()
        for marker in ("youtube.com", "youtu.be", "m.youtube"):
            assert marker not in lowered, f"{marker!r} on slide {i + 1}"


def test_one_spelling_of_anonymise_deck_wide(reopened):
    for i, slide in enumerate(reopened.slides):
        assert not re.search(r"anonymiz", _slide_corpus(slide), re.IGNORECASE), (
            f"slide {i + 1}: the other spelling of anonymise"
        )


# Forbidden synonyms from the CONTEXT.md Avoid-lists — the Lesson 2 list
# plus the two Lesson 3 additions (the retired wiki's Avoid-list applies
# to the KM System; "knowledge base" is what the plan's Gem line avoids).
FORBIDDEN_TERMS = [
    "homework week", "self-study week", "blank week", "blank lesson",
    "breakout team", "domain team",
    "knowledge base", "shared wiki",
    "use case", "case study", "team project",
    "home base", "primary AI environment",
    "workhorse", "main tool",
    "tool ecosystem", "satellite tool",
]


def test_terminology_audit_no_forbidden_synonyms(reopened):
    context = CONTEXT_PATH.read_text(encoding="utf-8").lower()
    for term in FORBIDDEN_TERMS:
        assert term.lower() in context, f"{term!r} is not on a CONTEXT.md Avoid-list"
    for i, slide in enumerate(reopened.slides):
        corpus = _slide_corpus(slide).lower()
        for term in FORBIDDEN_TERMS:
            assert term.lower() not in corpus, f"{term!r} on slide {i + 1}"


def test_copilot_for_obsidian_is_never_bare_copilot(reopened):
    """CONTEXT.md's Avoid-list for the vault plugins: never "Copilot"
    without the qualifier. The only other Copilot in the deck is Lesson
    2's Copilot Chat (Basic), and Microsoft Copilot in the name-clash line."""
    allowed = re.compile(r"Copilot (?:for Obsidian|Chat \(Basic\)|Plus)|Microsoft Copilot")
    for i, slide in enumerate(reopened.slides):
        corpus = _slide_corpus(slide)
        stripped = allowed.sub("", corpus)
        assert "Copilot" not in stripped, f"slide {i + 1}: bare Copilot in {corpus!r}"


CLOUD_TELLS = ("notebooklm", "notebook", "api key", "copilot for obsidian", "perplexity", "daily driver",
               "deep research", "gemini", "cloud")
LOCAL_TELLS = ("obsidian", "smart connections", "local", "on the machine")


def test_lane_colours_appear_only_where_the_slide_names_that_kind_of_tool(reopened):
    """Amber means a cloud call and green means local, deck-wide: a slide
    carrying either colour names a tool of that kind somewhere on it. The
    Sensitive Lane has no colour in this deck, so Lesson 2's teal never
    appears."""
    from deck.lesson2 import template as lesson2
    from deck.lesson3 import template

    amber, green = str(template.DAILY_DRIVER.color), str(template.LOCAL.color)
    for i, slide in enumerate(reopened.slides):
        fills = _fills(slide)
        lowered = _slide_text(slide).lower()
        if amber in fills:
            assert any(t in lowered for t in CLOUD_TELLS), f"slide {i + 1}: amber, no cloud tool"
        if green in fills:
            assert any(t in lowered for t in LOCAL_TELLS), f"slide {i + 1}: green, no local tool"
        assert str(lesson2.SENSITIVE_LANE.color) not in fills, f"slide {i + 1}: Lesson 2's teal"
