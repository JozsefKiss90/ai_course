"""Design template for the Lesson 1 deck — the single source of visual truth.

Every slide in the deck is created through the helpers here, so fonts,
colors, and the footer stay uniform no matter which segment module adds
the slide. The theme (fonts + color scheme) is written into the slide
master's theme part, and the footer is a master shape, so both apply to
all slides.

Visual language: the whole deck lives on a deep-navy canvas (set as the
master background). Content sits on elevated navy panels with a colored
edge bar — azure for neutral, orange for emphasis. Type is white and
light-azure; the accents are vivid azure, orange, and amber. Title and
divider slides go full-bleed navy with amber accents and progress pills.
"""

from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt

from deck import core
# The lesson-independent mechanics live in deck.core; these names are
# re-exported so segment modules keep calling template.solid(...) etc.
from deck.core import (  # noqa: F401
    SLIDE_HEIGHT, SLIDE_WIDTH, Design, add_bar, add_text as _add_text,
    blank_slide as _blank_slide, set_notes, solid,
)

# --- Design constants -------------------------------------------------------

HEADING_FONT = "Segoe UI Semibold"
BODY_FONT = "Segoe UI"

PRIMARY = RGBColor(0x1F, 0x38, 0x64)  # brand navy — theme color, legacy fill key
CANVAS = RGBColor(0x14, 0x24, 0x44)  # near-black navy — every slide's ground
PANEL = RGBColor(0x20, 0x36, 0x5E)  # elevated card surface on the canvas
PANEL_EDGE = RGBColor(0x31, 0x4E, 0x80)  # card borders, hairlines
NAVY_SOFT = RGBColor(0x35, 0x52, 0x82)  # inactive markers, ghosted fills
ACCENT = RGBColor(0x37, 0x8C, 0xE6)  # vivid azure — chips, bars, label blocks
HIGHLIGHT = RGBColor(0xE8, 0x59, 0x0C)  # vivid orange — demo cues, emphasis
AMBER = RGBColor(0xFF, 0xB1, 0x3D)  # warm amber — numbers, accents, headings
TEXT = RGBColor(0xE9, 0xF0, 0xFA)  # near-white body text
MUTED = RGBColor(0x9D, 0xB2, 0xD3)  # light slate — captions, credits
SLATE = RGBColor(0x8E, 0xA3, 0xBF)  # de-emphasized fills that stay legible
SKY = RGBColor(0xBD, 0xD7, 0xF5)  # light azure text
LINE = PANEL_EDGE  # legacy alias — borders on the dark canvas
PALE = RGBColor(0xE8, 0xF0, 0xFB)  # legacy fill key → neutral panel card
PALE_WARM = RGBColor(0xFD, 0xE9, 0xD8)  # legacy fill key → emphasis panel card
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

FOOTER_TEXT = "AI for Research Management · Lesson 1 · Week 1"

# Content area shared by all content slides (below the title block).
CONTENT_LEFT = Inches(0.6)
CONTENT_TOP = Inches(1.5)
CONTENT_WIDTH = Inches(12.13)
CONTENT_HEIGHT = Inches(5.3)

# --- Presentation setup -----------------------------------------------------

DESIGN = Design(
    heading_font=HEADING_FONT,
    body_font=BODY_FONT,
    scheme={"dk2": PRIMARY, "accent1": ACCENT, "accent2": HIGHLIGHT},
    canvas=CANVAS,
    footer_text=FOOTER_TEXT,
    footer_rule=PANEL_EDGE,
    footer_color=MUTED,
)


def new_presentation():
    """A 16:9 presentation with theme fonts/colors, dark canvas, and footer."""
    return core.new_presentation(DESIGN)


def outlined(shape, fill=PANEL, border=PANEL_EDGE, weight=1.0):
    """Elevated panel style: solid fill with a thin border, no shadow."""
    core.outlined(shape, fill, border, weight)


def panel(shape):
    """The standard elevated card surface."""
    outlined(shape, PANEL, PANEL_EDGE)


def style(run, size, color, font=BODY_FONT, bold=False, italic=False):
    core.style(run, size, color, font, bold=bold, italic=italic)


# --- Slide builders ---------------------------------------------------------


def add_title_slide(prs, title, course_name, session_line):
    """Full-bleed deep-navy cover: amber kicker, white title, tri-color base."""
    slide = _blank_slide(prs)

    bg = add_bar(slide, 0, 0, SLIDE_WIDTH, SLIDE_HEIGHT, CANVAS)
    bg.shadow.inherit = False

    add_bar(slide, Inches(0.9), Inches(1.55), Inches(1.0), Inches(0.06), AMBER)

    tf = _add_text(slide, Inches(0.9), Inches(1.75), Inches(11.5), Inches(0.5))
    run = tf.paragraphs[0].add_run()
    run.text = course_name.upper()
    style(run, size=16, color=AMBER, font=HEADING_FONT, bold=True)

    tf = _add_text(slide, Inches(0.9), Inches(2.5), Inches(11.5), Inches(2.0))
    run = tf.paragraphs[0].add_run()
    run.text = title
    style(run, size=40, color=WHITE, font=HEADING_FONT, bold=True)

    tf = _add_text(slide, Inches(0.9), Inches(4.75), Inches(11.5), Inches(0.5))
    run = tf.paragraphs[0].add_run()
    run.text = session_line
    style(run, size=16, color=SKY)

    # Tri-color baseline: azure, amber, orange — the deck's palette in one line.
    add_bar(slide, 0, Inches(7.26), Inches(8.0), Inches(0.24), ACCENT)
    add_bar(slide, Inches(8.0), Inches(7.26), Inches(3.2), Inches(0.24), AMBER)
    add_bar(slide, Inches(11.2), Inches(7.26), Inches(2.14), Inches(0.24), HIGHLIGHT)
    return slide


def add_section_divider(prs, number, total, title, duration, mode):
    """Full-bleed navy divider: giant number, amber rule, progress pills."""
    slide = _blank_slide(prs)

    bg = add_bar(slide, 0, 0, SLIDE_WIDTH, SLIDE_HEIGHT, CANVAS)
    bg.shadow.inherit = False

    tf = _add_text(
        slide, Inches(0.7), Inches(2.0), Inches(2.5), Inches(2.4),
        align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE,
    )
    run = tf.paragraphs[0].add_run()
    run.text = str(number)
    style(run, size=120, color=WHITE, font=HEADING_FONT, bold=True)

    # Vertical amber rule between the number and the text block.
    add_bar(slide, Inches(3.45), Inches(2.35), Inches(0.05), Inches(1.75), AMBER)

    tf = _add_text(
        slide, Inches(3.9), Inches(2.25), Inches(8.6), Inches(1.35),
        anchor=MSO_ANCHOR.BOTTOM,
    )
    run = tf.paragraphs[0].add_run()
    run.text = title
    style(run, size=32, color=WHITE, font=HEADING_FONT, bold=True)

    tf = _add_text(slide, Inches(3.9), Inches(3.75), Inches(8.6), Inches(0.5))
    run = tf.paragraphs[0].add_run()
    run.text = f"{duration} · {mode}"
    style(run, size=18, color=AMBER, font=HEADING_FONT, bold=True)

    tf = _add_text(slide, Inches(3.9), Inches(4.3), Inches(8.6), Inches(0.4))
    run = tf.paragraphs[0].add_run()
    run.text = f"Segment {number} of {total}"
    style(run, size=12, color=SKY)

    # Progress pills: one per segment, the current one amber.
    pill_w, pill_h, gap = 0.55, 0.1, 0.2
    row_w = total * pill_w + (total - 1) * gap
    x = (SLIDE_WIDTH.inches - row_w) / 2
    for i in range(1, total + 1):
        add_bar(slide, Inches(x), Inches(6.55), Inches(pill_w), Inches(pill_h),
                AMBER if i == number else NAVY_SOFT)
        x += pill_w + gap
    return slide


# How far the header block rides up on a compact slide, in inches. Chart
# slides carry a full-width image that needs the vertical room; everything
# else keeps the standard header.
COMPACT_HEADER_RISE = 0.44
COMPACT_CONTENT_TOP = Inches(0.9)  # first free inch below a compact header


def add_content_slide(prs, title, body_lines=None, notes=None, compact=False):
    """Standard content slide on the dark canvas: orange kicker bar, white
    title, layered rule with a vivid azure lead segment.

    body_lines is a list of strings, or (indent_level, string) tuples.
    Returns the slide so callers can add custom shapes in the content area
    (CONTENT_LEFT/TOP/WIDTH/HEIGHT).

    compact=True raises the whole header block by COMPACT_HEADER_RISE so an
    image-dominant slide can start its content at COMPACT_CONTENT_TOP.
    """
    slide = _blank_slide(prs)
    rise = COMPACT_HEADER_RISE if compact else 0.0

    add_bar(slide, CONTENT_LEFT, Inches(0.55 - rise), Inches(0.13), Inches(0.5),
            HIGHLIGHT)

    tf = _add_text(slide, Inches(0.95), Inches(0.45 - rise), Inches(11.78),
                   Inches(0.7))
    run = tf.paragraphs[0].add_run()
    run.text = title
    style(run, size=28, color=WHITE, font=HEADING_FONT, bold=True)

    # Layered rule: full-width hairline with a vivid azure lead segment.
    add_bar(slide, CONTENT_LEFT, Inches(1.21 - rise), CONTENT_WIDTH, Pt(1.2),
            PANEL_EDGE)
    add_bar(slide, CONTENT_LEFT, Inches(1.19 - rise), Inches(1.6), Pt(3), ACCENT)

    if body_lines:
        tf = _add_text(slide, CONTENT_LEFT, CONTENT_TOP, CONTENT_WIDTH, CONTENT_HEIGHT)
        first = True
        for entry in body_lines:
            level, text = entry if isinstance(entry, tuple) else (0, entry)
            para = tf.paragraphs[0] if first else tf.add_paragraph()
            first = False
            para.level = level
            run = para.add_run()
            run.text = text
            style(run, size=18 if level == 0 else 15, color=TEXT if level == 0 else MUTED)
            para.space_after = Pt(8)

    if notes:
        set_notes(slide, notes)
    return slide


def add_demo_cue(slide, text, top, left=CONTENT_LEFT, width=CONTENT_WIDTH):
    """Orange live-demo marker: the facilitator switches to a live tool here.

    Renders as a filled band reading "▶ LIVE DEMO — <text>". Callers must
    also mention the demo in the slide's speaker notes.
    """
    band = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, left, top, width, Inches(0.55)
    )
    solid(band, HIGHLIGHT)
    add_bar(slide, left, top, Inches(0.1), Inches(0.55), AMBER)
    tf = band.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.3)
    tf.margin_right = Inches(0.25)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    run = tf.paragraphs[0].add_run()
    run.text = f"▶ LIVE DEMO — {text}"
    style(run, size=16, color=WHITE, font=HEADING_FONT, bold=True)
    return band


def add_band(slide, top, height, text, size=20, fill=PRIMARY, sub=None):
    """Full-width emphasis band: centered bold white text on a vivid solid
    rectangle with symmetric amber end caps; optional sub-line. The legacy
    navy fill maps to vivid azure so bands pop on the dark canvas."""
    if fill == PRIMARY:
        fill = ACCENT
    band = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, CONTENT_LEFT, top, CONTENT_WIDTH, height
    )
    solid(band, fill)
    add_bar(slide, CONTENT_LEFT, top, Inches(0.1), height, AMBER)
    add_bar(slide, Inches(CONTENT_LEFT.inches + CONTENT_WIDTH.inches - 0.1),
            top, Inches(0.1), height, AMBER)
    tf = band.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.4)
    tf.margin_right = Inches(0.4)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.paragraphs[0].alignment = PP_ALIGN.CENTER
    run = tf.paragraphs[0].add_run()
    run.text = text
    style(run, size=size, color=WHITE, font=HEADING_FONT, bold=True)
    if sub:
        para = tf.add_paragraph()
        para.alignment = PP_ALIGN.CENTER
        para.space_before = Pt(6)
        run = para.add_run()
        run.text = sub
        style(run, size=13, color=PALE if fill == ACCENT else PALE_WARM)
    return band


def add_card(slide, left, top, width, height, fill, heading, heading_color,
             body=None):
    """Deck card on the dark canvas. The legacy navy fill renders as a
    solid vivid-azure label block (centered white heading). Light legacy
    fills render as elevated navy panels with a colored edge bar — azure
    for neutral cards, orange for emphasis — and headings remap to white
    (neutral) or amber (emphasis) so they read on the dark surface."""
    card = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    dark = fill == PRIMARY
    if dark:
        solid(card, ACCENT)
    else:
        panel(card)
        warm = fill == PALE_WARM or heading_color == HIGHLIGHT
        add_bar(slide, left, top, Inches(0.09), height,
                HIGHLIGHT if warm else ACCENT)
    if heading_color == PRIMARY:
        heading_color = WHITE
    elif heading_color == HIGHLIGHT:
        heading_color = AMBER
    tf = card.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.3)
    tf.margin_right = Inches(0.25)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    para = tf.paragraphs[0]
    if not dark:
        para.alignment = PP_ALIGN.LEFT
    run = para.add_run()
    run.text = heading
    style(run, size=18, color=heading_color, font=HEADING_FONT, bold=True)
    if body:
        para = tf.add_paragraph()
        para.alignment = PP_ALIGN.LEFT
        para.space_before = Pt(4)
        run = para.add_run()
        run.text = body
        style(run, size=14, color=TEXT)
    return card


def add_label(slide, left, top, width, text, size=14, color=MUTED, italic=False,
              align=PP_ALIGN.LEFT):
    """Small free-standing caption/label text."""
    box = slide.shapes.add_textbox(left, top, width, Inches(0.4))
    tf = box.text_frame
    tf.word_wrap = True
    tf.paragraphs[0].alignment = align
    run = tf.paragraphs[0].add_run()
    style(run, size=size, color=color, italic=italic)
    run.text = text
    return box


def add_placeholder_slide(prs, segment_label, planned_title, gist, ticket):
    """A visibly-unfinished slide standing in for one planned content slide."""
    slide = add_content_slide(prs, planned_title)

    card = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(1.6), Inches(2.3), Inches(10.1), Inches(2.9)
    )
    outlined(card, PANEL, ACCENT, weight=1.5)
    tf = card.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.4)
    tf.margin_right = Inches(0.4)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    para = tf.paragraphs[0]
    para.alignment = PP_ALIGN.CENTER
    run = para.add_run()
    run.text = f"[PLACEHOLDER] {segment_label}"
    style(run, size=20, color=WHITE, font=HEADING_FONT, bold=True)
    para = tf.add_paragraph()
    para.alignment = PP_ALIGN.CENTER
    run = para.add_run()
    run.text = gist
    style(run, size=14, color=MUTED)

    set_notes(
        slide,
        f"Placeholder — replaced by ticket {ticket} ({segment_label}). "
        f"Planned content: {gist}",
    )
    return slide
