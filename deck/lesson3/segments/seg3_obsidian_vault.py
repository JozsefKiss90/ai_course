"""Segment 3 — Obsidian: the Personal Vault and its two plugins (20 min,
Live facilitator demo). Local segment.

Eight slides, per lesson_3.md Segment 3: what Obsidian is (with the
installation flag); the vault's four folders and the note template; links,
backlinks, tags and the linking discipline; Graph View and Canvas, with the
Gemini Canvas name clash; the two community plugins side by side; why two
plugins; the setup order for the hands-on block; and the six-step demo
script.

Obsidian and Smart Connections never call a cloud model, so the segment
carries the local colour; the one cloud call from the vault — Copilot for
Obsidian, on the participant's own free Gemini API key — is amber wherever
it appears, and it is never called bare "Copilot".
"""

from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

from deck.lesson3 import template
from deck.lesson3.template import (
    DAILY_DRIVER, DAILY_TINT, HEADING_FONT, INK, LOCAL, LOCAL_TINT, MUTED,
    SLATE, TEXT, WHITE,
)

LANE = LOCAL
_LEFT = template.CONTENT_LEFT
_WIDTH = template.CONTENT_WIDTH


def _mono_card(slide, left, top, width, height, edge, heading, lines):
    """A card whose body is verbatim text (a folder tree, a template),
    line by line, in the body face at a fixed size."""
    card = template.add_open_card(slide, left, top, width, height, edge)
    tf = card.text_frame
    tf.vertical_anchor = MSO_ANCHOR.TOP
    tf.margin_top = Inches(0.14)
    run = tf.paragraphs[0].add_run()
    run.text = heading
    template.style(run, size=13, color=INK, font=HEADING_FONT, bold=True)
    for line in lines:
        para = tf.add_paragraph()
        para.space_before = Pt(2)
        run = para.add_run()
        run.text = line
        template.style(run, size=12.5, color=TEXT)
    return card


# --- 1. What Obsidian is -----------------------------------------------------


def _what_obsidian_is_slide(prs):
    slide = template.add_content_slide(prs, "What Obsidian is", lane=LANE)
    template.add_detail_card(
        slide, _LEFT, Inches(1.42), Inches(5.95), Inches(1.45), LANE.color,
        "Notes as plain Markdown files in a folder on your disk",
        ["No account, no cloud, no lock-in: the vault is a folder, and any "
         "text editor can open it. Nothing in Obsidian requires an account."],
        heading_size=14, line_size=12,
    )
    template.add_detail_card(
        slide, _LEFT + Inches(6.18), Inches(1.42), Inches(5.95), Inches(1.45), LANE.color,
        "Free for personal use",
        ["The paid add-ons (Sync, Publish) are not used."],
        heading_size=14, line_size=12,
    )
    template.add_detail_card(
        slide, _LEFT, Inches(3.02), Inches(5.95), Inches(1.45), LANE.color,
        "The fourth floor of the four-floor map",
        ["Where the notebook's output becomes durable, linked, navigable work."],
        heading_size=14, line_size=12,
    )
    template.add_detail_card(
        slide, _LEFT + Inches(6.18), Inches(3.02), Inches(5.95), Inches(1.45), LANE.color,
        "One connecting thread",
        ["Work scattered across four places has no connecting thread; the "
         "vault is the thread."],
        heading_size=14, line_size=12,
    )
    template.add_band(
        slide, Inches(4.72), Inches(1.15),
        "Installation flag",
        size=15, fill=SLATE,
        sub="Desktop application (Windows, macOS, Linux; mobile apps exist). "
            "On institutionally managed machines the installer may need "
            "administrator rights: confirmed the week before; if not, the "
            "portable build or a personal device for the hands-on block.",
    )
    template.add_label(
        slide, _LEFT, Inches(6.05), _WIDTH,
        "This segment stays local for everything that is Obsidian itself; "
        "the one cloud call — Copilot for Obsidian — is marked amber when it "
        "comes up.",
        size=11, color=MUTED, italic=True,
    )
    template.set_notes(slide, (
        "~2 min. A note-taking application whose notes are plain Markdown "
        "files in a folder on your disk: no account, no cloud, no lock-in. "
        "Free for personal use; Sync and Publish are paid and not used. The "
        "vault is the fourth floor of the four-floor map. Talk row E7, one "
        "line: work scattered across four places has no connecting thread; "
        "the vault is the thread.\n\n"
        "Installation flag: Obsidian is a desktop application. On "
        "institutionally managed machines the installer may need "
        "administrator rights — you confirmed the week before whether "
        "participants can install it; if not, the portable build or a "
        "personal device for the hands-on block. Nothing in Obsidian "
        "requires an account.\n\n"
        "Live demo in Obsidian on your own machine runs through this segment; "
        "the demo-script slide at the end carries the six steps and timings."
    ))
    return slide


# --- 2. The vault ------------------------------------------------------------

_TREE = [
    "Anchor Project vault/",
    "├── concepts/",
    "├── processes/",
    "├── guides/",
    "├── projects/",
    "└── _templates/      (one note template)",
]

_TEMPLATE = [
    "---",
    "type: concept | process | guide | project",
    "source: <notebook name> · <source document> · <section, if any>",
    "captured: YYYY-MM-DD",
    "status: draft | checked",
    "---",
    "# <Title>",
    "<The checked text from the notebook, edited by you.>",
    "## Related",
    "- [[ ]]",
]


def _vault_slide(prs):
    slide = template.add_content_slide(
        prs, "The vault: four folders and one note template", lane=LANE)
    _mono_card(slide, _LEFT, Inches(1.42), Inches(4.6), Inches(2.6), LANE.color,
               "One vault per participant, named after the Anchor Project",
               _TREE)
    template.add_detail_card(
        slide, _LEFT, Inches(4.17), Inches(4.6), Inches(2.1), SLATE,
        "The human gate, made visible",
        ["A note is draft until its text has been checked against the "
         "citation, then checked.",
         "Graph View and the plugins do not care; the participant does.",
         "The four folders are the four entity types — the wiki's content "
         "model, unchanged."],
        line_size=11,
    )
    _mono_card(slide, _LEFT + Inches(4.83), Inches(1.42), Inches(7.3), Inches(4.3),
               LANE.color,
               "The note template — one file in _templates/, inserted into every "
               "new note (core Templates plugin)",
               _TEMPLATE)
    template.add_label(
        slide, _LEFT + Inches(4.83), Inches(5.85), Inches(7.3),
        "The source and status fields are Segment 2's rule as frontmatter: "
        "the notebook's name, the source document and the section; draft "
        "until checked.",
        size=11, color=MUTED, italic=True,
    )
    template.add_demo_cue(
        slide, "steps 1–2: create the vault, paste the Segment 2 briefing, "
        "set draft, check one sentence, set checked", Inches(6.4), lane=LANE)
    template.set_notes(slide, (
        "~2 min of the five on the vault. Structure: one vault per "
        "participant, named after the Anchor Project, with four folders that "
        "are the four entity types, plus _templates/ holding one note "
        "template that Obsidian's core Templates plugin inserts into every "
        "new note.\n\n"
        "The source and status fields are the human gate from Segment 2 made "
        "visible: a note is draft until its text has been checked against "
        "the citation, then checked. Graph View and the plugins do not care; "
        "the participant does.\n\n"
        "Live: demo steps 1–2 run here — create the vault “Demo Anchor "
        "Project”, the four folders, the template; paste the Segment 2 "
        "briefing into processes/, fill the frontmatter, set status: draft; "
        "check one sentence against the notebook's citation, set checked."
    ))
    return slide


# --- 3. Links, backlinks, tags -----------------------------------------------


def _links_slide(prs):
    slide = template.add_content_slide(
        prs, "Links, backlinks, tags — and the linking discipline", lane=LANE)
    template.add_detail_card(
        slide, _LEFT, Inches(1.42), Inches(3.9), Inches(2.0), LANE.color,
        "Links",
        ["[[Note title]] creates a link; typing [[ offers every note in the "
         "vault; a link to a note that does not exist yet creates it when "
         "clicked."],
        heading_size=14, line_size=12,
    )
    template.add_detail_card(
        slide, _LEFT + Inches(4.12), Inches(1.42), Inches(3.9), Inches(2.0), LANE.color,
        "Backlinks",
        ["The right-hand pane shows every note that points at the current "
         "one."],
        heading_size=14, line_size=12,
    )
    template.add_detail_card(
        slide, _LEFT + Inches(8.23), Inches(1.42), Inches(3.9), Inches(2.0), LANE.color,
        "Tags",
        ["#eligible-cost, #deadline — tags cut across the four folders."],
        heading_size=14, line_size=12,
    )
    template.add_band(
        slide, Inches(3.65), Inches(0.95),
        "a note with no links is a file, not knowledge",
        size=20, fill=LANE.color, sub="the slide's one rule",
    )
    template.add_list_card(
        slide, _LEFT, Inches(4.8), _WIDTH, Inches(1.4), LOCAL_TINT,
        "Linking discipline",
        ["Every concept links to the process it appears in",
         "Every process links to its concepts and to the guide that explains it",
         "The project note links to everything — Graph View only pays off if "
         "entities reference each other"],
        size=12.5,
    )
    template.add_demo_cue(
        slide, "step 3: two concept notes, [[ ]] links from the process note, "
        "backlinks, one tag", Inches(6.35), lane=LANE)
    template.set_notes(slide, (
        "~2 min. Links: [[Note title]] creates one; typing [[ offers every "
        "note in the vault; a link to a note that does not exist yet creates "
        "it when clicked. Backlinks in the right-hand pane show every note "
        "pointing at the current one. Tags cut across the four folders.\n\n"
        "The one rule: a note with no links is a file, not knowledge. Every "
        "concept links to the process it appears in; every process to its "
        "concepts and to the guide that explains it; the project note to "
        "everything. Graph View only pays off if entities reference each "
        "other.\n\n"
        "Live: demo step 3 — create two concept notes from the same "
        "briefing, link them from the process note with [[ ]], show "
        "backlinks, add one tag."
    ))
    return slide


# --- 4. Graph View and Canvas ------------------------------------------------


def _graph_slide(prs):
    slide = template.add_content_slide(prs, "Graph View and Canvas", lane=LANE)
    template.add_detail_card(
        slide, _LEFT, Inches(1.42), Inches(5.95), Inches(2.3), LANE.color,
        "Graph View",
        ["The whole vault as nodes and edges",
         "Filters by folder and tag",
         "A local graph for one note"],
        heading_size=15, line_size=12.5,
    )
    template.add_detail_card(
        slide, _LEFT + Inches(6.18), Inches(1.42), Inches(5.95), Inches(2.3), LANE.color,
        "Canvas",
        ["A freeform board where notes, cards and images are arranged spatially",
         "For planning a proposal timeline or mapping stakeholders"],
        heading_size=15, line_size=12.5,
    )
    template.add_band(
        slide, Inches(3.95), Inches(1.15),
        "Name clash with Lesson 2",
        size=16, fill=SLATE,
        sub="Gemini Canvas is a drafting surface inside a chat; Obsidian "
            "Canvas is a whiteboard of your own notes. Different things.",
    )
    template.add_demo_cue(
        slide, "step 4: Graph View with the three nodes just made, then the "
        "prepared vault of ~25 linked notes with the folder filter",
        Inches(5.35), lane=LANE)
    template.set_notes(slide, (
        "~1 min. Graph View: the whole vault as nodes and edges; filters by "
        "folder and tag; local graph for one note. Canvas: a freeform board "
        "where notes, cards and images are arranged spatially, for planning "
        "a proposal timeline or mapping stakeholders.\n\n"
        "Name clash with Gemini's Canvas from Lesson 2: Gemini Canvas is a "
        "drafting surface inside a chat, Obsidian Canvas is a whiteboard of "
        "your own notes.\n\n"
        "Live: demo step 4 — Graph View with the three nodes just made, then "
        "open the prepared vault of ~25 linked notes and show the same view "
        "with the folder filter."
    ))
    return slide


# --- 5. Two community plugins ------------------------------------------------

# The plan's comparison table: (row label, Smart Connections, Copilot for Obsidian).
_PLUGIN_ROWS = [
    ("Author", "Brian Petro", "logancyang"),
    ("What it does",
     "Shows notes semantically related to the one you are reading "
     "(Connections view); semantic Lookup across the vault",
     "Chat with a model inside Obsidian; Vault QA: ask a question, get an "
     "answer assembled from your notes, with the notes cited; Relevant "
     "Notes surfaced in the chat"),
    ("How it works",
     "Builds embeddings of your notes with a local model that ships with "
     "the plugin: no API key, no account, nothing leaves the machine",
     "Indexes the vault locally, then sends the question and the relevant "
     "note chunks to the model you configure; you bring your own API key"),
    ("Cost",
     "Free core (v4). Paid “Pro” add-ons and the separate Smart Chat plugin "
     "exist and are not used",
     "Free plugin with your own key. “Copilot Plus” is a paid tier and is "
     "not used"),
    ("Which model",
     "Local embedding model, no choice needed",
     "A free Gemini API key from Google AI Studio; the plugin lists Gemini "
     "models to choose from"),
    ("Lane",
     "No cloud call: the only AI in the course that runs entirely on your "
     "device",
     "Daily Driver. The free API tier may use submitted content to improve "
     "Google's models; the vault holds anonymised or fictional Anchor "
     "Project content only, so this is acceptable; nothing from the "
     "Lesson 2 “never” list is ever in the vault"),
    ("Job in the KM System",
     "Finding: “what in my vault relates to this?” The link you forgot to make",
     "Asking: “what does my vault say about X?” The vault's own grounding "
     "floor, with citations to notes instead of to sources"),
    ("Name clash",
     "—",
     "Not Microsoft Copilot. Unrelated project, unrelated data path, never "
     "the Sensitive Lane"),
]

_ROW_HEIGHTS = [0.32, 0.62, 0.62, 0.5, 0.5, 0.76, 0.5, 0.42]
_LABEL_W = 1.75
_COL_W = 5.12
_COL_GAP = 0.1


def _plugins_slide(prs):
    slide = template.add_content_slide(
        prs, "Two community plugins: the vault's own AI", lane=LANE)
    template.add_label(
        slide, _LEFT, Inches(1.27), _WIDTH,
        "Settings → Community plugins → turn off Restricted mode. Only these "
        "two, from the official community list; read each plugin's data note "
        "before enabling.",
        size=11, color=MUTED,
    )
    sc_left = _LEFT + Inches(_LABEL_W + _COL_GAP)
    cfo_left = sc_left + Inches(_COL_W + _COL_GAP)
    template.add_card(slide, sc_left, Inches(1.66), Inches(_COL_W), Inches(0.4),
                      LOCAL.color, "Smart Connections", WHITE)
    template.add_card(slide, cfo_left, Inches(1.66), Inches(_COL_W), Inches(0.4),
                      DAILY_DRIVER.color, "Copilot for Obsidian", WHITE)
    top = 2.12
    for (label, sc, cfo), height in zip(_PLUGIN_ROWS, _ROW_HEIGHTS):
        box = template.add_label(slide, _LEFT, Inches(top), Inches(_LABEL_W),
                                 label, size=10.5, color=INK, bold=True)
        box.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
        box.height = Inches(height)
        for left, text, tint in ((sc_left, sc, LOCAL_TINT), (cfo_left, cfo, DAILY_TINT)):
            cell = template.add_open_card(slide, left, Inches(top), Inches(_COL_W),
                                          Inches(height),
                                          LOCAL.color if tint == LOCAL_TINT
                                          else DAILY_DRIVER.color)
            cell.text_frame.margin_top = 0
            cell.text_frame.margin_bottom = 0
            run = cell.text_frame.paragraphs[0].add_run()
            run.text = text
            template.style(run, size=9.5, color=TEXT)
        top += height + 0.06
    template.set_notes(slide, (
        "~8 min for the two plugins, this slide and the next two. Enabling "
        "community plugins means switching off Obsidian's Restricted mode "
        "(Settings → Community plugins). Show the switch and the warning, "
        "and state the rule: only these two plugins, from the official "
        "community list, and read the plugin's data note before enabling.\n\n"
        "Read the table by column, not by row: Smart Connections builds "
        "embeddings with a local model that ships with the plugin — no API "
        "key, no account, nothing leaves the machine — and surfaces related "
        "notes while you write. Copilot for Obsidian indexes the vault "
        "locally, then sends the question and the relevant note chunks to "
        "the model you configure on your own free Gemini API key from Google "
        "AI Studio; Vault QA answers with the notes cited. Free core in both "
        "cases; the paid tiers (Smart Connections Pro and the separate Smart "
        "Chat plugin, Copilot Plus) are not used.\n\n"
        "Lane: Smart Connections is the only AI in the course that runs "
        "entirely on the device. Copilot for Obsidian is Daily Driver: the "
        "free API tier may use submitted content to improve Google's models; "
        "the vault holds anonymised or fictional Anchor Project content "
        "only, so this is acceptable, and nothing from the Lesson 2 “never” "
        "list is ever in the vault.\n\n"
        "Say the name clash once: not Microsoft Copilot. Unrelated project, "
        "unrelated data path, never the Sensitive Lane."
    ))
    return slide


# --- 6. Why two plugins ------------------------------------------------------


def _why_two_slide(prs):
    slide = template.add_content_slide(prs, "Why two plugins", lane=LANE)
    template.add_label(
        slide, _LEFT, Inches(1.3), _WIDTH,
        "They answer different questions",
        size=14, color=INK, bold=True,
    )
    template.add_card(
        slide, _LEFT, Inches(1.85), Inches(5.95), Inches(2.5), LOCAL_TINT,
        "Smart Connections — finding", INK,
        "Surfaces relationships while you write, silently, offline — which "
        "is what makes the linking discipline sustainable past thirty notes.",
    )
    template.add_card(
        slide, _LEFT + Inches(6.18), Inches(1.85), Inches(5.95), Inches(2.5), DAILY_TINT,
        "Copilot for Obsidian — asking", INK,
        "Answers questions over the whole vault the way NotebookLM answers "
        "over sources — so once the vault is the source of truth, it can be "
        "queried like one.",
    )
    template.add_band(
        slide, Inches(4.6), Inches(1.2),
        "capture in the notebook, keep in the vault, query the vault",
        size=19,
        sub="together they close the loop the wiki was supposed to close",
    )
    template.set_notes(slide, (
        "~1 min. Smart Connections surfaces relationships while you write, "
        "silently, offline, which is what makes the linking discipline "
        "sustainable past thirty notes. Copilot for Obsidian answers "
        "questions over the whole vault the way NotebookLM answers over "
        "sources, so once the vault is the source of truth it can be queried "
        "like one. Together they close the loop the wiki was supposed to "
        "close: capture in the notebook, keep in the vault, query the vault."
    ))
    return slide


# --- 7. Setup order ----------------------------------------------------------

_SETUP_STEPS = [
    ("1", "Settings → Community plugins → turn off Restricted mode → Browse"),
    ("2", "Install and enable Smart Connections; wait for the first embedding "
          "run (a progress notice); open the Connections view from the right "
          "sidebar"),
    ("3", "Get a Gemini API key at Google AI Studio (personal Google account, "
          "free tier)"),
    ("4", "Install and enable Copilot for Obsidian; Settings → Copilot for "
          "Obsidian → add the Gemini key, pick a Gemini chat model and, if "
          "offered, a Gemini embedding model; run “Refresh vault index”"),
    ("5", "Open the Copilot for Obsidian pane, switch to Vault QA, ask one "
          "question"),
]


def _setup_slide(prs):
    slide = template.add_content_slide(
        prs, "Setup order for the hands-on block", lane=LANE)
    edges = [LOCAL.color, LOCAL.color, DAILY_DRIVER.color, DAILY_DRIVER.color,
             DAILY_DRIVER.color]
    top = 1.5
    for (number, text), edge in zip(_SETUP_STEPS, edges):
        template.add_badge(slide, _LEFT, Inches(top + 0.12), Inches(0.5), number)
        card = template.add_open_card(slide, _LEFT + Inches(0.65), Inches(top),
                                      _WIDTH - Inches(0.65), Inches(0.74), edge)
        run = card.text_frame.paragraphs[0].add_run()
        run.text = text
        template.style(run, size=12.5, color=TEXT)
        top += 0.84
    template.add_label(
        slide, _LEFT, Inches(5.75), _WIDTH,
        "Green: stays on the machine. Amber: the key and the Copilot for "
        "Obsidian calls go to Google's free API tier — Daily Driver.",
        size=11, color=MUTED, italic=True,
    )
    template.set_notes(slide, (
        "~1 min, read once; this slide comes back in the hands-on block as "
        "steps 6–8 of the checklist. Typical snags, said now so they are "
        "familiar later: the API key pasted with a trailing space; the "
        "refresh of the vault index not run, so Vault QA answers “no "
        "relevant notes”; Smart Connections still embedding on a slow "
        "machine (wait, it is a one-off).\n\n"
        "Preparation: both plugins installed fresh on a clean vault the day "
        "before, to catch settings changes, and the Google AI Studio key "
        "flow walked through once."
    ))
    return slide


# --- 8. Demo script ----------------------------------------------------------

_DEMO_STEPS = [
    ("1", "Create a vault “Demo Anchor Project”, the four folders, the "
          "template", "1 min"),
    ("2", "Paste the Segment 2 briefing into processes/, fill the "
          "frontmatter, set status: draft; check one sentence against the "
          "notebook's citation, set checked", "1.5 min"),
    ("3", "Create two concept notes from the same briefing, link them from "
          "the process note with [[ ]], show backlinks; add one tag", "1.5 min"),
    ("4", "Graph View: three nodes; then open a prepared vault of ~25 linked "
          "notes and show the same view, with the folder filter", "1 min"),
    ("5", "Smart Connections on the prepared vault: open a note, read the "
          "Connections view, click one related note you “forgot” to link, "
          "add the link", "1.5 min"),
    ("6", "Copilot for Obsidian on the prepared vault, Vault QA: “Which "
          "deadlines apply to the periodic report, and which notes say so?” "
          "Read the cited notes; ask a question the vault cannot answer and "
          "show the answer say so", "1.5 min"),
]


def _demo_script_slide(prs):
    slide = template.add_content_slide(
        prs, "Demo script — six steps, one prepared vault", lane=LANE)
    template.add_demo_cue(
        slide, "Obsidian on the facilitator's machine, ≈8 of the 20 minutes, "
        "interleaved with the slides", Inches(1.4), lane=LANE)
    edges = [LOCAL.color] * 5 + [DAILY_DRIVER.color]
    top = 2.1
    for (number, text, minutes), edge in zip(_DEMO_STEPS, edges):
        template.add_badge(slide, _LEFT, Inches(top + 0.04), Inches(0.5), number)
        card = template.add_open_card(slide, _LEFT + Inches(0.65), Inches(top),
                                      _WIDTH - Inches(1.75), Inches(0.58), edge)
        run = card.text_frame.paragraphs[0].add_run()
        run.text = text
        template.style(run, size=10.5, color=TEXT)
        tail = template.add_open_card(slide, _LEFT + _WIDTH - Inches(1.0),
                                      Inches(top), Inches(1.0), Inches(0.58), edge)
        tail.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
        run = tail.text_frame.paragraphs[0].add_run()
        run.text = minutes
        template.style(run, size=11, color=INK, bold=True)
        top += 0.66
    template.add_band(
        slide, Inches(6.12), Inches(0.62),
        "The vault contains only anonymised or fictional content, so the "
        "Copilot for Obsidian call is a Daily Driver call  ·  Smart "
        "Connections never calls anything",
        size=12, fill=SLATE,
    )
    template.set_notes(slide, (
        "The six live demo steps in one place, ≈8 minutes of the 20, "
        "interleaved with the slides before this one (steps 1–2 on the vault "
        "slide, 3 on the links slide, 4 on the Graph View slide, 5–6 here). "
        "Two vaults: the “Demo Anchor Project” built live from the Segment 2 "
        "briefing, and the prepared vault of ~25 linked notes with Smart "
        "Connections embedded and Copilot for Obsidian configured.\n\n"
        "Step 6 is the one cloud call: Vault QA on the prepared vault — "
        "“Which deadlines apply to the periodic report, and which notes say "
        "so?” — read the cited notes, then ask a question the vault cannot "
        "answer and show the answer say so.\n\n"
        "Rule stated while demoing: the vault contains only anonymised or "
        "fictional content, so the Copilot for Obsidian call is a Daily "
        "Driver call; Smart Connections never calls anything."
    ))
    return slide


def add_slides(prs):
    _what_obsidian_is_slide(prs)
    _vault_slide(prs)
    _links_slide(prs)
    _graph_slide(prs)
    _plugins_slide(prs)
    _why_two_slide(prs)
    _setup_slide(prs)
    _demo_script_slide(prs)
