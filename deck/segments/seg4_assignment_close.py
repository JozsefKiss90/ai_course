"""Segment 4 — Assignment briefing + close (10 min, Presentation + Q&A).

Four slides per lesson_1.md Segment 4: the Anchor Project (the
anonymize-or-fictional rule and its role as the course-long thread), the
five-part weekly assignment on one screenshot-friendly slide, the effort
expectation with the Lesson 2 debrief format, and the close — a teaser of
next week's two-lane rule using the CONTEXT.md lane names without
teaching the rule itself.
"""

from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

from deck import template
from deck.template import (
    CONTENT_WIDTH, HIGHLIGHT, LINE, PALE, PALE_WARM, PRIMARY, TEXT, WHITE,
)


# --- Slide 1: your Anchor Project --------------------------------------------


def _anchor_project_slide(prs):
    slide = template.add_content_slide(prs, "Your Anchor Project")

    template.add_card(
        slide, template.CONTENT_LEFT, Inches(1.6), Inches(5.96), Inches(1.7),
        PALE, "A real project — anonymized", PRIMARY,
        "One research project from your environment. Strip names, partners, "
        "amounts — anonymize anything you couldn't publish.",
    )
    template.add_card(
        slide, Inches(template.CONTENT_LEFT.inches + 6.17), Inches(1.6),
        Inches(5.96), Inches(1.7), PALE_WARM, "…or a fictional one", HIGHLIGHT,
        "Invent a plausible project instead — same funding scheme, same "
        "paperwork, none of the confidentiality worries.",
    )

    template.add_band(
        slide, Inches(3.6), Inches(1.1),
        "You will build your entire knowledge system around it during the course.",
        size=18,
        sub="One project, followed from Lesson 1 to Week 8 — pick one with a "
            "real lifecycle: funding scheme, reporting duties, stakeholders.",
    )
    template.add_label(
        slide, template.CONTENT_LEFT, Inches(5.0), CONTENT_WIDTH,
        "First step this week: write 5 sentences describing it.",
        color=HIGHLIGHT, align=PP_ALIGN.CENTER,
    )
    template.set_notes(slide, (
        "Segment 4 opens the briefing with the one decision that outlives "
        "this week: the Anchor Project. Either flavor works — a real project "
        "anonymized (the publish test from the demo rules applies to the "
        "anonymization) or a fictional one for anyone uneasy about even "
        "anonymized detail. What matters is a real lifecycle to mine: a "
        "funding scheme with rules, reporting duties, a stakeholder "
        "landscape — every later lesson populates the knowledge system "
        "from it.\n\n"
        "The 5-sentence description is deliberately small: enough to give "
        "every later prompt its context block, not a project report. "
        "Reassure them the choice isn't final-final — but changing it later "
        "means redoing groundwork, so a day of thought this week is well "
        "spent."
    ))


# --- Slide 2: the five-part assignment ---------------------------------------

# The five items from lesson_1.md Segment 4, one row each with its own
# height. Kept to one slide so participants can screenshot the whole
# assignment — which is why item 2 spells out the four pattern names.
_ASSIGNMENT_ITEMS = [
    ("Pick your Anchor Project",
     "One real (anonymized) or fictional research project — write 5 "
     "sentences describing it.", 0.85),
    ("Three ICC prompts in Gemini",
     "About your Anchor Project's domain, at least one using a pattern — "
     "persona, example-driven, step-by-step, or critique-then-improve. "
     "Save each prompt, the answer, and one sentence on what you'd "
     "improve.", 1.2),
    ("One ask-me-first prompt",
     "Let Gemini interview you about a task from your actual work, then "
     "complete the task with it.", 0.85),
    ("One meta-prompt",
     "Describe a recurring task and ask Gemini to write the optimal "
     "prompt for it; run the result and save the prompt.", 0.95),
    ("Bring all of this to Lesson 2",
     "2–3 volunteers will present at the debrief.", 0.75),
]

_ROW_GAP = 0.1


def _assignment_row(slide, number, top, heading, detail, height):
    """Numbered assignment row: navy number chip + pale item card."""
    chip = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, template.CONTENT_LEFT, Inches(top),
        Inches(0.6), Inches(height),
    )
    template.solid(chip, PRIMARY)
    tf = chip.text_frame
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.paragraphs[0].alignment = PP_ALIGN.CENTER
    run = tf.paragraphs[0].add_run()
    run.text = str(number)
    template.style(run, size=20, color=WHITE, font=template.HEADING_FONT,
                   bold=True)

    card = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(template.CONTENT_LEFT.inches + 0.8), Inches(top),
        Inches(CONTENT_WIDTH.inches - 0.8), Inches(height),
    )
    template.outlined(card, WHITE, LINE)
    tf = card.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.25)
    tf.margin_right = Inches(0.25)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    para = tf.paragraphs[0]
    para.alignment = PP_ALIGN.LEFT
    run = para.add_run()
    run.text = heading
    template.style(run, size=15, color=PRIMARY, font=template.HEADING_FONT,
                   bold=True)
    run = para.add_run()
    run.text = f"  —  {detail}"
    template.style(run, size=13, color=TEXT)


def _assignment_slide(prs):
    slide = template.add_content_slide(prs, "This week's assignment")

    top = 1.45
    for number, (heading, detail, height) in enumerate(_ASSIGNMENT_ITEMS, start=1):
        _assignment_row(slide, number, top, heading, detail, height)
        top += height + _ROW_GAP

    template.add_label(
        slide, template.CONTENT_LEFT, Inches(top + 0.02), CONTENT_WIDTH,
        "Screenshot this slide — it is the whole assignment.",
        italic=True, align=PP_ALIGN.CENTER,
    )
    template.set_notes(slide, (
        "Read all five items aloud, slowly — then pause and say the caption "
        "out loud: screenshot this slide, it is the whole assignment. "
        "Everything on it was demoed in Segment 3; nothing here is new, it "
        "is the same moves pointed at their own Anchor Project.\n\n"
        "Two clarifications worth volunteering before Q&A: for item 2, "
        "'save' means prompt + answer + one improvement sentence — a doc or "
        "notes file is fine, no special tool; the four pattern names on the "
        "slide are the ones from the patterns slide. For items 3 and 4, "
        "the task should come from their actual work — the demo rules "
        "still apply: nothing they couldn't publish goes into Gemini. "
        "Then open the floor for questions."
    ))


# --- Slide 3: effort and debrief ---------------------------------------------


def _effort_slide(prs):
    slide = template.add_content_slide(prs, "Effort and debrief")

    template.add_card(
        slide, template.CONTENT_LEFT, Inches(1.95), Inches(5.96), Inches(1.9),
        PALE, "~1–2 hours across the week", PRIMARY,
        "Not one sitting. Twenty minutes here and there — three ICC "
        "prompts one day, the ask-me-first another — beats a Sunday "
        "marathon.",
    )
    template.add_card(
        slide, Inches(template.CONTENT_LEFT.inches + 6.17), Inches(1.95),
        Inches(5.96), Inches(1.9), PALE, "Due: Lesson 2's debrief", PRIMARY,
        "Lesson 2 opens with the assignment debrief — 2–3 volunteers "
        "present one prompt and what they'd improve; everyone else "
        "brings theirs along.",
    )

    template.add_band(
        slide, Inches(4.55), Inches(1.1),
        "Volunteering is the fast lane — presenting a prompt gets you "
        "direct feedback on it.",
        size=18,
    )
    template.set_notes(slide, (
        "Set the effort expectation out loud: one to two hours across the "
        "whole week, spread out — this is practice, not a deliverable to "
        "polish. If someone is over two hours, they are overdoing it and "
        "should stop.\n\n"
        "Explain the debrief so nobody is ambushed in Lesson 2: every "
        "session from now on opens with an assignment debrief (the first "
        "slot of the Session Template), 2–3 volunteers show one prompt and what "
        "they'd improve, five minutes total, no grading. Ask for tentative "
        "volunteers now — a named volunteer shows up prepared. This slide "
        "and the previous one are the Q&A anchor; take remaining questions "
        "here before the close."
    ))


# --- Slide 4: close — the two-lane teaser ------------------------------------


def _close_slide(prs):
    slide = template.add_content_slide(prs, "Next week: the two-lane rule")

    template.add_card(
        slide, template.CONTENT_LEFT, Inches(2.15), Inches(5.96), Inches(1.6),
        PRIMARY, "The Sensitive Lane", WHITE,
    )
    card = template.add_card(
        slide, Inches(template.CONTENT_LEFT.inches + 6.17), Inches(2.15),
        Inches(5.96), Inches(1.6), PALE_WARM, "The Daily Driver", HIGHLIGHT,
    )
    # Mirror the centered heading of the dark lane card beside it.
    card.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
    template.add_label(
        slide, template.CONTENT_LEFT, Inches(4.35), CONTENT_WIDTH,
        "Which AI is allowed to see sensitive institutional data — Copilot "
        "or Gemini? And why?",
        size=18, color=PRIMARY, align=PP_ALIGN.CENTER,
    )
    template.add_label(
        slide, template.CONTENT_LEFT, Inches(5.15), CONTENT_WIDTH,
        "Lesson 2 answers that. Until then: nothing you couldn't publish.",
        italic=True, align=PP_ALIGN.CENTER,
    )
    template.set_notes(slide, (
        "The close, two minutes, no new teaching. Name the two lanes and "
        "leave them as a question: one AI is authorized for sensitive "
        "institutional data, the other handles everything else — which is "
        "which, and why, is Lesson 2's opening topic. Resist explaining "
        "the rule now; the names on the cards are the whole teaser.\n\n"
        "Repeat the interim rule one last time — until the lanes are "
        "taught, nothing you couldn't publish goes into any AI tool. Then "
        "hand over to Segment 5: thirty minutes, hands-on, everyone "
        "leaves with working tools."
    ))


def add_slides(prs):
    _anchor_project_slide(prs)
    _assignment_slide(prs)
    _effort_slide(prs)
    _close_slide(prs)
