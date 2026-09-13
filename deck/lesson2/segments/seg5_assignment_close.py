"""Segment 5 — Assignment briefing + close (8 min, Presentation + Q&A).

Eight slides, per lesson_2.md Segment 5 and in its order: the assignment,
the roadmap of the whole route, the eight steps of the blueprint in three
screenfuls, the outcome on the 7 Levels map, the effort with the two places
the work is shown, and the close.

The segment *is* the blueprint, not a description of one. "Produce a
blueprint of the tool chaining" is a sentence a participant can read twice
and still not know what to write down, so the plan authors the route in
full — eight steps, each with its tool, what goes in, what to do and what
comes out — and these slides render it. Participants do not invent an
orchestration from a sentence; they run this route on a topic of their own
and write down what they skipped or changed, and why.

Two rules make the route readable, and both are the reason the first
version of this segment was reworked. First, every step says where its
input comes from and where its output goes — "→ step 5" on the OUT card,
"from step 2" on the IN card — so the chain reads as a chain and nobody has
to guess how one step's result reaches the next. Second, only the step
that handles sensitive data is required, and only when the participant's
own material contains any; every other step is optional. A step that is
optional says so on its own row, so the route is a recommendation with one
rule in it, not a nine-item quota.

The roadmap slide is the diagram the plan draws in Mermaid: the eight
steps as two rows of four, read left to right, the fifth column holding
the two cards that are not steps — Prompt Coach, which runs beside any
step, and the deliverable the chain ends in. Drawing it natively keeps it
in the deck's own type and colour and editable in PowerPoint; the Mermaid
source stays in lesson_2.md, because the participants' own blueprint is
expected to be a diagram too and forking that source is the cheapest way
to start one.

Colour carries the routing argument, so the lane rule is load-bearing: a
step's card, its head bar and its detail cards take the colour of the
lane its tool is in — violet for the five Daily Driver steps, teal for the
one Sensitive Lane step and for Prompt Coach, neutral slate for the two
steps that use no lane at all: the manual anonymisation step, which uses
no AI tool, and the build step, whose tools sit outside both lanes and
only ever see content step 2 cleared.

Adobe Acrobat is not in the chain. The Teams app needs administrator
approval participants cannot get, there are no institutional PDFs in Teams
to review, and anonymisation is by-hand work in Word or Excel that has to
finish before any file reaches any AI tool — so the step that used to
carry Acrobat now carries no tool, on purpose.
"""

from dataclasses import dataclass

from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

from deck.lesson2 import template
from deck.lesson2.template import (
    DAILY_DRIVER, INK, MUTED, PANEL, SENSITIVE_LANE, SLATE, TEXT, WHITE,
)

# The content area, in inches, as the geometry every slide here is laid out
# on — the same grid the demo segments use, so the deck reads as one.
_LEFT = template.CONTENT_LEFT.inches
_WIDTH = template.CONTENT_WIDTH.inches
_GAP = 0.23
_COL_WIDTH = (_WIDTH - _GAP) / 2
_RIGHT_COL = _LEFT + _COL_WIDTH + _GAP
_COL3_WIDTH = (_WIDTH - 2 * _GAP) / 3

# A numbered row: a solid chip carrying the number, then the row's card.
_CHIP_WIDTH = 0.62
_CHIP_GAP = 0.18
_ROW_LEFT = _LEFT + _CHIP_WIDTH + _CHIP_GAP
_ROW_WIDTH = _WIDTH - _CHIP_WIDTH - _CHIP_GAP

# The two labels a step carries when it is in neither lane. CONTEXT.md
# defines two lanes and neither of these is a third one: the first is the
# absence of any AI tool, the second is a tool outside the tenant, which is
# why it carries the clearance condition with it wherever it appears.
_NO_TOOL = "No lane"
_NO_LANE = "Outside both lanes · anonymised content only"

_OPTIONAL = "Optional"


def _numbered_row(slide, number, top, height, heading, detail, accent=SLATE):
    """One numbered row: the number in a solid chip, then a card whose lead
    is bold and whose detail continues on the same line. Returns the card,
    so a caller that needs more than one line — the outcome map's levels do,
    the assignment's items do not — can add its own paragraphs."""
    template.add_card(
        slide, template.CONTENT_LEFT, Inches(top), Inches(_CHIP_WIDTH),
        Inches(height), accent, str(number), WHITE)

    card = template.add_open_card(
        slide, Inches(_ROW_LEFT), Inches(top), Inches(_ROW_WIDTH),
        Inches(height), accent)
    para = card.text_frame.paragraphs[0]
    run = para.add_run()
    run.text = heading
    template.style(run, size=15, color=INK, font=template.HEADING_FONT,
                   bold=True)
    run = para.add_run()
    run.text = f"  —  {detail}"
    template.style(run, size=13, color=TEXT)
    return card


# --- The blueprint the whole segment renders ---------------------------------


@dataclass(frozen=True)
class _Step:
    """One step of the workflow orchestration, as lesson_2.md's blueprint
    table and its per-step paragraph define it.

    `verb` and `node_tool` are the short forms the roadmap's cards carry;
    `title` and `tool` are the full ones the blueprint rows use. `lane` is
    the lane the step's tool belongs to, or None for a step that has no
    lane — the colour of every shape the step owns is read off it. `need`
    is the Required / Optional tag, in full. `inp` names what goes in and
    which step it comes from; `out` names what comes out and which step
    takes it; `do` is the slide-length instruction and `how` the full one
    that goes to the speaker notes.
    """

    number: int
    verb: str
    title: str
    tool: str
    node_tool: str
    lane: object
    need: str
    inp: str
    do: str
    out: str
    how: str

    @property
    def accent(self):
        return self.lane.color if self.lane else SLATE

    @property
    def lane_label(self):
        if self.lane:
            return self.lane.name
        return _NO_TOOL if self.number == 2 else _NO_LANE

    @property
    def required(self):
        return self.need != _OPTIONAL


_STEPS = [
    _Step(
        1, "Frame the topic", "Frame the topic",
        "Gemini · Guided Learning", "Gemini · Guided Learning", DAILY_DRIVER,
        _OPTIONAL,
        "A topic idea, your audience, and the length of the presentation.",
        "Give Guided Learning the ask-me-first prompt from Lesson 1 and "
        "answer its questions. Keep only outcomes you can observe: “decide "
        "whether a cost is eligible and say why”, not “understand eligible "
        "costs”. Stop at three.",
        "A topic statement and three learning outcomes → steps 5, 6, 7.",
        "Open Guided Learning and give it the ask-me-first prompt from "
        "Lesson 1: before you answer, ask me what you need to know. Answer "
        "its questions. Keep only outcomes you can observe: “decide whether "
        "a cost is eligible and say why” is one, “understand eligible "
        "costs” is not. Stop at three. Output goes to: step 5 (the "
        "outline), step 6 (the learning material) and step 7 (the visuals) "
        "are all built against these three outcomes.",
    ),
    _Step(
        2, "Anonymise your material", "Anonymise your own material",
        "No AI tool · Word or Excel, by hand", "Word or Excel, by hand", None,
        "Required if the material contains sensitive data",
        "Your own documents: notes, tables, call texts, reports.",
        "Work on a copy. Go through each file by hand with the Lesson 1 "
        "anonymisation checklist: replace names, IDs, amounts, dates, "
        "anything that identifies a person or a project. Record per file: "
        "anonymised, public, or not usable. Nothing reaches any AI tool "
        "before this is done. No material of your own? Skip this step and "
        "say so.",
        "Files that may be uploaded anywhere → steps 4, 5.",
        "Work on a copy. Go through each file by hand in Word or Excel with "
        "the Lesson 1 anonymisation checklist: replace names, IDs, amounts, "
        "dates and anything that identifies a person, a project or a "
        "partner. Record per file: anonymised, public, or not usable. No AI "
        "tool is involved, because nothing may reach an AI tool before this "
        "step is done. Output goes to: step 4 (the notebook) and step 5 "
        "(the Gem's reference files). If you use no material of your own, "
        "skip this step and say so in the blueprint.",
    ),
    _Step(
        3, "Research the topic", "Research the topic",
        "Gemini · Deep Research", "Gemini · Deep Research", DAILY_DRIVER,
        _OPTIONAL,
        "One research question: what your material does not cover, or the "
        "topic itself if you have no material.",
        "One question, not five. Ask for a source per claim. Open at least "
        "three sources and check they say what the report says. The free "
        "tier limits reports per month, so write the question before you "
        "spend one.",
        "A cited briefing, saved as a file → step 4.",
        "Write one research question, not five. If you have your own "
        "material, ask what it does not cover; if not, ask about the topic "
        "itself. Ask for a source per claim. Open at least three of the "
        "sources and check they say what the report says. The free tier "
        "limits reports per month, so write the question before you spend "
        "one. Output goes to: step 4, as one more file in the notebook.",
    ),
    _Step(
        4, "Collect the content", "Collect what the presentation must cover",
        "NotebookLM", "NotebookLM", DAILY_DRIVER, _OPTIONAL,
        "The files from step 2 and the briefing from step 3.",
        "One notebook per topic. Add the files, then ask: what does a "
        "newcomer need to know about this, and which source says it? Drop "
        "any answer without a citation. No NotebookLM? Upload the same "
        "files to Gemini and say so in your blueprint.",
        "A coverage list, each point with its source → step 5.",
        "One notebook per topic. Add the files from step 2 and the briefing "
        "from step 3, then ask: what does a newcomer need to know about "
        "this, and which source says it? Drop any answer without a "
        "citation. If NotebookLM is not available to you this fortnight, "
        "upload the same files to Gemini, ask the same question, and say so "
        "in your blueprint. Output goes to: step 5, as the list of what the "
        "outline has to contain.",
    ),
    _Step(
        5, "Draft the outline", "Draft the outline",
        "Gemini · a Gem with file upload, long context and Canvas",
        "Gemini · Gem + Canvas", DAILY_DRIVER, _OPTIONAL,
        "The outcomes from step 1, the coverage list from step 4, the "
        "files from step 2.",
        "Build a Gem: topic and outcomes as its instructions, the files as "
        "its reference files. Ask for an outline in Canvas and edit one "
        "section at a time. Tag each section with the outcome it serves; "
        "cut any section that serves none.",
        "A section-by-section outline → steps 6, 7, 8.",
        "Build a Gem first: the topic statement and the three outcomes as "
        "its instructions, the files from step 2 as its reference files; "
        "long context means it reads all of them. Ask for an outline in "
        "Canvas, then edit one section at a time instead of regenerating "
        "the whole thing. Tag each section with the outcome it serves and "
        "cut any section that serves none. Output goes to: step 6 "
        "(learning material), step 7 (visuals) and step 8 (the deck).",
    ),
    _Step(
        6, "Add learning material", "Add the learning material",
        "Study and Learn · Copilot Chat (Basic)", "Study and Learn",
        SENSITIVE_LANE, _OPTIONAL,
        "The outline from step 5, one learning outcome at a time.",
        "Per outcome: ask for an explanation and read it as a learner "
        "would; then flashcards on the key terms; then a five-question "
        "quiz. Drop any question that can be answered without seeing the "
        "presentation. Needs a Microsoft 365 Education account, no Copilot "
        "licence.",
        "An explanation, a flashcard set and a quiz per outcome → step 8.",
        "One outcome at a time. Ask for an explanation and read it as a "
        "learner would; then flashcards on the key terms; then a "
        "five-question quiz. Drop any question that can be answered "
        "without having seen the presentation. Needs a Microsoft 365 "
        "Education account; no Copilot licence. Output goes to: step 8, as "
        "its own section of the deck.",
    ),
    _Step(
        7, "Make the visuals", "Make the visuals",
        "Gemini · image generation", "Gemini · image generation",
        DAILY_DRIVER, _OPTIONAL,
        "One section of the outline from step 5 and a one-sentence "
        "description of its diagram.",
        "Describe the relationship, not the picture: “three stages, each "
        "feeding the next, the middle one highlighted”. Read the labels in "
        "the result; text in generated images is often wrong. Regenerate, "
        "do not patch. Not every section needs an image.",
        "One image per section that needs one → step 8.",
        "Describe the relationship you want shown, not the picture: “three "
        "stages, each feeding the next, the middle one highlighted”. Read "
        "the labels in the result; text inside generated images is often "
        "wrong, and the fix is to regenerate, not to patch. Not every "
        "section needs an image. Output goes to: step 8.",
    ),
    _Step(
        8, "Build the presentation", "Build the presentation",
        "Canva (Teams app) → genspark.ai", "Canva → genspark.ai", None,
        "Optional tools, required output",
        "The outline (5), the learning material (6), the visuals (7).",
        "Prototype the slide structure in Canva and let a colleague comment "
        "before content goes in. Cut a 20–40 second video excerpt there for "
        "the opening. Give outline, material and visuals to genspark.ai to "
        "build the deck. Rehearse once out loud. AI-use log on the last "
        "slide.",
        "The finished presentation and a short video excerpt: the "
        "deliverable.",
        "Prototype the slide structure in Canva and let a colleague comment "
        "on it before any content goes in. Cut a 20–40 second video excerpt "
        "in Canva for the opening. Then give the outline, the learning "
        "material and the visuals to genspark.ai and build the deck. "
        "Rehearse once out loud. Put the AI-use log on the last slide. Both "
        "tools sit outside the institutional tenant: only content that "
        "passed step 2 goes into them. Output goes to: nowhere. This is the "
        "deliverable.",
    ),
]

# Prompt Coach is not a step: it runs beside whichever step the participant
# is on. The plan gives it its own paragraph after the eight steps.
_PROMPT_COACH = (
    "Before you run a prompt in steps 3 to 6, paste it into Prompt Coach "
    "with one sentence about the job it has to do, ask for a critique and "
    "an improved version, and keep the improved one only if you can say "
    "what changed. Copilot Chat (Basic), Sensitive Lane; nothing but the "
    "prompt goes in."
)


# --- The assignment, on one slide -------------------------------------------

# The plan's three numbered items, verbatim in substance: what to produce,
# what the blueprint has to carry, and how every step is routed. The
# blueprint item points at the roadmap rather than describing one, because
# the roadmap is the next four slides.
_ASSIGNMENT_ITEMS = [
    ("Deliverable",
     "A presentation on a hand-picked topic, designed for educational "
     "purposes and intended to facilitate learning outcomes."),
    ("Blueprint",
     "The roadmap on the next four slides, adapted to your topic. For each "
     "step you run: the tool, the input, the output and how you ran it. "
     "For each step you skip or change: one line saying why."),
    ("Routing",
     "Every step that touches sensitive data goes to the correct lane. "
     "Documents are anonymised before upload. The blueprint states the "
     "lane per step."),
]

_ITEM_TOP = 2.26
_ITEM_HEIGHT = 0.92
_ITEM_PITCH = 1.00

# What actually gets handed in, as three things rather than one word.
_HAND_IN = [
    ("The blueprint",
     "The step table, a line per skipped or changed step, and the chain "
     "drawn as a diagram."),
    ("The presentation",
     "On your topic, teaching something, with the learning material from "
     "step 6 if you ran it."),
    ("The AI-use log",
     "Tool, date and purpose per step, as the declaration on your last "
     "slide."),
]


def _assignment_slide(prs):
    slide = template.add_content_slide(
        prs, "Assignment: your first workflow orchestration")

    template.add_band(
        slide, Inches(1.28), Inches(0.82),
        "Run the roadmap on a topic of your own", size=21,
        sub="Chain the tools the course has introduced so far into one "
            "multi-step workflow orchestration, end to end, on a topic you "
            "pick.",
    )

    top = _ITEM_TOP
    for number, (heading, detail) in enumerate(_ASSIGNMENT_ITEMS, start=1):
        _numbered_row(slide, number, top, _ITEM_HEIGHT, heading, detail)
        top += _ITEM_PITCH

    template.add_label(
        slide, template.CONTENT_LEFT, Inches(5.32), template.CONTENT_WIDTH,
        "What you hand in", size=11, color=MUTED,
    )

    for index, (heading, body) in enumerate(_HAND_IN):
        template.add_detail_card(
            slide, Inches(_LEFT + index * (_COL3_WIDTH + _GAP)), Inches(5.62),
            Inches(_COL3_WIDTH), Inches(0.94), SLATE, heading, [body],
            heading_size=12.5, line_size=11)

    template.add_label(
        slide, template.CONTENT_LEFT, Inches(6.66), template.CONTENT_WIDTH,
        "The roadmap is next. Only one step is required, and only if your "
        "material contains sensitive data; every other step is optional.",
        size=12.5, color=MUTED, italic=True, align=PP_ALIGN.CENTER,
    )

    template.set_notes(slide, (
        "~1 minute. Read the three items aloud and then get off this slide "
        "— the roadmap is where the time goes.\n\n"
        "The one thing to say plainly: the blueprint is not homework you "
        "invent from scratch. It is on the next four slides, written out, "
        "step by step. What they owe is a run of it on their own topic and "
        "a line of reasoning for every step they skip or change. Say that "
        "out loud — last year's version of this assignment asked people to "
        "produce a blueprint from the word “blueprint”, and what came back "
        "was nine rows of tool names with no reasoning in them.\n\n"
        "Item 1 is the deliverable: a presentation that teaches, on a topic "
        "they pick. Say the topic is theirs — it does not have to be the "
        "Anchor Project, and a topic they already know well makes the "
        "learning-outcome part easier to judge.\n\n"
        "Item 3 is where this can go wrong quietly. A chain does not "
        "inherit a lane from its first step: each step is routed on its "
        "own. If a step cannot be cleared, the step changes — not the "
        "lane.\n\n"
        "Say the footer line: only step 2 is required, and only when their "
        "own material contains sensitive data. Everything else is optional, "
        "and skipping a step is fine as long as the blueprint says why.\n\n"
        "Paste this slide's text and the blueprint into the Teams chat "
        "straight after the session, together with the Mermaid source from "
        "the plan, so nobody is working from a screenshot alone."
    ))
    return slide


# --- The roadmap -------------------------------------------------------------

# The roadmap grid: five columns of cards with an arrow glyph in each
# gutter, read left to right, top row then bottom. Steps 1–4 fill the top
# row and 5–8 the bottom one; the fifth column holds the two cards that
# are not steps — Prompt Coach above, the deliverable below.
_NODE_COLUMNS = 5
_NODE_GUTTER = 0.40
_NODE_WIDTH = (_WIDTH - (_NODE_COLUMNS - 1) * _NODE_GUTTER) / _NODE_COLUMNS
_NODE_HEIGHT = 1.42
_ROW_ONE_TOP = 1.90
_ROW_TWO_TOP = 3.72


def _node_left(column):
    return _LEFT + column * (_NODE_WIDTH + _NODE_GUTTER)


def _node_card(slide, column, top, edge, kicker, heading, tool, tag=None,
               tag_color=MUTED):
    """One card in the roadmap grid: a small kicker, the bold heading, the
    tool in the lane's colour, and an optional last line — the step's
    Required / Optional tag."""
    card = template.add_open_card(
        slide, Inches(_node_left(column)), Inches(top), Inches(_NODE_WIDTH),
        Inches(_NODE_HEIGHT), edge)
    tf = card.text_frame
    tf.vertical_anchor = MSO_ANCHOR.TOP
    tf.margin_top = Inches(0.12)

    run = tf.paragraphs[0].add_run()
    run.text = kicker
    template.style(run, size=8.5, color=MUTED, bold=True)

    para = tf.add_paragraph()
    para.space_before = Pt(2)
    run = para.add_run()
    run.text = heading
    template.style(run, size=12.5, color=INK, font=template.HEADING_FONT,
                   bold=True)

    para = tf.add_paragraph()
    para.space_before = Pt(3)
    run = para.add_run()
    run.text = tool
    template.style(run, size=9.5, color=edge, bold=True)

    if tag:
        para = tf.add_paragraph()
        para.space_before = Pt(3)
        run = para.add_run()
        run.text = tag
        template.style(run, size=9, color=tag_color, bold=True, italic=True)
    return card


def _step_node(slide, column, top, step):
    """One step as a card in the roadmap grid: its number, the short verb
    for what it does, the tool it runs in — in the lane's colour — and
    whether it is required."""
    tag = "Required if sensitive data" if step.required else _OPTIONAL
    if step.number == 8:
        tag = "Optional tools · required output"
    return _node_card(
        slide, column, top, step.accent, f"STEP {step.number}", step.verb,
        step.node_tool, tag, INK if step.required else MUTED)


def _helper_node(slide, column, top):
    """The card for Prompt Coach, which is not a step: it runs beside
    whichever step the participant is on."""
    return _node_card(
        slide, column, top, SENSITIVE_LANE.color, "ANY STEP",
        "Improve the prompt first", "Prompt Coach",
        "Optional · steps 3 to 6")


def _terminal_node(slide, column, top):
    """The card the chain ends in — the deliverable, not a step, so it is a
    solid block rather than another white card."""
    block = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(_node_left(column)), Inches(top),
        Inches(_NODE_WIDTH), Inches(_NODE_HEIGHT))
    template.solid(block, INK)
    tf = block.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.22)
    tf.margin_right = Inches(0.18)
    tf.margin_top = Inches(0.12)
    tf.vertical_anchor = MSO_ANCHOR.TOP

    run = tf.paragraphs[0].add_run()
    run.text = "OUT"
    template.style(run, size=8.5, color=WHITE, bold=True)

    para = tf.add_paragraph()
    para.space_before = Pt(2)
    run = para.add_run()
    run.text = "The presentation"
    template.style(run, size=12.5, color=WHITE, font=template.HEADING_FONT,
                   bold=True)

    para = tf.add_paragraph()
    para.space_before = Pt(3)
    run = para.add_run()
    run.text = "+ the blueprint + the AI-use log"
    template.style(run, size=9.5, color=WHITE)
    return block


def _arrow(slide, column, top):
    """The flow marker in the gutter to the right of a column."""
    template.add_label(
        slide, Inches(_node_left(column) + _NODE_WIDTH),
        Inches(top + _NODE_HEIGHT / 2 - 0.17), Inches(_NODE_GUTTER), "→",
        size=15, color=MUTED, align=PP_ALIGN.CENTER)


# The colour key for the roadmap, read off the lanes themselves so a lane's
# name and its colour cannot drift apart.
_LEGEND = [
    (DAILY_DRIVER.color,
     f"{DAILY_DRIVER.name} · {DAILY_DRIVER.tool} — steps 1, 3, 4, 5, 7"),
    (SENSITIVE_LANE.color,
     f"{SENSITIVE_LANE.name} · {SENSITIVE_LANE.tool} — step 6, Prompt Coach"),
    (SLATE, "No lane — step 2 (no AI tool), step 8 (outside both lanes)"),
]


def _roadmap_slide(prs):
    slide = template.add_content_slide(
        prs, "The roadmap: eight steps, end to end")

    template.add_label(
        slide, template.CONTENT_LEFT, Inches(1.24), template.CONTENT_WIDTH,
        "Three phases: prepare (1–2), gather (3–4), produce (5–8). Each "
        "card says which tool runs the step and whether you must run it. "
        "The colour is the lane: each step is routed on its own, and no "
        "step inherits a lane from the one before it.",
        size=12, color=MUTED,
    )

    for column, step in enumerate(_STEPS[:4]):
        _step_node(slide, column, _ROW_ONE_TOP, step)
        if column < 3:
            _arrow(slide, column, _ROW_ONE_TOP)
    _helper_node(slide, 4, _ROW_ONE_TOP)

    template.add_label(
        slide, template.CONTENT_LEFT, Inches(3.42), template.CONTENT_WIDTH,
        "…THEN", size=10.5, color=SLATE)
    for column, step in enumerate(_STEPS[4:]):
        _step_node(slide, column, _ROW_TWO_TOP, step)
        _arrow(slide, column, _ROW_TWO_TOP)
    _terminal_node(slide, 4, _ROW_TWO_TOP)

    template.add_label(
        slide, template.CONTENT_LEFT, Inches(5.22), template.CONTENT_WIDTH,
        "Only step 2 is required, and only when your material contains "
        "sensitive data. Skip any other step your topic does not need, and "
        "say why in the blueprint. Steps 6 and 7 can run in either order.",
        size=11, color=MUTED, italic=True,
    )

    template.add_band(
        slide, Inches(5.64), Inches(0.66),
        "Before every step: check every output before it becomes the next "
        "step's input", size=15,
        sub="…and log the tool, the date and what it was used for. That log "
            "is the AI-use declaration on your last slide.",
    )

    left = _LEFT
    for color, text in _LEGEND:
        template.add_bar(slide, Inches(left), Inches(6.44), Inches(0.16),
                         Inches(0.16), color)
        template.add_label(
            slide, Inches(left + 0.26), Inches(6.36), Inches(3.9), text,
            size=10, color=TEXT)
        left += 4.10

    template.add_label(
        slide, template.CONTENT_LEFT, Inches(6.76), template.CONTENT_WIDTH,
        "Screenshot this slide and the next three — together they are the "
        "assignment.",
        size=12, color=MUTED, italic=True, align=PP_ALIGN.CENTER,
    )

    template.set_notes(slide, (
        "~2 minutes, and this is the slide the segment exists for. Walk the "
        "route once, left to right, naming the tool at each card and "
        "nothing else — eight names in ninety seconds. The detail is on the "
        "next three slides and in the chat; do not read it here.\n\n"
        "Say the three phases: prepare, gather, produce. Then say the one "
        "rule: only step 2 is required, and only when their own material "
        "contains sensitive data. Everything else is optional, and a "
        "skipped step with a reason is a better blueprint than a step run "
        "for the sake of it.\n\n"
        "Then make the argument the colours make for you: violet is "
        "Gemini, teal is Copilot, slate is no lane at all. Step 2 has no "
        "AI tool on purpose: anonymisation is by-hand work in Word or "
        "Excel, and it has to be finished before any file reaches any AI "
        "tool. If someone asks why Adobe Acrobat is gone from the chain: "
        "the Teams app needs administrator approval nobody here can get, "
        "and there are no institutional PDFs in Teams to review.\n\n"
        "Step 8 is in neither lane: Canva and genspark.ai sit outside the "
        "institutional tenant, which is exactly why only what step 2 "
        "cleared may go into them.\n\n"
        "Prompt Coach is the card in the top right, and it is not a step: "
        + _PROMPT_COACH + "\n\n"
        "The diagram is part of the deliverable. Say that the Mermaid "
        "source for this picture is in the session plan and goes into the "
        "chat: they are expected to redraw it for their own chain, and a "
        "chain that cannot be drawn has not been thought through.\n\n"
        "Leave this slide up for the Q&A at the end of the segment."
    ))
    return slide


# --- The blueprint, three steps to a slide -----------------------------------

# The three fields a blueprint row shows, in proportion so they always fill
# the content width: what goes in and which step it comes from, what to do,
# and what comes out and which step takes it. The full instruction is in
# the speaker notes and the handout, because a paragraph does not fit a
# card.
_FIELD_SHARES = [("IN", "inp", 3), ("DO", "do", 5), ("OUT", "out", 3)]
_FIELD_GAP = 0.12

_BLUEPRINT_TOP = 1.60
_BLUEPRINT_PITCH = 1.76
_HEAD_HEIGHT = 0.38
_FIELD_TOP_GAP = 0.06
_FIELD_HEIGHT = 1.24


def _blueprint_row(slide, top, step):
    """One step, whole: a solid head bar in the lane's colour carrying the
    step, its tool, its lane and whether it is required, and under it the
    three fields as detail cards."""
    head = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, template.CONTENT_LEFT, Inches(top),
        template.CONTENT_WIDTH, Inches(_HEAD_HEIGHT))
    template.solid(head, step.accent)
    tf = head.text_frame
    tf.word_wrap = False
    tf.margin_left = Inches(0.22)
    tf.margin_right = Inches(0.22)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    para = tf.paragraphs[0]
    para.alignment = PP_ALIGN.LEFT
    run = para.add_run()
    run.text = f"{step.number} · {step.title}"
    template.style(run, size=14, color=WHITE, font=template.HEADING_FONT,
                   bold=True)

    tag = template.add_label(
        slide, template.CONTENT_LEFT, Inches(top + 0.07),
        template.CONTENT_WIDTH - Inches(0.22),
        f"{step.tool}   ·   {step.lane_label}   ·   {step.need}",
        size=10.5, color=WHITE, align=PP_ALIGN.RIGHT)
    tag.text_frame.word_wrap = False

    span = _WIDTH - (len(_FIELD_SHARES) - 1) * _FIELD_GAP
    shares = sum(share for _, _, share in _FIELD_SHARES)
    left = _LEFT
    field_top = top + _HEAD_HEIGHT + _FIELD_TOP_GAP
    for heading, attribute, share in _FIELD_SHARES:
        width = span * share / shares
        template.add_detail_card(
            slide, Inches(left), Inches(field_top), Inches(width),
            Inches(_FIELD_HEIGHT), step.accent, heading,
            [getattr(step, attribute)], heading_size=10, line_size=10)
        left += width + _FIELD_GAP


# The three blueprint slides: which steps each carries, and the caption
# that says what that group of steps is for.
_BLUEPRINT_SLIDES = [
    ((1, 3), "The blueprint · steps 1–3",
     "Prepare and start gathering: fix the topic, make your own material "
     "safe to upload, and research what it does not cover"),
    ((4, 6), "The blueprint · steps 4–6",
     "Collect what the presentation must say, draft the outline in a "
     "workspace that knows your topic, and add the material that makes it "
     "teach"),
    ((7, 8), "The blueprint · steps 7–8",
     "Make the visuals and build the deck. Prompt Coach runs beside any "
     "step: improve the prompt before you use it"),
]


def _blueprint_slide_builder(span, title, caption):
    first, last = span
    steps = _STEPS[first - 1:last]

    def build(prs):
        slide = template.add_content_slide(prs, title)
        template.add_label(
            slide, template.CONTENT_LEFT, Inches(1.22),
            template.CONTENT_WIDTH, caption, size=12.5, color=MUTED)

        top = _BLUEPRINT_TOP
        for step in steps:
            _blueprint_row(slide, top, step)
            top += _BLUEPRINT_PITCH

        if last == 8:
            template.add_detail_card(
                slide, template.CONTENT_LEFT, Inches(top),
                template.CONTENT_WIDTH, Inches(1.24), SENSITIVE_LANE.color,
                "Prompt Coach · Copilot Chat (Basic) · Sensitive Lane · "
                "at any step, optional",
                [_PROMPT_COACH], heading_size=11, line_size=10.5)

        template.set_notes(slide, (
            "~30 seconds on this slide. Do not read the cards — point at "
            "the IN and OUT columns and say that every output names the "
            "step that takes it. The full instructions below go into the "
            "Teams chat, not onto the screen.\n\n"
            + "\n\n".join(
                f"Step {step.number} — {step.title} "
                f"({step.tool}; {step.lane_label}; {step.need}). "
                f"In: {step.inp} How: {step.how} Out: {step.out}"
                for step in steps
            )
            + ("\n\nPrompt Coach, at any step. " + _PROMPT_COACH
               if last == 8 else "")
        ))
        return slide

    return build


# --- The outcome, on the 7 Levels map ---------------------------------------

# Three rows of the 7 Levels map from lesson_1.md: the level just left, the
# level reached this week, and the level the roadmap is the foundation of.
# The names and the one-line descriptions are Lesson 1's, so the room
# recognises the map rather than meeting a new one.
@dataclass(frozen=True)
class _Level:
    """One row of the 7 Levels map as this slide shows it: the level as
    Lesson 1 named and described it, plus where the course stands on it.
    `chip` is the number chip's fill and the card's edge; `status_color`
    is the weight the standing line carries — the level reached this week
    is the darkest of the three, the level already behind them the
    lightest."""

    number: int
    name: str
    description: str
    status: str
    chip: object
    status_color: object


_LEVELS = [
    _Level(2, "Prompt Crafter",
           "Structured prompts, examples, iteration.",
           "Done — that was last week's assignment.", SLATE, MUTED),
    _Level(3, "Power User",
           "Persistent context: dedicated workspaces that already know your "
           "domain.",
           "Reached this week — as soon as the Gem you saw built is your "
           "own, which is step 5 of the roadmap.", INK, INK),
    _Level(4, "Tool Orchestrator",
           "The right specialised tool per step: NotebookLM for documents, "
           "Perplexity for cited research.",
           "Founded, not yet reached — the roadmap is the orchestration "
           "written down; Lessons 3 and 4 hand you the specialised tools to "
           "put in it.", SLATE, SLATE),
]

_LEVEL_TOP = 1.78
_LEVEL_HEIGHT = 1.06
_LEVEL_PITCH = 1.18


def _outcome_slide(prs):
    slide = template.add_content_slide(
        prs, "The outcome, on the 7 Levels map")

    template.add_label(
        slide, template.CONTENT_LEFT, Inches(1.30), template.CONTENT_WIDTH,
        "Three rows of the map from Lesson 1 — where last week left you, "
        "where this week puts you, and what the roadmap is the foundation "
        "of",
        size=13, color=MUTED,
    )

    top = _LEVEL_TOP
    for level in _LEVELS:
        card = _numbered_row(
            slide, level.number, top, _LEVEL_HEIGHT,
            f"Level {level.number} · {level.name}", level.description,
            accent=level.chip)
        para = card.text_frame.add_paragraph()
        run = para.add_run()
        run.text = level.status
        template.style(run, size=12, color=level.status_color, bold=True)
        top += _LEVEL_PITCH

    template.add_band(
        slide, Inches(5.42), Inches(0.94),
        "Level 3 achieved · foundation for Level 4 established", size=20,
        sub="A persistent, domain-aware workspace — and the right tool per "
            "step, chained.",
    )

    template.add_label(
        slide, template.CONTENT_LEFT, Inches(6.46), template.CONTENT_WIDTH,
        "Nobody reaches a level by owning a tool. Level 3 is a workspace you "
        "work inside; Level 4 is a chain you can explain.",
        size=11.5, color=MUTED, italic=True,
    )

    template.set_notes(slide, (
        "~45 seconds. This is the answer to “what did this week get me?”, "
        "and it is worth stating because the answer is not “two more "
        "tools”.\n\n"
        "Point at the map, not at the slide's wording: they saw these rows "
        "in Lesson 1 and have not seen them since. Level 2 was last week — "
        "structured prompts, iteration, done. Level 3 is today, and the "
        "thing that earns it is the Gem, not the number of tools they now "
        "have accounts for. Be honest about the tense: they watched one "
        "built, and they are on Level 3 once they have their own — which is "
        "step 5 of the roadmap.\n\n"
        "Be precise about Level 4: the assignment does not reach it, it "
        "founds it. Orchestration is the right specialised tool per step, "
        "and the specialised tools arrive in Lessons 3 and 4 — NotebookLM, "
        "Perplexity. What the roadmap gives them is the habit of saying "
        "why a step is there, which is the part that does not come with a "
        "tool.\n\n"
        "Say the closing line as written if the room looks pleased with "
        "itself: nobody reaches a level by owning a tool.\n\n"
        "If someone asks where Levels 5–7 are, one sentence: Level 5 is "
        "Lesson 5's Agent Builder, and 6 and 7 are the final session's "
        "outlook. Do not open the map here."
    ))
    return slide


# --- Effort and where the work is shown (≈1 min) -----------------------------

_WHERE_SHOWN = [
    "Practice Week 1's drop-in (W3) — troubleshooting a half-built chain",
    "The Lesson 3 debrief (W4) — 2–3 volunteers present blueprint and "
    "presentation",
]

# What "done" looks like, as three things that can be checked rather than a
# feeling. The middle one is the deliverable; the other two are the parts
# people skip.
_DONE = [
    ("A blueprint you can follow",
     "Every step you ran says what went in, what came out and which tool "
     "ran it. Every step you skipped or changed says why. The chain is "
     "drawn, not just listed."),
    ("A presentation that teaches",
     "It serves three stated learning outcomes, and carries the learning "
     "material from step 6 if you ran it: an explanation, flashcards and "
     "a quiz."),
    ("An AI-use log",
     "Tool, date and purpose per step, ending up as the declaration on the "
     "presentation's last slide."),
]


def _effort_slide(prs):
    slide = template.add_content_slide(
        prs, "Effort and where the work is shown")

    template.add_label(
        slide, template.CONTENT_LEFT, Inches(1.26), template.CONTENT_WIDTH,
        "Two weeks, the blueprint first — and two places where the work gets "
        "looked at",
        size=13, color=MUTED,
    )

    effort = template.add_card(
        slide, template.CONTENT_LEFT, Inches(1.62), Inches(_COL_WIDTH),
        Inches(1.52), PANEL, "3–4 hours across the two weeks", INK,
        "The blueprint first, then the presentation. Eight steps is the "
        "recommended route, not a quota: only step 2 is required, and only "
        "if your material contains sensitive data. Twenty minutes at a time "
        "beats one long Sunday.",
    )
    # Top-anchored, so its heading shares a baseline with the list card
    # beside it — add_card centres its text, add_list_card does not.
    effort.text_frame.vertical_anchor = MSO_ANCHOR.TOP
    effort.text_frame.margin_top = Inches(0.16)
    template.add_list_card(
        slide, Inches(_RIGHT_COL), Inches(1.62), Inches(_COL_WIDTH),
        Inches(1.52), PANEL, "Where the work is shown", _WHERE_SHOWN,
        size=12.5)

    template.add_label(
        slide, template.CONTENT_LEFT, Inches(3.32), template.CONTENT_WIDTH,
        "What “done” looks like", size=11, color=MUTED,
    )

    for index, (heading, body) in enumerate(_DONE):
        template.add_detail_card(
            slide, Inches(_LEFT + index * (_COL3_WIDTH + _GAP)), Inches(3.62),
            Inches(_COL3_WIDTH), Inches(1.52), SLATE, heading, [body],
            heading_size=13, line_size=11.5)

    template.add_band(
        slide, Inches(5.36), Inches(0.86),
        "A chain does not inherit a lane from its first step", size=19,
        sub="Every step is routed on its own — and nothing that skipped "
            "step 2 goes into Canva or genspark.ai.",
    )

    template.add_label(
        slide, template.CONTENT_LEFT, Inches(6.38), template.CONTENT_WIDTH,
        "Stuck on a step? That is what Practice Week 1's drop-in is for — "
        "bring the blueprint even if nothing runs yet.",
        size=11.5, color=MUTED, italic=True, align=PP_ALIGN.CENTER,
    )

    template.set_notes(slide, (
        "~1 minute. Set the effort expectation out loud: 3–4 hours across "
        "the two weeks, spread out, blueprint first. Anyone past four hours "
        "is polishing the presentation instead of writing the blueprint, "
        "and the blueprint is the assignment.\n\n"
        "Say the quota line, because somebody always asks whether all "
        "eight steps are mandatory: they are not. Only step 2 is, and only "
        "when their own material contains sensitive data. Skipping or "
        "merging steps is fine; the line saying why is what is not "
        "optional.\n\n"
        "“Done” is three things and the middle one is the only one people "
        "remember. Read the first and third aloud — a blueprint you can "
        "follow, and an AI-use log. The log is the directive's "
        "transparency point, not bureaucracy: where AI was used, which "
        "tool, when, for what.\n\n"
        "Both venues, plainly. Practice Week 1's drop-in in W3 is optional "
        "and is the place for a half-built chain. The Lesson 3 debrief in "
        "W4 is where 2–3 volunteers present blueprint and presentation. Ask "
        "for tentative volunteers now, in the chat; a named volunteer "
        "arrives prepared.\n\n"
        "Anyone who did not finish: the drop-in, not an apology in the "
        "chat. Say it once, lightly."
    ))
    return slide


# --- Close (≈1 min) ----------------------------------------------------------

_NEXT_SESSIONS = [
    ("Practice Week 1 · W3",
     ["An optional, facilitated drop-in — no new material, and no "
      "obligation to come",
      "Bring the blueprint and the lane questions it raised",
      "Bring any Gemini free-tier surprise you hit — that is what the hour "
      "is for"]),
    ("Lesson 3 · W4",
     ["NotebookLM properly, not at scenario level",
      "Perplexity for cited research you can trace back to a source",
      "The Wiki Sandbox — the first piece of your KM System, and what the "
      "rest of the course is built on"]),
]


def _close_slide(prs):
    slide = template.add_content_slide(
        prs, "Next: Practice Week 1 and Lesson 3")

    template.add_label(
        slide, template.CONTENT_LEFT, Inches(1.30), template.CONTENT_WIDTH,
        "No new material next week. The week after, the first piece of the "
        "system this course is building",
        size=13, color=MUTED,
    )

    for index, (heading, items) in enumerate(_NEXT_SESSIONS):
        card = template.add_list_card(
            slide, Inches(_LEFT if index == 0 else _RIGHT_COL), Inches(1.86),
            Inches(_COL_WIDTH), Inches(2.00), PANEL, heading, items,
            size=13)
        card.text_frame.margin_top = Inches(0.18)

    template.add_band(
        slide, Inches(4.45), Inches(1.06),
        "A workflow you can explain, step by step", size=21,
        sub="Bring the blueprint to the drop-in, and the whole chain to the "
            "Lesson 3 debrief.",
    )

    template.add_label(
        slide, template.CONTENT_LEFT, Inches(5.80), template.CONTENT_WIDTH,
        "Questions now — the roadmap stays up.",
        size=13, color=MUTED, italic=True, align=PP_ALIGN.CENTER,
    )

    template.set_notes(slide, (
        "~1 minute, no new teaching, and then Q&A on whatever is left.\n\n"
        "Practice Week 1 is next week and it is exactly what its name says: "
        "no new material, the session slot runs as an optional facilitated "
        "drop-in, and the reason it exists is that this course runs during "
        "office hours. Say that anyone stuck on a lane question or a "
        "half-built chain should come; anyone whose chain already runs does "
        "not have to.\n\n"
        "Lesson 3 in W4 is the payoff for the NotebookLM scenarios they "
        "just watched — the tool properly this time, Perplexity beside it "
        "for cited research, and the Wiki Sandbox, which is the first piece "
        "of the KM System they will still be building in Week 8.\n\n"
        "Go back to the roadmap slide for the questions, and post the "
        "blueprint and the Mermaid source in the chat before you close the "
        "meeting."
    ))
    return slide


# --- Segment assembly --------------------------------------------------------

# Slide order per lesson_2.md Segment 5: the assignment, the roadmap, the
# blueprint in three screenfuls, the outcome it produces, the effort it
# costs and where it is shown, then the close.
_BUILDERS = [
    _assignment_slide,
    _roadmap_slide,
    *[_blueprint_slide_builder(*entry) for entry in _BLUEPRINT_SLIDES],
    _outcome_slide,
    _effort_slide,
    _close_slide,
]


def add_slides(prs):
    for builder in _BUILDERS:
        builder(prs)
