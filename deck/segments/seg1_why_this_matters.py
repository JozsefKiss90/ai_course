"""Segment 1 — Opening: why this matters (10 min, Presentation).

Four slides per lesson_1.md Segment 1: the Tufts Digital Planet headline
figures, the sector chart showing where the risk concentrates, the
augmentation–displacement pipeline — claim, the researchers' scatter as its
evidence, and the verbatim quote on one slide — and the turn, the course
thesis framed for research management administrators. Not fear — skill.

Both chart slides place the researchers' own graphics (.scratch/
aijobloss.png, aijobloss2.png) on a white mat under a compact header;
everything else on the slide is drawn in the deck's own language.
"""

import pathlib

from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

from deck import template
from deck.template import (
    ACCENT, AMBER, HEADING_FONT, HIGHLIGHT, MUTED, PALE, PALE_WARM, PANEL_EDGE,
    PRIMARY, SKY, SLATE, TEXT, WHITE,
)

TUFTS_CREDIT = (
    "Source: Tufts Digital Planet — AI and the Emerging Geography of "
    "American Job Risk"
)


def _stat_tile(slide, left, width, number, caption, detail=None):
    """Big-number tile: the figure dominates, the caption explains it."""
    tile = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, left, Inches(1.95), width, Inches(1.62)
    )
    template.panel(tile)
    template.add_bar(slide, left, Inches(1.95), width, Inches(0.09), AMBER)
    tf = tile.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.4)
    tf.margin_right = Inches(0.4)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    para = tf.paragraphs[0]
    para.alignment = PP_ALIGN.CENTER
    run = para.add_run()
    run.text = number
    template.style(run, size=38, color=AMBER, font=HEADING_FONT, bold=True)
    para = tf.add_paragraph()
    para.alignment = PP_ALIGN.CENTER
    para.space_before = Pt(6)
    run = para.add_run()
    run.text = caption
    template.style(run, size=14, color=WHITE)
    if detail:
        para = tf.add_paragraph()
        para.alignment = PP_ALIGN.CENTER
        para.space_before = Pt(4)
        run = para.add_run()
        run.text = detail
        template.style(run, size=12, color=SKY, italic=True)
    return tile


# The two ends of the Tufts "AI Job Risk Index" occupation ranking: the
# three most vulnerable of 784 occupations analyzed, and the three least.
# Percentages are predicted job loss; the whole tail of the ranking sits
# at 0%, so the bottom three are the ranking's own final three.
_MOST_VULNERABLE = [
    (1, "Writers and Authors", "57%"),
    (2, "Computer Programmers", "55%"),
    (3, "Web and Digital Interface Designers", "55%"),
]

_LEAST_VULNERABLE = [
    (782, "Foundry Mold and Coremakers", "0%"),
    (783, "Cleaners of Vehicles and Equipment", "0%"),
    (784, "Dishwashers", "0%"),
]


def _rank_panel(slide, left, width, heading, rows, edge, value_color):
    """Ranked-occupation card: heading, rule, then rank / name / percent."""
    top = 3.75
    height = 2.45
    card = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, left, Inches(top), width, Inches(height)
    )
    template.panel(card)
    template.add_bar(slide, left, Inches(top), Inches(0.09), Inches(height), edge)

    inner = Inches(left.inches + 0.34)
    template.add_label(slide, inner, Inches(top + 0.16), width, heading,
                       size=15, color=WHITE)
    template.add_bar(slide, inner, Inches(top + 0.62),
                     Inches(width.inches - 0.68), Pt(1), PANEL_EDGE)

    name_left = Inches(left.inches + 0.86)
    name_width = Inches(width.inches - 0.86 - 1.25)
    value_left = Inches(left.inches + width.inches - 1.3)
    row_top = top + 0.76
    for rank, name, value in rows:
        template.add_label(slide, inner, Inches(row_top), Inches(0.5),
                           str(rank), size=13, color=MUTED)
        template.add_label(slide, name_left, Inches(row_top - 0.02), name_width,
                           name, size=14, color=TEXT)
        template.add_label(slide, value_left, Inches(row_top - 0.03),
                           Inches(0.95), value, size=17, color=value_color,
                           align=PP_ALIGN.RIGHT)
        row_top += 0.52
    return card


def _headline_figures_slide(prs):
    slide = template.add_content_slide(
        prs,
        "AI and the geography of job risk",
        ["What new research says about where AI lands in the labor market:"],
    )
    _stat_tile(
        slide, template.CONTENT_LEFT, Inches(5.96),
        "9.3 million",
        "US jobs at risk under the median AI-adoption scenario",
        "range across scenarios: 2.7–19.5 million",
    )
    _stat_tile(
        slide, Inches(6.77), Inches(5.96),
        "$757 billion",
        "in annual income attached to those jobs",
    )
    _rank_panel(
        slide, template.CONTENT_LEFT, Inches(5.96),
        "Most vulnerable occupations", _MOST_VULNERABLE,
        HIGHLIGHT, AMBER,
    )
    _rank_panel(
        slide, Inches(6.77), Inches(5.96),
        "Least vulnerable occupations", _LEAST_VULNERABLE,
        ACCENT, SKY,
    )
    template.add_label(
        slide, template.CONTENT_LEFT, Inches(6.35), template.CONTENT_WIDTH,
        TUFTS_CREDIT + " · AI Job Risk Index, share of jobs lost, "
        "ranked across 784 occupations",
        size=11,
    )
    template.set_notes(slide, (
        "~2.5 min. Open the session here — but set the tone in the first "
        "sentence: the point of these numbers is not fear. The Tufts Digital "
        "Planet research maps where AI lands in the US labor market, and it "
        "lands exactly on administrative and knowledge work — our kind of "
        "work. 9.3 million US jobs at risk under the median AI-adoption "
        "scenario (the honest range across scenarios is 2.7 to 19.5 "
        "million), representing 757 billion dollars in annual income.\n\n"
        "Then drop to the two panels — this is the part that lands. The "
        "three most vulnerable occupations of the 784 analyzed are Writers "
        "and Authors at 57%, Computer Programmers at 55%, and Web and "
        "Digital Interface Designers at 55%. Writing, coding, design: all "
        "desk work, all of it work we would have called creative or highly "
        "skilled.\n\n"
        "The other end of the ranking is the punchline. Foundry mould "
        "makers, vehicle cleaners, dishwashers — 0% predicted job loss. "
        "The whole tail sits at zero: 311 of the 784 occupations come out "
        "at 0%. If the work is physical, this wave does not touch it. If "
        "the work arrives as text on a screen, it does.\n\n"
        "Do not let anyone read this as 'programmers are doomed'. Read it "
        "as: exposure follows the task, not the job title. Hold that "
        "thought — it is the argument of the next three slides."
    ))


# The two Tufts charts ship as PNGs rather than redrawn shapes: they carry
# hundreds of plotted occupations, which no native-shape recreation would
# survive. The files live with the rest of the lesson's source material and
# are versioned alongside it.
_ASSETS = pathlib.Path(__file__).resolve().parents[2] / ".scratch"
_SECTOR_CHART = _ASSETS / "aijobloss.png"
_AUGMENTATION_CHART = _ASSETS / "aijobloss2.png"

_MAT_INSET = 0.09  # white margin the chart sits on, so the PNG never bleeds


def _chart_on_mat(slide, image, left, top, width, height):
    """A chart PNG on a white mat — the charts have white grounds, so the
    mat extends them into a deliberate panel instead of a ragged edge."""
    mat = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(left - _MAT_INSET), Inches(top - _MAT_INSET),
        Inches(width + 2 * _MAT_INSET), Inches(height + 2 * _MAT_INSET),
    )
    template.solid(mat, WHITE)
    return slide.shapes.add_picture(
        str(image), Inches(left), Inches(top), Inches(width), Inches(height),
    )


def _concentration_slide(prs):
    """The sector chart itself, full width — the figures live in the PNG."""
    slide = template.add_content_slide(
        prs, "Percentage of Jobs at Risk Aggregated by Sector", compact=True)

    width, height = 7.9, 5.09  # the chart's own 854×550 aspect ratio
    _chart_on_mat(slide, _SECTOR_CHART, (13.333 - width) / 2, 0.95, width, height)

    # The caption runs two lines at this width; the credit clears both.
    template.add_label(
        slide, template.CONTENT_LEFT, Inches(6.15), template.CONTENT_WIDTH,
        "High-skill, analytical, office-based occupations face the steepest "
        "exposure. The jobs historically “safe” from automation are now the "
        "most exposed.",
        size=15, color=TEXT,
    )
    template.add_label(slide, template.CONTENT_LEFT, Inches(6.64),
                       template.CONTENT_WIDTH, TUFTS_CREDIT, size=11)
    template.set_notes(slide, (
        "~2 min. Walk the bars top-down: Information 18%, Finance and "
        "Insurance 16%, Professional Services 16% — against a 6% "
        "all-industry average. Three times the average, and all three are "
        "knowledge-work sectors.\n\n"
        "Land the reversal: for decades, automation threatened manual and "
        "routine work while analytical office jobs were the 'safe' ones. "
        "This research says the opposite is now true — high-skill, "
        "analytical, office-based occupations face the steepest exposure. "
        "Research management sits in exactly this category; say so, but "
        "save the full framing for the thesis slide in a minute."
    ))


_CHART_READING = [
    "Across: how much AI augments the job's tasks.",
    "Up: predicted job loss for that job.",
    "The cloud climbs a diagonal — the more augmented, the more displaced.",
    "The flat blue band along the bottom is physical work: neither.",
]


def _reading_panel(slide, left, top, width, height, heading, lines):
    """A small how-to-read-it panel — the step between chart and claim."""
    box = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(left), Inches(top), Inches(width),
        Inches(height),
    )
    template.panel(box)
    template.add_bar(slide, Inches(left), Inches(top), Inches(0.09),
                     Inches(height), ACCENT)
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.28)
    tf.margin_right = Inches(0.2)
    tf.margin_top = Inches(0.14)
    tf.vertical_anchor = MSO_ANCHOR.TOP
    para = tf.paragraphs[0]
    para.alignment = PP_ALIGN.LEFT
    run = para.add_run()
    run.text = heading
    template.style(run, size=14, color=WHITE, font=HEADING_FONT, bold=True)
    for line in lines:
        para = tf.add_paragraph()
        para.space_before = Pt(5)
        run = para.add_run()
        run.text = line
        template.style(run, size=12, color=MUTED)
    return box


def _pipeline_slide(prs):
    """The core finding and the researchers' own scatter on one slide: the
    claim on the left, the measurement beside it, the quote under both."""
    slide = template.add_content_slide(
        prs, "The augmentation–displacement pipeline", compact=True)

    template.add_label(
        slide, template.CONTENT_LEFT, Inches(0.86), template.CONTENT_WIDTH,
        "The research's counterintuitive core finding: help and risk are the "
        "same list — and here is the measurement that says so.",
        size=15, color=SKY,
    )

    # Right: the evidence. 949×750 on a white mat, sized to the column.
    chart_width = 5.69
    chart_height = chart_width / 1.2653  # the chart's own aspect ratio
    _chart_on_mat(slide, _AUGMENTATION_CHART, 12.63 - chart_width, 1.28,
                  chart_width, chart_height)

    # Left: the claim, read top-down — augmentation leads to displacement.
    column = 6.05
    template.add_card(
        slide, template.CONTENT_LEFT, Inches(1.28), Inches(column), Inches(1.05),
        PALE, "Highest augmentation potential", PRIMARY,
        "The occupations where AI helps the most with everyday tasks…",
    )
    template.add_card(
        slide, template.CONTENT_LEFT, Inches(2.72), Inches(column), Inches(1.05),
        PALE_WARM, "Highest displacement risk", PRIMARY,
        "…are the same occupations AI can most readily replace.",
    )
    template.add_label(slide, Inches(0.6 + column / 2 - 0.4), Inches(2.33),
                       Inches(0.8), "↓", size=20, color=AMBER,
                       align=PP_ALIGN.CENTER)
    _reading_panel(slide, 0.6, 3.97, column, 1.81,
                   "Reading the chart", _CHART_READING)

    template.add_band(
        slide, Inches(5.95), Inches(0.78),
        "“The more AI helps you do your job, the more expendable you become.”",
        size=18, sub="— the researchers' own framing",
    )
    template.add_label(slide, template.CONTENT_LEFT, Inches(6.82),
                       template.CONTENT_WIDTH, TUFTS_CREDIT, size=11)

    template.set_notes(slide, (
        "~3 min. This is the slide the whole opening builds to, and it is "
        "built to be taught in three moves — claim, evidence, consequence. "
        "Do not read it all at once.\n\n"
        "1. THE CLAIM (left, top-down). The naive expectation is that "
        "AI-helps and AI-replaces are two different lists of jobs. The "
        "research found they are one list: the occupations with the highest "
        "augmentation potential also carry the highest displacement risk. "
        "Say it before showing why it is true.\n\n"
        "2. THE EVIDENCE (right). Now turn to the chart and show that this "
        "is measured, not asserted. Give the axes first — across is how much "
        "AI augments the job's tasks, up is predicted job loss — then let "
        "the shape do the work: if the two were unrelated, the cloud would "
        "be shapeless; instead it climbs a diagonal. Point at two named dots "
        "and let the room read them aloud: Writers and Authors, Computer "
        "Programmers, top right — most augmented and most displaced. Then "
        "the flat blue band along the bottom: physical work, barely "
        "augmented and barely displaced. Ninety seconds, no more.\n\n"
        "3. THE CONSEQUENCE (the quote). Read it out loud, then pause and "
        "let it sit — this is the uncomfortable moment of the opening and "
        "it should feel that way. Then plant the question the next slide "
        "answers: if being helped by AI makes you expendable, what makes "
        "the difference between at-risk and empowered? Do not answer it "
        "here — the turn is the whole point of the course.\n\n"
        "If the room is quiet, ask where they would put their own job on "
        "the chart before you move on."
    ))


def _turn_slide(prs):
    slide = template.add_content_slide(
        prs,
        "The turn: skill, not job title",
        ["The point is not fear. The same research implies the way through — "
         "what matters is how you work with AI, not what your job is called."],
    )
    template.add_card(
        slide, template.CONTENT_LEFT, Inches(2.5), Inches(5.96), Inches(1.9),
        PALE_WARM, "Displacement attaches to shallow use", PRIMARY,
        "Letting AI do the job for you — unchecked, unexamined, at a "
        "shallow level.",
    )
    template.add_card(
        slide, Inches(6.77), Inches(2.5), Inches(5.96), Inches(1.9),
        PALE, "Value attaches to moving up the skill curve", PRIMARY,
        "Directing AI, verifying it, and building systems with it. This "
        "course exists to move you up that curve.",
    )
    template.add_band(
        slide, Inches(4.85), Inches(1.35),
        "Research management is administrative, analytical knowledge work — "
        "squarely in the exposed zone.\nThat is precisely why learning this "
        "well matters for you, now.",
        size=18,
    )
    template.set_notes(slide, (
        "~3 min, including a beat for questions. Answer the question the "
        "quote raised: displacement risk attaches to people who let AI do "
        "their job for them at a shallow level. People who move up the "
        "skill curve — who direct AI, verify it, and build systems with it "
        "— become more valuable, not less. That is the course thesis, and "
        "the reason this course exists.\n\n"
        "Then make it personal for this audience: research management is "
        "administrative, analytical knowledge work — squarely in the "
        "exposed zone, which is precisely why learning this well matters "
        "now. Keep the register matter-of-fact, not alarmist: the numbers "
        "were the diagnosis, the course is the response.\n\n"
        "Bridge to Segment 2: 'so let's start by understanding the tool "
        "itself — what these models actually are, and why they behave so "
        "strangely.'"
    ))


def add_slides(prs):
    _headline_figures_slide(prs)
    _concentration_slide(prs)
    _pipeline_slide(prs)
    _turn_slide(prs)
