"""Structural tests for the Lesson 2 deck skeleton (ticket 03 of the
Lesson 2 effort).

Every assertion runs against a saved-and-reopened .pptx, so it checks what
PowerPoint will load. The Timing Overview in lesson_2.md is the oracle for
segment names, durations and modes, as lesson_1.md is for the Lesson 1
deck.
"""

import os
import re
import tempfile
from pathlib import Path

import pytest
from pptx import Presentation

from deck.lesson2.build import build_presentation


def _slide_text(slide):
    """All visible text on a slide, joined."""
    return "\n".join(
        shape.text_frame.text for shape in slide.shapes if shape.has_text_frame
    )


def _slide_corpus(slide):
    """Everything a slide says: its visible text and its speaker notes."""
    text = _slide_text(slide)
    if slide.has_notes_slide:
        text += "\n" + slide.notes_slide.notes_text_frame.text
    return text


@pytest.fixture(scope="module")
def reopened():
    """Build the deck, save it, and reopen it the way PowerPoint would."""
    prs = build_presentation()
    fd, path = tempfile.mkstemp(suffix=".pptx")
    os.close(fd)
    try:
        prs.save(path)
        yield Presentation(path)
    finally:
        os.unlink(path)


# --- Title slide ------------------------------------------------------------


def test_title_slide_names_the_lesson_course_week_and_length(reopened):
    text = _slide_text(reopened.slides[0])
    assert "Lesson 2 — The Two-Lane Tool Landscape" in text
    assert "ai for research management" in text.lower()
    assert "Week 2" in text
    assert "60 minutes" in text


# --- Section dividers ---------------------------------------------------------

# The Timing Overview table from lesson_2.md — the independent source of truth.
TIMING_OVERVIEW = [
    ("Recap + assignment debrief", "15 min", "Participant demos, facilitated"),
    ("The sensitive-data rule and data-protection principles", "10 min", "Presentation"),
    ("Copilot Chat (Basic) in practice", "12 min", "Live facilitator demo"),
    ("Gemini in depth (+ NotebookLM preview)", "15 min", "Live facilitator demo"),
    ("Assignment briefing + close", "8 min", "Presentation + Q&A"),
]


def test_each_segment_has_a_divider_with_duration_and_mode(reopened):
    all_slides = [_slide_text(s) for s in reopened.slides]
    for number, (name, duration, mode) in enumerate(TIMING_OVERVIEW, start=1):
        dividers = [t for t in all_slides if name in t and f"{duration} · {mode}" in t]
        assert dividers, f"no divider found for segment {number} ({name})"
        assert any(f"Segment {number} of 5" in t for t in dividers)


def test_dividers_appear_in_session_order_after_the_title(reopened):
    all_slides = [_slide_text(s) for s in reopened.slides]
    positions = [
        next(i for i, t in enumerate(all_slides) if name in t)
        for name, _, _ in TIMING_OVERVIEW
    ]
    assert positions == sorted(positions)
    assert positions[0] > 0


# --- Segments as slide ranges ------------------------------------------------

# The five segments, in the order their dividers appear. Tickets 04–08 have
# all landed, so no segment carries a placeholder any more — the integration
# section below is where that is asserted deck-wide.
SEGMENT_NAMES = [name for name, _, _ in TIMING_OVERVIEW]


def _segment_ranges(reopened):
    """Map divider name -> the slides between that divider and the next."""
    all_slides = [_slide_text(s) for s in reopened.slides]
    starts = {
        name: next(i for i, t in enumerate(all_slides) if name in t)
        for name in SEGMENT_NAMES
    }
    ordered = sorted(starts.items(), key=lambda kv: kv[1])
    ranges = {}
    for (name, start), nxt in zip(ordered, [s for _, s in ordered[1:]] + [len(all_slides)]):
        ranges[name] = list(reopened.slides)[start + 1 : nxt]
    return ranges


def _dividers(reopened):
    """The section dividers, found where they are rather than by what they
    say: the slide immediately before each segment's first content slide."""
    all_slides = list(reopened.slides)
    return [
        all_slides[all_slides.index(slides[0]) - 1]
        for slides in _segment_ranges(reopened).values()
    ]


def test_total_slide_count(reopened):
    # The title slide, a divider per segment, and the per-segment counts
    # pinned by EXPECTED_SEGMENT_SLIDE_COUNTS in the integration section.
    expected = (
        1 + len(TIMING_OVERVIEW) + sum(EXPECTED_SEGMENT_SLIDE_COUNTS.values())
    )
    assert len(reopened.slides) == expected


# --- Segment 1: Recap + assignment debrief (ticket 04) -----------------------

# The plan is the oracle. Unlike TIMING_OVERVIEW above (transcribed from
# lesson_2.md, so the copy is independent), the recap rows are read out of
# the plan's Part A table at test time: the ticket asks for a test that
# checks the pairs "against the plan", and the table's seven rows are too
# long to transcribe without the copy itself becoming the thing that
# drifts. A row added to or dropped from the plan fails the deck here.
PLAN_PATH = Path(__file__).resolve().parents[1] / "lesson_2.md"

RECAP_TITLE = "Lesson 1 → Lesson 2: the recap map"
DEBRIEF_TITLE = "Assignment debrief — the next 12 minutes"
CAPTURE_TITLE = "Common problems — captured live"
SEG1_TITLES = [RECAP_TITLE, DEBRIEF_TITLE, CAPTURE_TITLE]

# One (concept token, link token) pair per row of the plan's recap table,
# in the plan's order. Each token is asserted to appear in its plan cell
# first, so the pairs cannot drift from the plan; matching is
# case-sensitive, which is what keeps "Meta-prompting" (the concept) from
# matching "meta-prompting" inside its own link cell.
RECAP_TOKENS = [
    ("7 Levels", "Level 3"),
    ("ICC", "lane decides"),
    ("Ask-me-first", "standing instructions"),
    ("Meta-prompting", "packaged agent"),
    ("two rules", "two-lane rule"),
    ("Anchor Project", "anonymise-before-upload"),
    ("Context window", "curation"),
]


def _slide_by_title(reopened, title):
    """The slide whose text contains `title`; exactly one must."""
    matches = [s for s in reopened.slides if title in _slide_text(s)]
    assert len(matches) == 1, f"{title!r}: {len(matches)} slides matched"
    return matches[0]


def _shape_containing(slide, token):
    """The single shape whose text contains `token` (case-sensitive)."""
    matches = [
        sh for sh in slide.shapes
        if sh.has_text_frame and token in sh.text_frame.text
    ]
    assert len(matches) == 1, f"{token!r}: {len(matches)} shapes matched"
    return matches[0]


def _slide_title(slide):
    """A content slide's title — the first text box add_content_slide adds.
    A segment names its own subjects in more places than the title, so a
    title match has to be the title, not a mention anywhere on the slide."""
    for shape in slide.shapes:
        if shape.has_text_frame:
            return shape.text_frame.text
    return ""


def _segment_slide(reopened, segment, title):
    """The slide inside one segment whose title is exactly this."""
    matches = [
        s for s in _segment_ranges(reopened)[segment]
        if _slide_title(s) == title
    ]
    assert len(matches) == 1, f"{title!r}: {len(matches)} slides matched"
    return matches[0]


def _recap_rows_from_plan():
    """The (concept, leads-into) rows of the Part A table in lesson_2.md."""
    text = PLAN_PATH.read_text(encoding="utf-8")
    part_a = text[text.index("### Part A: The recap map"):text.index("### Part B")]
    rows = []
    for line in part_a.splitlines():
        line = line.strip()
        if not line.startswith("|") or set(line) <= set("|- "):
            continue
        cells = [c.strip() for c in line.strip("|").split("|")]
        if cells[0].startswith("Lesson 1 concept"):
            continue
        rows.append((cells[0], cells[1]))
    return rows


def _segment_refs(text):
    """The segment numbers a piece of text points forward to — "Segment 4",
    "Segments 3 · 4", "Segments 3 and 4"."""
    refs = set()
    for chunk in re.findall(r"Segments?\s+(\d+(?:\s*(?:·|,|and)\s*\d+)*)", text):
        refs.update(int(d) for d in re.findall(r"\d+", chunk))
    return refs


def test_seg1_slides_present_in_order_and_not_placeholders(reopened):
    """All three slides live between Segment 1's divider and Segment 2's,
    in the plan's order, with no placeholder left."""
    segment = _segment_ranges(reopened)["Recap + assignment debrief"]
    texts = [_slide_text(s) for s in segment]
    positions = []
    for title in SEG1_TITLES:
        matches = [i for i, t in enumerate(texts) if title in t]
        assert matches, f"{title!r} is not in Segment 1"
        assert "[PLACEHOLDER]" not in texts[matches[0]]
        positions.append(matches[0])
    assert positions == sorted(positions)


def test_recap_map_pairs_every_plan_concept_with_its_lesson_2_link(reopened):
    """Every row of the plan's Part A table is on the slide as a concept
    sitting to the left of the Lesson 2 topic it leads into, and the link
    keeps the plan's forward references."""
    slide = _slide_by_title(reopened, RECAP_TITLE)
    plan_rows = _recap_rows_from_plan()
    assert len(plan_rows) == len(RECAP_TOKENS), (
        f"plan has {len(plan_rows)} recap rows, tokens cover {len(RECAP_TOKENS)}"
    )

    tops = []
    for (concept_cell, link_cell), (concept, link) in zip(plan_rows, RECAP_TOKENS):
        assert concept in concept_cell, f"{concept!r} not in the plan's concept cell"
        assert link in link_cell, f"{link!r} not in the plan's leads-into cell"

        concept_shape = _shape_containing(slide, concept)
        link_shape = _shape_containing(slide, link)
        assert concept_shape.top == link_shape.top, f"{concept!r}: link not alongside"
        assert concept_shape.left < link_shape.left, f"{concept!r}: link not to its right"

        planned = _segment_refs(concept_cell) | _segment_refs(link_cell)
        assert planned <= _segment_refs(link_shape.text_frame.text), (
            f"{concept!r}: slide points to "
            f"{_segment_refs(link_shape.text_frame.text)}, plan to {planned}"
        )
        tops.append(concept_shape.top)

    assert tops == sorted(tops), "recap rows are not in the plan's order"


# Wording that would mean a Lesson 1 concept is being explained again
# rather than referenced — the recap map links, it never re-teaches.
DEFINITION_MARKERS = [
    " is a ", " is an ", " is the ", " are the ", " means ", " refers to ",
    " stands for ", "in other words", "that is, ", "which is ",
]


def test_recap_map_restates_no_lesson_1_definition(reopened):
    """No cell on the recap slide explains a concept: no definition
    wording and no multi-sentence prose — each cell is a link, not a
    lesson."""
    slide = _slide_by_title(reopened, RECAP_TITLE)
    for shape in slide.shapes:
        if not shape.has_text_frame:
            continue
        text = shape.text_frame.text
        if RECAP_TITLE in text:
            continue
        lowered = text.lower()
        for marker in DEFINITION_MARKERS:
            assert marker not in lowered, f"{marker!r} in {text!r}"
        assert not re.search(r"\.\s+[A-Z“]", text), f"explanation prose in {text!r}"


def test_debrief_slide_states_the_format_and_the_four_items(reopened):
    text = _slide_text(_slide_by_title(reopened, DEBRIEF_TITLE))
    assert "2–3 volunteers" in text
    assert "4 minutes" in text          # the per-volunteer time box
    assert "12 minutes" in text         # the segment's time box
    assert "Gemini" in text             # screen-shared from their own account

    # The four assignment items, in the plan's order.
    items = ["Anchor Project", "ICC prompt", "Ask-me-first", "meta-prompt"]
    positions = []
    for item in items:
        assert item.lower() in text.lower(), f"{item!r} missing from the debrief slide"
        positions.append(text.lower().index(item.lower()))
    assert positions == sorted(positions), "the four items are out of plan order"


def test_debrief_slide_carries_the_facilitator_questions(reopened):
    text = _slide_text(_slide_by_title(reopened, DEBRIEF_TITLE)).lower()
    assert "anonymise or invent" in text
    assert "five sentences" in text


def test_capture_slide_tags_the_four_items_and_routes_the_problems(reopened):
    """The capture slide is empty space the facilitator fills live, tagged
    by assignment item, with the routing the plan gives."""
    slide = _slide_by_title(reopened, CAPTURE_TITLE)
    text = _slide_text(slide)
    for tag in ("Anchor Project", "ICC", "Ask-me-first", "Meta-prompt"):
        assert tag in text, f"capture lane {tag!r} missing"
    assert _segment_refs(text) >= {2, 3, 4}, "the routing loses a segment"
    assert "Practice Week 1" in text

    # Blank lanes to write in: cards with no text of their own.
    blank = [
        sh for sh in slide.shapes
        if sh.has_text_frame and not sh.text_frame.text.strip()
    ]
    assert len(blank) >= 4, "no blank capture lanes to write in"


def test_every_seg1_slide_has_speaker_notes(reopened):
    for title in SEG1_TITLES:
        slide = _slide_by_title(reopened, title)
        assert slide.has_notes_slide, f"{title!r} has no notes slide"
        assert slide.notes_slide.notes_text_frame.text.strip(), (
            f"{title!r} has empty speaker notes"
        )


def test_seg1_notes_carry_the_time_keeping_script(reopened):
    """The facilitation script that keeps the demos on time lives in the
    notes: the 3-minute chat warning, the cut at 4, and where an
    unfinished assignment goes."""
    notes = " ".join(
        _slide_by_title(reopened, t).notes_slide.notes_text_frame.text
        for t in SEG1_TITLES
    ).lower()
    assert "3-minute" in notes or "three-minute" in notes
    assert "cut at 4" in notes
    assert "practice week 1" in notes
    assert "backup" in notes  # the spare volunteer if a screen-share fails


# --- Segment 2: The sensitive-data rule (ticket 05) --------------------------

ROUTING_TITLE = "The two-lane rule: one routing question"
FRAMEWORK_TITLE = "AI data safety and ethical use — the framework"
DATA_SAFETY_TITLE = "Key pillars of data safety"
ETHICS_TITLE = "Ethical AI principles"
SOURCES_TITLE = "Sources"
ANONYMISE_TITLE = "Anonymise before upload"
NEVER_TITLE = "What must never enter Gemini"
SEG2_TITLES = [
    ROUTING_TITLE, FRAMEWORK_TITLE, DATA_SAFETY_TITLE, ETHICS_TITLE,
    SOURCES_TITLE, ANONYMISE_TITLE, NEVER_TITLE,
]

CONTEXT_PATH = PLAN_PATH.parent / "CONTEXT.md"


def _plan_section(heading):
    """The body of one subsection of lesson_2.md, up to the next heading."""
    text = PLAN_PATH.read_text(encoding="utf-8")
    rest = text[text.index(heading) + len(heading):]
    end = rest.index("\n#") if "\n#" in rest else len(rest)
    return rest[:end]


def test_seg2_slides_present_in_plan_order_and_not_placeholders(reopened):
    """All seven slides live between Segment 2's divider and Segment 3's,
    in the plan's order, with no placeholder left."""
    segment = _segment_ranges(reopened)[
        "The sensitive-data rule and data-protection principles"]
    texts = [_slide_text(s) for s in segment]
    positions = []
    for title in SEG2_TITLES:
        matches = [i for i, t in enumerate(texts) if title in t]
        assert matches, f"{title!r} is not in Segment 2"
        assert "[PLACEHOLDER]" not in texts[matches[0]]
        positions.append(matches[0])
    assert positions == sorted(positions)


# The four kinds of work CONTEXT.md puts in the Sensitive Lane. Each is
# asserted against CONTEXT.md's own Sensitive Lane entry first, so the
# routing slide cannot drift from the ubiquitous language.
SENSITIVE_CATEGORIES = [
    "Internal administration", "Unsubmitted grant proposals",
    "Personnel information", "Strategic financial planning",
]


def test_routing_slide_asks_one_question_and_names_both_lanes(reopened):
    text = _slide_text(_slide_by_title(reopened, ROUTING_TITLE))
    assert "Does the task involve sensitive or IP-sensitive institutional data?" in text
    for lane in ("Sensitive Lane", "Daily Driver"):
        assert lane in text, f"{lane!r} missing from the routing slide"
    assert "Copilot Chat (Basic)" in text
    assert "Gemini" in text
    # Routed per task, not per person and not per day.
    assert "per task" in text.lower()


# What the plan's routing table puts in the Daily Driver column. CONTEXT.md
# is the oracle for the Sensitive Lane's four — they are its own list; the
# plan is the oracle for the Daily Driver's, which it words for the slide.
DAILY_CATEGORIES = [
    "Drafting", "Prompting exercises", "Long-document analysis",
    "General research management work",
]


def test_routing_slide_lists_what_the_daily_driver_covers(reopened):
    plan = _plan_section("### The two-lane rule as a routing decision").lower()
    text = _slide_text(_slide_by_title(reopened, ROUTING_TITLE)).lower()
    for category in DAILY_CATEGORIES:
        assert category.lower() in plan, (
            f"{category!r} is not in the plan's routing table"
        )
        assert category.lower() in text, f"{category!r} missing from the routing slide"


def test_routing_slide_lists_the_four_sensitive_categories(reopened):
    context = CONTEXT_PATH.read_text(encoding="utf-8")
    sensitive = context[
        context.index("**Sensitive Lane**"):context.index("**Daily Driver**")
    ].lower()
    text = _slide_text(_slide_by_title(reopened, ROUTING_TITLE)).lower()
    for category in SENSITIVE_CATEGORIES:
        assert category.lower() in sensitive, (
            f"{category!r} is not a CONTEXT.md Sensitive Lane category"
        )
        assert category.lower() in text, f"{category!r} missing from the routing slide"


def test_routing_slide_carries_the_three_worked_examples(reopened):
    """The plan's three routing examples, each landing in the lane the plan
    routes it to."""
    plan = _plan_section("### The two-lane rule as a routing decision")
    slide = _slide_by_title(reopened, ROUTING_TITLE)
    for token, lane in (("Cluster 4", "Sensitive Lane"),
                        ("Horizon Europe", "Daily Driver"),
                        ("workload plan", "Sensitive Lane")):
        assert token in plan, f"{token!r} is not one of the plan's examples"
        shape = _shape_containing(slide, token)
        assert lane in shape.text_frame.text, f"{token!r} is not routed to {lane}"


def test_routing_slide_colour_codes_both_lanes(reopened):
    from deck.lesson2 import template

    fills = _fills(_slide_by_title(reopened, ROUTING_TITLE))
    assert str(template.SENSITIVE_LANE.color) in fills
    assert str(template.DAILY_DRIVER.color) in fills


# The four pillars of the plan's framework table. Each is asserted against
# the plan first, so the slide cannot drift from what the plan authors.
FRAMEWORK_PILLARS = [
    "Data governance", "Algorithmic fairness",
    "Transparency", "Human control",
]


def test_framework_slide_draws_the_plans_four_pillars(reopened):
    plan = _plan_section("### AI data safety and ethical use: the framework")
    slide = _slide_by_title(reopened, FRAMEWORK_TITLE)
    text = _slide_text(slide)
    assert "PILLARS OF TRUSTWORTHY AI" in text, "the pillars block has no heading"

    lefts = []
    for pillar in FRAMEWORK_PILLARS:
        assert pillar in plan, f"{pillar!r} is not a pillar in the plan"
        shape = _shape_containing(slide, pillar)
        lefts.append(shape.left)
    # Drawn as one row, in the plan's order — the block, not a list.
    assert lefts == sorted(lefts), "the pillars are not laid out left to right"
    assert len(set(lefts)) == len(lefts), "two pillars share a column"

    lowered = text.lower()
    assert "sensitive information" in lowered
    assert "algorithmic bias" in lowered
    assert "human control" in lowered


def test_framework_slide_keeps_the_directives_as_a_reference_strip(reopened):
    """The two institutional documents are cited on the framework slide —
    titles and the directive's headline rules, never the document."""
    text = _slide_text(_slide_by_title(reopened, FRAMEWORK_TITLE))
    assert "1/2026 (III.26.)" in text, "the dean's directive is not cited"
    assert "4/2025. (X. 28.)" in text, "the rector-level parent is not named"
    lowered = text.lower()
    assert "rector" in lowered and "dean" in lowered
    # The rules the recap map promises this segment delivers.
    assert "human oversight" in lowered
    assert "explicit permission" in lowered
    assert "critical checking of every output" in lowered


# The section references the plan attaches to the directive's key points.
# They belong to the facilitator, not to the slide.
_KEY_POINT_SECTIONS = ["2", "2, 4, 5", "2, 4", "6", "4, 6", "7"]


def test_directive_section_references_live_in_the_notes_not_on_the_slide(reopened):
    """The slide points at the documents; the facilitator holds the page
    numbers."""
    slide = _slide_by_title(reopened, FRAMEWORK_TITLE)
    assert "section" not in _slide_text(slide).lower()
    notes = slide.notes_slide.notes_text_frame.text.lower()
    for reference in _KEY_POINT_SECTIONS:
        word = "sections" if "," in reference else "section"
        assert f"{word} {reference}" in notes, (
            f"the notes do not place a key point at {word} {reference}"
        )


# The plan's four data-safety practices and four ethical principles, as the
# lead of each card. Both slides take the same 2×2 shape.
DATA_SAFETY_PILLARS = [
    "Data minimisation", "Privacy by design", "Anonymisation",
    "Regulatory compliance",
]
ETHICAL_PRINCIPLES = [
    "Fairness and non-discrimination", "Transparency", "Accountability",
    "Human oversight",
]


@pytest.mark.parametrize("title, heading, leads", [
    (DATA_SAFETY_TITLE, "### Key pillars of data safety", DATA_SAFETY_PILLARS),
    (ETHICS_TITLE, "### Ethical AI principles", ETHICAL_PRINCIPLES),
])
def test_quad_slide_carries_the_plans_four_items(reopened, title, heading, leads):
    plan = _plan_section(heading)
    slide = _slide_by_title(reopened, title)
    seats = []
    for lead in leads:
        assert lead in plan, f"{lead!r} is not in the plan's {heading!r}"
        shape = _shape_containing(slide, lead)
        seats.append((shape.top, shape.left))
    assert len(set(seats)) == len(seats), "two cards share a position"
    assert seats == sorted(seats), "the cards are out of the plan's order"


def test_data_safety_slide_keeps_the_lesson_1_and_gdpr_anchors(reopened):
    text = _slide_text(_slide_by_title(reopened, DATA_SAFETY_TITLE))
    lowered = text.lower()
    assert "gdpr" in lowered
    assert "eu ai act" in lowered
    assert "personally identifiable information" in lowered
    # The Lesson 1 tie-back: minimisation and curation want the same thing.
    assert "more context is not better context" in lowered


def test_ethics_slide_closes_on_human_responsibility(reopened):
    text = _slide_text(_slide_by_title(reopened, ETHICS_TITLE))
    assert "the person signs" in text.lower()
    assert "critical checking of every output" in text.lower()


# Every bracketed source marker used on the three content slides has to
# resolve to a numbered entry on the sources slide.
def _cited_numbers(slide):
    numbers = set()
    for chunk in re.findall(r"\[([\d,\s–-]+)\]", _slide_text(slide)):
        parts = re.findall(r"(\d+)\s*[–-]\s*(\d+)|(\d+)", chunk)
        for low, high, single in parts:
            if single:
                numbers.add(int(single))
            else:
                numbers.update(range(int(low), int(high) + 1))
    return numbers


def test_every_source_marker_resolves_to_the_sources_slide(reopened):
    """The bracketed numbers on the framework, data-safety and ethics
    slides point at entries that exist — a claim can always be traced."""
    sources = _slide_by_title(reopened, SOURCES_TITLE)
    listed = _cited_numbers(sources)
    assert listed == set(range(1, 10)), f"the sources slide lists {sorted(listed)}"

    cited = set()
    for title in (FRAMEWORK_TITLE, DATA_SAFETY_TITLE, ETHICS_TITLE):
        cited |= _cited_numbers(_slide_by_title(reopened, title))
    assert cited, "no slide cites a source"
    assert cited <= listed, f"unresolved markers: {sorted(cited - listed)}"


def test_sources_slide_matches_the_plans_reading_list(reopened):
    """One row per source in the plan's list, each carrying a live link to
    the address the plan gives."""
    plan = _plan_section("### Sources")
    planned = re.findall(r"^\d+\. (.+?) — (https?://\S+)$", plan, re.MULTILINE)
    assert len(planned) == 9, f"the plan lists {len(planned)} sources"

    slide = _slide_by_title(reopened, SOURCES_TITLE)
    addresses = {
        run.hyperlink.address
        for shape in slide.shapes if shape.has_text_frame
        for para in shape.text_frame.paragraphs for run in para.runs
        if run.hyperlink.address
    }
    for _, url in planned:
        assert url in addresses, f"{url} is not linked from the sources slide"


def test_sources_slide_also_names_the_institutional_documents(reopened):
    text = _slide_text(_slide_by_title(reopened, SOURCES_TITLE))
    assert "4/2025. (X. 28.)" in text and "1/2026 (III.26.)" in text


def test_anonymise_slide_binds_every_upload_in_either_lane(reopened):
    text = _slide_text(_slide_by_title(reopened, ANONYMISE_TITLE))
    lowered = text.lower()
    assert "anonymised before it is uploaded" in lowered
    assert "any ai tool" in lowered
    assert "either lane" in lowered
    # Both lanes are named, so "either lane" is not an abstraction.
    assert "Sensitive Lane" in text and "Daily Driver" in text
    # The Lesson 1 rule it generalises, as a reference not a re-teach.
    assert "Anchor Project" in text


# One token per item of the plan's "what must never enter Gemini" list.
NEVER_TOKENS = [
    "Personal data of any kind",
    "Unsubmitted grant proposals",
    "Internal financial planning",
    "Unpublished research results",
    "industrial or contract research",
    "confidentiality or data-processing agreement",
    "not anonymised",
]


def test_never_slide_lists_the_plans_items_and_routes_them_onward(reopened):
    plan = _plan_section("### What must never enter Gemini")
    slide = _slide_by_title(reopened, NEVER_TITLE)
    text = _slide_text(slide)
    tops = []
    for token in NEVER_TOKENS:
        assert token in plan, f"{token!r} is not on the plan's list"
        tops.append(_shape_containing(slide, token).top)
    assert tops == sorted(tops), "the list is out of the plan's order"

    # Where the list sends the work, and the floor under it.
    assert "Sensitive Lane" in text
    assert "Copilot Chat (Basic)" in text
    assert "does not go into an ai tool at all" in text.lower()


def test_never_slide_is_not_coloured_as_a_gemini_recommendation(reopened):
    """Violet means "Gemini is the right tool" everywhere in this deck, so
    the list of what must never go to Gemini carries none of it. The only
    lane colour on the slide is the teal of the lane it routes to."""
    from deck.lesson2 import template

    fills = _fills(_slide_by_title(reopened, NEVER_TITLE))
    assert str(template.DAILY_DRIVER.color) not in fills
    assert str(template.SENSITIVE_LANE.color) in fills


def test_every_seg2_slide_has_speaker_notes(reopened):
    for title in SEG2_TITLES:
        slide = _slide_by_title(reopened, title)
        assert slide.has_notes_slide, f"{title!r} has no notes slide"
        assert slide.notes_slide.notes_text_frame.text.strip(), (
            f"{title!r} has empty speaker notes"
        )


def test_seg2_notes_supersede_the_lesson_1_interim_rule(reopened):
    """The tie-back to Lesson 1's interim rule is a note to the
    facilitator, not a slide that teaches the old rule again."""
    notes = " ".join(
        _slide_by_title(reopened, t).notes_slide.notes_text_frame.text
        for t in SEG2_TITLES
    ).lower()
    assert "nothing you couldn't publish" in notes
    assert "lesson 1" in notes
    for title in SEG2_TITLES:
        on_slide = _slide_text(_slide_by_title(reopened, title)).lower()
        assert "couldn't publish" not in on_slide, f"{title!r} re-teaches the old rule"


# --- Segment 3: Copilot Chat (Basic) in practice (ticket 06) -----------------

SEG3_NAME = "Copilot Chat (Basic) in practice"

# The compiled list of agents and apps available under the institutional
# licence — the oracle for which tools this segment must introduce.
TOOLS_PATH = PLAN_PATH.parent / ".scratch" / "copilot_free_to_use_tools.md"

SCENARIOS_TITLE = "Four scenarios in the Sensitive Lane"
LIMITS_TITLE = "Limits and guardrails"
ACROBAT_TITLE = "Adobe Acrobat in Teams and Copilot"
COACH_TITLE = "Prompt Coach — your own prompt, critiqued"
STUDY_TITLE = "Study and Learn — concepts, flashcards, quizzes"
CANVA_TITLE = "Canva in Teams — prototype the presentation"
GENSPARK_TITLE = "genspark.ai — build the actual presentation"
SEG3_TITLES = [
    SCENARIOS_TITLE, LIMITS_TITLE, ACROBAT_TITLE, COACH_TITLE, STUDY_TITLE,
    CANVA_TITLE, GENSPARK_TITLE,
]


def _seg3_demo_script():
    """The plan's Segment 3 demo script — which step runs live."""
    text = PLAN_PATH.read_text(encoding="utf-8")
    start = text.index("**Demo script (which use case runs live):**")
    return text[start:text.index("**Rule stated while demoing:**", start)]


def _segment_corpus(reopened, name):
    """One segment's divider and slides, each as slide text plus its
    notes — everything that segment says, in one list."""
    slides = _segment_ranges(reopened)[name]
    all_slides = list(reopened.slides)
    divider = all_slides[all_slides.index(slides[0]) - 1]
    return [_slide_corpus(slide) for slide in [divider, *slides]]


def test_seg3_slides_present_in_plan_order_and_not_placeholders(reopened):
    """All seven slides live between Segment 3's divider and Segment 4's, in
    the plan's order, with no placeholder left."""
    segment = _segment_ranges(reopened)[SEG3_NAME]
    assert len(segment) == len(SEG3_TITLES)
    texts = [_slide_text(s) for s in segment]
    positions = []
    for title in SEG3_TITLES:
        matches = [i for i, t in enumerate(texts) if title in t]
        assert matches, f"{title!r} is not in Segment 3"
        assert "[PLACEHOLDER]" not in texts[matches[0]]
        positions.append(matches[0])
    assert positions == sorted(positions)


# The plan's four Sensitive Lane scenarios: the category, and a token from
# the one-line prompt the plan gives it.
SEG3_SCENARIOS = [
    ("Internal administration", "decision log and action list"),
    ("Unsubmitted proposals", "evaluation criteria"),
    ("Personnel matters", "workload-distribution note"),
    ("Strategic financial planning", "cost projection table"),
]


def test_scenarios_slide_carries_the_plans_four_scenarios_in_order(reopened):
    plan = _plan_section("### Sensitive Lane use cases")
    slide = _slide_by_title(reopened, SCENARIOS_TITLE)
    tops = []
    for category, prompt in SEG3_SCENARIOS:
        assert category in plan, f"{category!r} is not one of the plan's scenarios"
        assert prompt in plan, f"{prompt!r} is not in the plan's prompt for {category}"
        shape = _shape_containing(slide, category)
        assert prompt in shape.text_frame.text, (
            f"{category!r} is not shown with its prompt"
        )
        tops.append(shape.top)
    assert tops == sorted(tops), "the scenarios are out of the plan's order"


def test_scenarios_slide_names_the_icc_prompt_shape(reopened):
    """Each scenario is a routing example plus the prompt shape from
    Lesson 1 — named, never re-taught."""
    text = _slide_text(_slide_by_title(reopened, SCENARIOS_TITLE))
    assert "ICC" in text
    assert "Instruction · Context · Constraints" in text
    assert "Lesson 1" in text


def test_limits_slide_states_the_no_grounding_limit_and_the_rationale(reopened):
    plan = _plan_section("### Limits and guardrails")
    text = _slide_text(_slide_by_title(reopened, LIMITS_TITLE))
    lowered = text.lower()
    for token in ("SharePoint", "OneDrive"):
        assert token in plan, f"{token!r} is not in the plan's limits"
        assert token in text, f"{token!r} missing from the limits slide"
    # The limit itself: it does not read the institution's files by itself.
    assert "does not read the institution" in lowered
    assert "paste or attach" in lowered
    # The enterprise behaviour that makes this the Sensitive Lane.
    assert "tenant" in lowered
    assert "not used to train" in lowered
    # And the rationale for reserving it.
    assert "reserve it for the sensitive work" in lowered


# One slide per tool in the compiled list, each introduced through
# something a research manager actually does — the token is asserted
# against the plan's tool table first.
COMPILED_TOOLS = [
    ("Adobe Acrobat", ACROBAT_TITLE, "Teams channel"),
    ("Prompt Coach", COACH_TITLE, "Lesson 1 assignment"),
    ("Study and Learn", STUDY_TITLE, "eligible costs"),
    ("Canva", CANVA_TITLE, "multi-page"),
    ("genspark.ai", GENSPARK_TITLE, "presentation"),
]


def test_every_tool_in_the_compiled_list_has_its_own_scenario_slide(reopened):
    tools = TOOLS_PATH.read_text(encoding="utf-8")
    plan = _plan_section("### One use case per tool in the compiled tool list")
    segment = _segment_ranges(reopened)[SEG3_NAME]
    for tool, title, scenario_token in COMPILED_TOOLS:
        assert tool.lower() in tools.lower(), (
            f"{tool!r} is not in the compiled tool list"
        )
        assert scenario_token.lower() in plan.lower(), (
            f"{scenario_token!r} is not in the plan's table for {tool}"
        )
        slide = _slide_by_title(reopened, title)
        assert slide in segment, f"{tool}: the slide is not inside Segment 3"
        text = _slide_text(slide)
        assert tool in text, f"{tool!r} is not named on its own slide"
        assert scenario_token.lower() in text.lower(), (
            f"{tool}: the slide does not say what the tool is for"
        )


# Acrobat's five free-to-use features, named as the compiled list names
# them, and the paid capabilities that must be flagged as out of reach.
ACROBAT_FREE = [
    "High-quality viewing", "Central access to PDFs", "Streamline PDF review",
    "Keep reviewers informed", "Collaborate with confidence",
]
ACROBAT_PAID = [
    "Convert PDFs to Word", "combine multiple documents", "add a password",
    "attributions", "meeting transcript",
]


def test_acrobat_slide_shows_the_free_features(reopened):
    tools = TOOLS_PATH.read_text(encoding="utf-8").lower()
    text = _slide_text(_slide_by_title(reopened, ACROBAT_TITLE)).lower()
    assert "free to use" in text
    for feature in ACROBAT_FREE:
        assert feature.lower() in tools, (
            f"{feature!r} is not a free-to-use feature in the compiled list"
        )
        assert feature.lower() in text, f"{feature!r} missing from the Acrobat slide"


def test_acrobat_premium_and_ai_features_are_marked_not_available(reopened):
    """Every paid capability appears only inside a block that says it is
    not available under the institutional licence."""
    tools = TOOLS_PATH.read_text(encoding="utf-8").lower()
    slide = _slide_by_title(reopened, ACROBAT_TITLE)
    assert "AI Assistant" in _slide_text(slide)
    for capability in ACROBAT_PAID:
        assert capability.lower() in tools, (
            f"{capability!r} is not a paid capability in the compiled list"
        )
        shape = _shape_containing(slide, capability)
        assert "not available" in shape.text_frame.text.lower(), (
            f"{capability!r} is not marked as not available"
        )


# The plan runs three of Segment 3 on screen: the first Sensitive Lane
# scenario, Prompt Coach, and Study and Learn.
SEG3_LIVE_TITLES = [SCENARIOS_TITLE, COACH_TITLE, STUDY_TITLE]


def test_live_demo_cues_mark_exactly_the_slides_the_plan_runs_live(reopened):
    script = _seg3_demo_script()
    for tool in ("Prompt Coach", "Study and Learn"):
        assert f"{tool} live" in script, f"the plan does not run {tool} live"
    assert "use case 1 live" in script, "the plan does not run scenario 1 live"

    for title in SEG3_TITLES:
        slide = _slide_by_title(reopened, title)
        cued = "LIVE DEMO" in _slide_text(slide)
        assert cued == (title in SEG3_LIVE_TITLES), (
            f"{title!r}: live-demo cue {'present' if cued else 'missing'}"
        )
        if cued:
            notes = slide.notes_slide.notes_text_frame.text.lower()
            assert "live" in notes, f"{title!r}: the cue is not mirrored in the notes"


def test_seg3_notes_carry_the_demo_script(reopened):
    """The facilitator's running order for the segment lives in the notes,
    step for step with the plan's demo script."""
    script = _seg3_demo_script().lower()
    notes = " ".join(
        _slide_by_title(reopened, t).notes_slide.notes_text_frame.text
        for t in SEG3_TITLES
    ).lower()
    for token in ("enterprise notice", "meeting notes", "prompt coach",
                  "permission", "flashcards", "quiz", "slide only"):
        assert token in script, f"{token!r} is not in the plan's demo script"
        assert token in notes, f"{token!r} is not in the Segment 3 notes"
    # The rule the facilitator states while demoing.
    assert "anonymis" in notes and "check" in notes


# The tools whose slide is shown rather than demonstrated. The plan wants
# each of them to carry what the facilitator would type or do, not only
# prose about the tool.
SEG3_SLIDE_ONLY_TITLES = [ACROBAT_TITLE, CANVA_TITLE, GENSPARK_TITLE]

# The shortest a "What you type" / "What you do" line can be and still be
# something a participant could reproduce.
REPRODUCIBLE_MIN_CHARS = 60


def test_slide_only_tools_show_what_to_type_or_do(reopened):
    """Two tools are demonstrated; the rest are shown as slides, and the
    plan wants the prompt on those slides rather than a description of
    one."""
    plan = _plan_section("### One use case per tool in the compiled tool list")
    assert "the prompt on the slide" in plan, (
        "the plan no longer asks the slide-only tools to carry their prompt"
    )
    for title in [COACH_TITLE, *SEG3_SLIDE_ONLY_TITLES]:
        slide = _slide_by_title(reopened, title)
        blocks = [
            sh for sh in slide.shapes
            if sh.has_text_frame
            and sh.text_frame.text.startswith(("What you type", "What you do"))
        ]
        assert len(blocks) == 1, f"{title!r}: {len(blocks)} what-you-do blocks"
        body = blocks[0].text_frame.text
        assert len(body) >= REPRODUCIBLE_MIN_CHARS, (
            f"{title!r}: {body!r} is too thin to reproduce"
        )


def _fill_count(slide, color):
    """How many shapes on the slide are filled with `color`."""
    count = 0
    for shape in slide.shapes:
        try:
            if shape.fill.type is not None and shape.fill.fore_color.rgb == color:
                count += 1
        except (AttributeError, TypeError):
            continue
    return count


def test_tools_outside_the_tenant_carry_no_lane_colour_of_their_own(reopened):
    """Teal says "this is the Sensitive Lane's tool". Canva and genspark.ai
    sit outside the institutional tenant and say so on their own slides, so
    none of their content blocks may carry it — the only teal left is what
    add_content_slide paints on every lane slide (the chip and the accent
    lead under the title)."""
    from deck.lesson2 import template

    teal = template.SENSITIVE_LANE.color
    chrome = _fill_count(
        template.add_content_slide(
            template.new_presentation(), "baseline", lane=template.SENSITIVE_LANE,
        ),
        teal,
    )
    assert chrome, "add_content_slide paints no lane colour — the test is moot"

    for title in (CANVA_TITLE, GENSPARK_TITLE):
        assert _fill_count(_slide_by_title(reopened, title), teal) == chrome, (
            f"{title!r} colours a content block as the Sensitive Lane's tool"
        )
    # Control: a tool that does run inside the tenant does colour its blocks.
    assert _fill_count(_slide_by_title(reopened, STUDY_TITLE), teal) > chrome


def test_every_seg3_slide_has_speaker_notes(reopened):
    for title in SEG3_TITLES:
        slide = _slide_by_title(reopened, title)
        assert slide.has_notes_slide, f"{title!r} has no notes slide"
        assert slide.notes_slide.notes_text_frame.text.strip(), (
            f"{title!r} has empty speaker notes"
        )


# Lane separation — that this segment never names the other lane — is
# asserted for every lane-bound segment at once in the integration section
# (test_every_lane_declared_segment_stays_in_its_lane).


# --- Segment 4: Gemini in depth (+ NotebookLM preview) (ticket 07) -----------

SEG4_NAME = "Gemini in depth (+ NotebookLM preview)"

SEG4_OPENER_TITLE = "What the Daily Driver adds"
SEG4_GEMS_TITLE = "Gems — persistent context"
SEG4_CHAIN_TITLE = "Chaining what you have just seen"

# One (plan capability, slide title, scenario token) triple per row of the
# plan's capability table, in the plan's order. The token is asserted to
# appear in the plan's own use-case cell first, so a slide cannot drift
# from the scenario the plan gave it; the rows themselves are read out of
# the plan at test time, so a capability added there fails the deck here.
SEG4_CAPABILITIES = [
    ("File and image upload", "File and image upload",
     "anonymised call document"),
    ("Long context", "Long context", "structured summary"),
    ("Gems — persistent context", SEG4_GEMS_TITLE, "Anchor Project assistant"),
    ("Deep Research", "Deep Research", "funding landscape"),
    ("Canvas", "Canvas", "Draft the outline"),
    ("Google Workspace hooks", "Google Workspace hooks", "Drive"),
    ("Image generation and editing", "Image generation and editing",
     "diagram"),
    ("Gemini Live and mobile", "Gemini Live and mobile", "rehears"),
    ("Guided Learning mode", "Guided Learning mode", "concept"),
    ("Saved info / personalisation", "Saved info — the lightweight Gem",
     "lightweight version of a Gem"),
]

# The three NotebookLM scenarios: the plan's numbered heading, the slide
# title, and a token that has to survive from one to the other.
SEG4_NOTEBOOK_SCENARIOS = [
    ("YouTube link turned into a textual source",
     "NotebookLM: a YouTube link becomes a source", "captions"),
    ("Several documents aggregated in one notebook",
     "NotebookLM: several documents, one notebook", "FAQ"),
    ("Notebooks imported into Gemini as project folders",
     "NotebookLM: a notebook as the project folder", "sources travel with"),
]

SEG4_NOTEBOOK_TITLES = [title for _, title, _ in SEG4_NOTEBOOK_SCENARIOS]

SEG4_TITLES = (
    [SEG4_OPENER_TITLE]
    + [title for _, title, _ in SEG4_CAPABILITIES]
    + SEG4_NOTEBOOK_TITLES
    + [SEG4_CHAIN_TITLE]
)


def _seg4_slide(reopened, title):
    """The Segment 4 slide with exactly this title."""
    return _segment_slide(reopened, SEG4_NAME, title)


def _seg4_capability_rows():
    """The (capability, scenario, demo) rows of the plan's Segment 4
    capability table."""
    section = _plan_section(
        "### Free-tier capabilities beyond the Sensitive Lane tool")
    rows = []
    for line in section.splitlines():
        line = line.strip()
        if not line.startswith("|"):
            continue
        cells = [c.strip().replace("**", "") for c in line.strip("|").split("|")]
        if len(cells) != 3 or cells[0] in ("Capability",):
            continue
        if set(cells[0]) <= set("-: "):
            continue
        rows.append(tuple(cells))
    return rows


def _seg4_demo_script():
    """The plan's Segment 4 demo script — the facilitator's running order."""
    text = PLAN_PATH.read_text(encoding="utf-8")
    start = text.index("### Demo script\n\n1. Upload demo")
    return text[start:text.index("**Rule stated while demoing:**", start)]


def test_seg4_slides_present_in_plan_order_and_not_placeholders(reopened):
    """All fifteen slides live between Segment 4's divider and Segment 5's,
    in the plan's order — the opener, the capability table row by row, the
    three NotebookLM scenarios, the chaining close — with no placeholder
    left."""
    segment = _segment_ranges(reopened)[SEG4_NAME]
    assert len(segment) == len(SEG4_TITLES)
    titles = [_slide_title(s) for s in segment]
    assert titles == SEG4_TITLES
    for slide in segment:
        assert "[PLACEHOLDER]" not in _slide_text(slide)


def test_every_capability_in_the_plan_has_a_slide_with_its_scenario(reopened):
    """One slide per row of the plan's capability table, each carrying the
    scenario the plan gave that capability."""
    rows = _seg4_capability_rows()
    assert len(rows) == len(SEG4_CAPABILITIES), (
        f"plan lists {len(rows)} capabilities, the deck covers "
        f"{len(SEG4_CAPABILITIES)}"
    )
    for (plan_name, scenario_cell, _), (name, title, token) in zip(
            rows, SEG4_CAPABILITIES):
        assert plan_name == name, f"plan capability {plan_name!r} is not {name!r}"
        assert token.lower() in scenario_cell.lower(), (
            f"{token!r} is not in the plan's scenario for {name!r}"
        )
        text = _slide_text(_seg4_slide(reopened, title))
        assert token.lower() in text.lower(), (
            f"{title!r} does not carry the plan's scenario ({token!r})"
        )


def test_gems_slide_states_the_level_3_link(reopened):
    """The plan calls the Gem "the Level 3 move": a workspace that already
    knows the Anchor Project, built out of the Lesson 1 prompts."""
    text = _slide_text(_seg4_slide(reopened, SEG4_GEMS_TITLE))
    assert "Level 3" in text
    assert "Anchor Project" in text
    # The ICC prompt and the reverse prompt from Lesson 1, promoted into the
    # Gem's standing instructions.
    lowered = text.lower()
    assert "standing instructions" in lowered
    assert "icc prompt" in lowered
    assert "reverse prompt" in lowered
    assert "Lesson 1" in text


SEG4_LIVE_TITLES = [
    "File and image upload", "Long context", SEG4_GEMS_TITLE, "Canvas",
]


def test_seg4_live_demo_cues_mark_exactly_the_plans_live_capabilities(reopened):
    """The plan's Demo column decides which slides carry a cue: the four it
    marks **Live**, and no others."""
    live = [
        name for name, _, demo in _seg4_capability_rows()
        if demo.startswith("Live")
    ]
    titles = dict((name, title) for name, title, _ in SEG4_CAPABILITIES)
    assert [titles[name] for name in live] == SEG4_LIVE_TITLES

    for title in SEG4_TITLES:
        slide = _seg4_slide(reopened, title)
        cued = "LIVE DEMO" in _slide_text(slide)
        assert cued == (title in SEG4_LIVE_TITLES), (
            f"{title!r}: live-demo cue {'present' if cued else 'missing'}"
        )
        if cued:
            notes = slide.notes_slide.notes_text_frame.text.lower()
            assert "live" in notes, f"{title!r}: the cue is not mirrored in the notes"


def test_seg4_notes_carry_the_demo_script(reopened):
    """The facilitator's running order for the segment lives in the notes,
    step for step with the plan's demo script."""
    script = _seg4_demo_script().lower()
    notes = " ".join(
        _seg4_slide(reopened, t).notes_slide.notes_text_frame.text
        for t in SEG4_TITLES
    ).lower()
    for token in ("pdf", "icc prompt", "reference files", "first chat",
                  "one iteration", "notebooklm"):
        assert token in script, f"{token!r} is not in the plan's demo script"
        assert token in notes, f"{token!r} is not in the Segment 4 notes"
    # The rule the plan has the facilitator state while demoing.
    assert "anonymised" in notes and "second monitor" in notes


def test_seg4_notes_carry_the_reverification_reminder(reopened):
    """The plan's availability flag: the free tier moves, so every slide
    that promises a capability tells the facilitator to check it the week
    before, and the flag itself is on the opener slide."""
    opener = _slide_text(_seg4_slide(reopened, SEG4_OPENER_TITLE))
    assert "free tier at the time of writing" in opener.lower()
    assert "September 2026" in opener

    promised = [SEG4_OPENER_TITLE] + [t for _, t, _ in SEG4_CAPABILITIES]
    for title in promised:
        notes = _seg4_slide(reopened, title).notes_slide.notes_text_frame.text.lower()
        assert "re-verify" in notes, f"{title!r}: no re-verification reminder"
        assert "week before" in notes, f"{title!r}: the reminder has no deadline"


def test_notebooklm_slides_carry_exactly_the_plans_three_scenarios(reopened):
    """Three NotebookLM slides, one per scenario in the plan's order, each
    flagged as the Lesson 3 tool."""
    section = _plan_section("### NotebookLM at use-case level")
    for (plan_scenario, title, token), number in zip(
            SEG4_NOTEBOOK_SCENARIOS, range(1, 4)):
        assert f"{number}. **{plan_scenario}**" in section, (
            f"scenario {number} in the plan is not {plan_scenario!r}"
        )
        case_body = section[section.index(f"{number}. **{plan_scenario}**"):]
        assert token.lower() in case_body.lower(), (
            f"{token!r} is not in the plan's scenario {number}"
        )
        text = _slide_text(_seg4_slide(reopened, title))
        assert token.lower() in text.lower(), (
            f"{title!r} does not carry the plan's scenario ({token!r})"
        )
        assert "Lesson 3" in text, f"{title!r} does not flag the Lesson 3 tool"

    # Exactly three: no fourth NotebookLM slide crept in.
    notebook_slides = [
        s for s in _segment_ranges(reopened)[SEG4_NAME]
        if _slide_title(s).startswith("NotebookLM")
    ]
    assert len(notebook_slides) == len(SEG4_NOTEBOOK_TITLES)


def test_notebooklm_stays_at_scenario_level(reopened):
    """Scenarios today, the walkthrough in Lesson 3 — so the slides carry
    no step-by-step of their own."""
    for title in SEG4_NOTEBOOK_TITLES:
        lowered = _slide_text(_seg4_slide(reopened, title)).lower()
        for marker in ("step 1", "step 2", "click the", "click on",
                       "sign in", "open the menu"):
            assert marker not in lowered, f"{title!r} walks through {marker!r}"


def test_notebooklm_slides_carry_no_lane_colour_of_their_own(reopened):
    """Violet says "this is the Daily Driver's tool". NotebookLM is a
    different product and the course does not open it until Lesson 3, so
    none of its content blocks may carry it — the only violet left is what
    add_content_slide paints on every lane slide (the chip and the accent
    lead under the title)."""
    from deck.lesson2 import template

    violet = template.DAILY_DRIVER.color
    chrome = _fill_count(
        template.add_content_slide(
            template.new_presentation(), "baseline", lane=template.DAILY_DRIVER,
        ),
        violet,
    )
    assert chrome, "add_content_slide paints no lane colour — the test is moot"

    for title in SEG4_NOTEBOOK_TITLES:
        assert _fill_count(_seg4_slide(reopened, title), violet) == chrome, (
            f"{title!r} colours a content block as the Daily Driver's tool"
        )
    # Control: a capability that is the Daily Driver's does colour its blocks.
    assert _fill_count(_seg4_slide(reopened, SEG4_GEMS_TITLE), violet) > chrome


def test_chaining_slide_chains_what_the_segment_demonstrated(reopened):
    """The closing slide is the seed of the assignment's workflow thinking:
    the two chains the segment just built, and nothing new."""
    text = _slide_text(_seg4_slide(reopened, SEG4_CHAIN_TITLE))
    assert "Upload → Gem → Canvas" in text
    for step in ("Upload", "Gem", "Canvas", "notebook"):
        assert step.lower() in text.lower(), f"the chains lose {step!r}"
    assert "workflow" in text.lower()
    assert "assignment" in text.lower()


def test_every_seg4_slide_has_speaker_notes(reopened):
    for title in SEG4_TITLES:
        slide = _seg4_slide(reopened, title)
        assert slide.has_notes_slide, f"{title!r} has no notes slide"
        assert slide.notes_slide.notes_text_frame.text.strip(), (
            f"{title!r} has empty speaker notes"
        )


# Lane separation — that this segment never names the other lane — is
# asserted for every lane-bound segment at once in the integration section
# (test_every_lane_declared_segment_stays_in_its_lane).


def test_seg4_names_youtube_without_linking_one(reopened):
    """The plan's first scenario is a "YouTube link turned into a textual
    source": the deck names YouTube — the banned thing is the link, not the
    word — and carries no URL, where the temptation is greatest."""
    slide = _seg4_slide(reopened, SEG4_NOTEBOOK_TITLES[0])
    corpus = _slide_corpus(slide)
    assert "YouTube" in corpus, "the plan's scenario is not recognisable"
    assert "captions" in corpus.lower()
    assert not re.search(r"https?://|www\.", corpus), (
        "the YouTube scenario carries a link"
    )


# --- Segment 5: Assignment briefing + close (tickets 08, 10, 12) -------------

SEG5_NAME = "Assignment briefing + close"

SEG5_ASSIGNMENT_TITLE = "Assignment: your first workflow orchestration"
SEG5_ROADMAP_TITLE = "The roadmap: eight steps, end to end"
SEG5_BLUEPRINT_TITLES = [
    "The blueprint · steps 1–3",
    "The blueprint · steps 4–6",
    "The blueprint · steps 7–8",
]
SEG5_OUTCOME_TITLE = "The outcome, on the 7 Levels map"
SEG5_EFFORT_TITLE = "Effort and where the work is shown"
SEG5_CLOSE_TITLE = "Next: Practice Week 1 and Lesson 3"

SEG5_TITLES = [
    SEG5_ASSIGNMENT_TITLE, SEG5_ROADMAP_TITLE, *SEG5_BLUEPRINT_TITLES,
    SEG5_OUTCOME_TITLE, SEG5_EFFORT_TITLE, SEG5_CLOSE_TITLE,
]

LESSON1_PATH = PLAN_PATH.parent / "lesson_1.md"

ASSIGNMENT_SECTION = "### Assignment slide (as it will appear)"
ROADMAP_SECTION = "### The roadmap (≈2 min)"
BLUEPRINT_SECTION = "### The blueprint, step by step"
EFFORT_SECTION = "### Effort and where the work is shown"
CLOSE_SECTION = "### Close (≈1 min)"

# The two labels a step carries when it is in neither lane: the absence of
# any AI tool, and a tool outside the tenant with its clearance condition.
SEG5_NO_TOOL = "No lane"
SEG5_NO_LANE = "Outside both lanes · anonymised content only"


def _seg5_slide(reopened, title):
    """The Segment 5 slide with exactly this title."""
    return _segment_slide(reopened, SEG5_NAME, title)


def _blueprint_slides(reopened):
    return [_seg5_slide(reopened, title) for title in SEG5_BLUEPRINT_TITLES]


def _blueprint_text(reopened):
    """Everything the three blueprint slides show, joined — the blueprint
    is split across slides for room, so what it says is what the three of
    them say together."""
    return "\n".join(_slide_text(s) for s in _blueprint_slides(reopened))


def _blueprint_corpus(reopened):
    """The blueprint slides' text and notes. The full instruction for a
    step does not fit a card and lives in the notes, so a test about the
    instructions has to read those too."""
    return "\n".join(_slide_corpus(s) for s in _blueprint_slides(reopened))


def _blueprint_steps_from_plan():
    """The rows of the blueprint table in lesson_2.md, as dicts keyed by the
    table's own column names. The plan authors the chain; the slides render
    it, so the steps are read out of the plan rather than transcribed
    here."""
    section = _plan_section(BLUEPRINT_SECTION)
    header, steps = None, []
    for line in section.splitlines():
        line = line.strip()
        if not line.startswith("|") or set(line) <= set("|- "):
            continue
        cells = [c.strip() for c in line.strip("|").split("|")]
        if header is None:
            header = cells
            continue
        steps.append(dict(zip(header, cells)))
    return steps


def _seven_levels_from_lesson_1():
    """Level number -> level name, from the 7 Levels table in lesson_1.md.
    The map is Lesson 1's, so the outcome slide is checked against where it
    was taught rather than against a name transcribed here."""
    text = LESSON1_PATH.read_text(encoding="utf-8")
    section = text[text.index("### Part B: The 7 Levels of AI"):
                   text.index("**Vibe coding")]
    levels = {}
    for line in section.splitlines():
        line = line.strip()
        if not line.startswith("|") or set(line) <= set("|- "):
            continue
        cells = [c.strip() for c in line.strip("|").split("|")]
        if not cells[0].isdigit():
            continue
        levels[int(cells[0])] = cells[1]
    return levels


def test_seg5_slides_present_in_plan_order_and_not_placeholders(reopened):
    """All eight slides live after Segment 5's divider, in the plan's order
    — the assignment, the roadmap, the blueprint in three screenfuls, the
    outcome on the 7 Levels map, the effort and its two venues, the close —
    with no placeholder left."""
    segment = _segment_ranges(reopened)[SEG5_NAME]
    assert len(segment) == len(SEG5_TITLES)
    assert [_slide_title(s) for s in segment] == SEG5_TITLES
    for slide in segment:
        assert "[PLACEHOLDER]" not in _slide_text(slide)


# Every element of the plan's assignment, as a token asserted to appear in
# the plan's own assignment section first, so the slide cannot drift from
# the wording participants are told to screenshot. The first item of each
# pair only names the element in the failure message.
SEG5_ASSIGNMENT_ELEMENTS = [
    ("chain", "multi-step workflow orchestration"),
    ("blueprint", "the tool, the input, the output"),
    ("blueprint", "how you ran it"),
    ("blueprint", "skip or change"),
    ("deliverable", "hand-picked topic"),
    ("deliverable", "educational purposes"),
    ("deliverable", "learning outcomes"),
    ("routing", "touches sensitive data"),
    ("routing", "anonymised before upload"),
    ("routing", "the lane per step"),
]


def test_assignment_slide_carries_every_element_of_the_plans_assignment(reopened):
    """The plan's three numbered items, whole, on the slide that opens the
    segment."""
    plan = _plan_section(ASSIGNMENT_SECTION).lower()
    text = _slide_text(_seg5_slide(reopened, SEG5_ASSIGNMENT_TITLE)).lower()
    for element, token in SEG5_ASSIGNMENT_ELEMENTS:
        assert token in plan, f"{token!r} is not in the plan's assignment"
        assert token in text, f"the assignment slide drops the {element}: {token!r}"


def test_the_assignment_points_at_the_blueprint_rather_than_describing_one(reopened):
    """The fix ticket 10 exists for: the briefing is the chain, not a slide
    about producing one. The assignment slide sends the room to the roadmap
    that follows, and the roadmap is what they are told to screenshot."""
    assignment = _slide_text(_seg5_slide(reopened, SEG5_ASSIGNMENT_TITLE))
    assert "roadmap on the next four slides" in assignment.lower(), (
        "the assignment does not point at the blueprint that follows it"
    )
    roadmap = _slide_text(_seg5_slide(reopened, SEG5_ROADMAP_TITLE))
    assert "Screenshot this slide" in roadmap
    assert "Screenshot" not in assignment, (
        "the screenshot cue is on the assignment slide, not on the roadmap"
    )


def test_roadmap_slide_draws_every_step_of_the_plans_blueprint(reopened):
    """The roadmap is the plan's Mermaid diagram, drawn: one card per step,
    numbered, each naming the tool it runs in, a card for Prompt Coach
    beside the steps, and a terminal card for the deliverable."""
    steps = _blueprint_steps_from_plan()
    assert len(steps) == 8, f"the plan's blueprint has {len(steps)} steps"
    text = _slide_text(_seg5_slide(reopened, SEG5_ROADMAP_TITLE))
    for step in steps:
        assert f"STEP {step['#']}" in text, f"step {step['#']} is not on the roadmap"
    assert "STEP 9" not in text
    assert "Prompt Coach" in text, "Prompt Coach is not on the roadmap"
    assert "The presentation" in text, "the chain ends nowhere"


def test_roadmap_shows_the_chain_crossing_lanes(reopened):
    """The routing argument the picture makes: both lanes appear in the
    roadmap's key, in their own colours, and the steps in neither lane
    are keyed too."""
    from deck.lesson2 import template

    slide = _seg5_slide(reopened, SEG5_ROADMAP_TITLE)
    text, fills = _slide_text(slide), _fills(slide)
    for lane in (template.SENSITIVE_LANE, template.DAILY_DRIVER):
        assert lane.name in text, f"the roadmap's key omits {lane.name!r}"
        assert str(lane.color) in fills, f"{lane.name!r} is keyed in no colour"
    assert "No lane" in text, "the steps in neither lane are unkeyed"
    assert "inherits a lane" in text, (
        "the roadmap draws the lane changes but never says what they mean"
    )


def test_roadmap_says_which_steps_are_required(reopened):
    """Only the step that handles sensitive data is required, and only when
    the participant's material contains any; every other step says it is
    optional on its own card, so the route reads as a recommendation with
    one rule in it rather than a quota."""
    steps = _blueprint_steps_from_plan()
    required = [s["#"] for s in steps if s["Required"].startswith("Required")]
    assert required == ["2"], f"required steps in the plan: {required}"
    assert "sensitive data" in steps[1]["Required"]
    plan = _plan_section(ROADMAP_SECTION)
    assert "only step 2" in plan, "the plan's roadmap never says which step is required"
    text = _slide_text(_seg5_slide(reopened, SEG5_ROADMAP_TITLE))
    assert "Only step 2 is required" in text
    assert text.count("Optional") >= len(steps) - 1, (
        "the optional steps are not tagged optional on their cards"
    )
    assert "Required if sensitive data" in text


def test_roadmap_and_plan_agree_on_the_gate_that_is_not_a_step(reopened):
    """Checking and logging are per step, not per project, so they are a
    gate across the whole chain rather than another card."""
    plan = _plan_section(ROADMAP_SECTION)
    text = _slide_text(_seg5_slide(reopened, SEG5_ROADMAP_TITLE))
    for token in ("check every output", "log the tool"):
        assert token in plan, f"{token!r} is not in the plan's roadmap section"
        assert token in text, f"the roadmap slide drops {token!r}"
    assert "AI-use declaration" in text


def test_the_plan_hands_out_a_mermaid_source_for_the_chain(reopened):
    """The diagram is part of the deliverable, so the plan carries a source
    participants can fork — and the deck says so rather than leaving it in
    a file nobody opens."""
    plan = _plan_section(ROADMAP_SECTION)
    assert "```mermaid" in plan, "the plan draws no diagram"
    assert "flowchart" in plan
    steps = _blueprint_steps_from_plan()
    for step in steps:
        assert f'S{step["#"]}[' in plan, f"step {step['#']} is not in the diagram"
    assert "PC[" in plan, "Prompt Coach is not in the diagram"
    notes = _slide_corpus(_seg5_slide(reopened, SEG5_ROADMAP_TITLE))
    assert "Mermaid" in notes, "the facilitator is never told the source exists"


# The three fields every blueprint row shows: what goes in and where from,
# what to do, what comes out and where to. The full instruction is a
# paragraph per step and lives in the speaker notes and the handout.
SEG5_BLUEPRINT_FIELDS = ["IN", "DO", "OUT"]


def test_blueprint_slides_carry_every_step_with_all_of_its_fields(reopened):
    """The blueprint the plan authors, rendered: every step, with its tool,
    its lane, whether it is required, what goes in, what to do and what
    comes out."""
    steps = _blueprint_steps_from_plan()
    text = _blueprint_text(reopened)
    for step in steps:
        assert f"{step['#']} · {step['Step']}" in text, (
            f"step {step['#']} ({step['Step']!r}) is not in the blueprint"
        )
        assert step["Tool"] in text, f"step {step['#']} names no tool"
        assert step["Required"] in text, f"step {step['#']} does not say if it is required"
    for field in SEG5_BLUEPRINT_FIELDS:
        found = len(re.findall(rf"\b{re.escape(field)}\b", text))
        assert found == len(steps), (
            f"the {field!r} field is on {found} steps, not {len(steps)}"
        )
    assert "WHY THIS TOOL" not in text, "the justification boxes are back"


def test_every_step_says_where_its_output_goes_and_its_input_comes_from(reopened):
    """The chain has to read as a chain: every step but the last names the
    step that takes its output, and every step that consumes an earlier
    result names the step it came from — in the plan's table and on the
    slides, so nobody has to guess how one result reaches the next step."""
    steps = _blueprint_steps_from_plan()
    text = _blueprint_text(reopened)
    for step in steps[:-1]:
        assert "→ step" in step["Out"], (
            f"step {step['#']}'s output names no step that takes it"
        )
        assert step["Out"] in text, f"step {step['#']}'s OUT is not on the slide"
    assert "deliverable" in steps[-1]["Out"]
    for step in steps[3:]:
        assert "step" in step["In"].lower() or "(5)" in step["In"], (
            f"step {step['#']}'s input names no step it comes from"
        )


def test_every_step_states_its_lane(reopened):
    """The assignment's third item is routing per step, so the blueprint
    states a lane on every row — either one CONTEXT.md defines, the
    absence of any AI tool, or the explicit absence of both lanes, which
    carries the clearance condition."""
    from deck.lesson2 import template

    context = CONTEXT_PATH.read_text(encoding="utf-8")
    steps = _blueprint_steps_from_plan()
    text = _blueprint_text(reopened)
    lanes = {
        "Sensitive Lane": template.SENSITIVE_LANE.name,
        "Daily Driver": template.DAILY_DRIVER.name,
        "None (manual)": SEG5_NO_TOOL,
        "Outside both lanes": SEG5_NO_LANE,
    }
    for lane in ("Sensitive Lane", "Daily Driver"):
        assert f"**{lane}**" in context, f"{lane!r} is not a CONTEXT.md lane"
    for step in steps:
        expected = lanes[step["Lane"]]
        assert expected in text, f"step {step['#']} states no lane"
    assert text.count(SEG5_NO_TOOL) == 1, "more than one step uses no AI tool, or none does"
    assert text.count(SEG5_NO_LANE) == 1, (
        "more than one step sits outside both lanes, or none does"
    )


def test_the_anonymisation_step_uses_no_ai_tool_and_acrobat_is_out(reopened):
    """Anonymisation happens before any file reaches any AI tool, so the
    step that does it names no AI tool — and Adobe Acrobat, which needs
    administrator approval participants cannot get, is out of the chain
    and the plan says why."""
    steps = _blueprint_steps_from_plan()
    step2 = steps[1]
    assert step2["Step"] == "Anonymise your own material"
    assert step2["Tool"].startswith("No AI tool")
    plan = _plan_section(BLUEPRINT_SECTION)
    assert "Acrobat" not in plan
    assert "Acrobat" not in _blueprint_text(reopened)
    roadmap = _plan_section(ROADMAP_SECTION)
    assert "Why Adobe Acrobat is not in the chain" in roadmap


def test_the_full_instructions_reach_the_facilitator_notes(reopened):
    """The full instruction is a paragraph per step; it does not fit a
    card, so it has to be somewhere a facilitator reading only the notes
    still finds it."""
    corpus = _blueprint_corpus(reopened)
    section = _plan_section(BLUEPRINT_SECTION)
    for step in _blueprint_steps_from_plan():
        assert f"Step {step['#']} — {step['Step']}" in section, (
            f"the plan writes no instructions for step {step['#']}"
        )
        assert f"Step {step['#']} — {step['Step']}" in corpus, (
            f"step {step['#']}'s instructions never reach the notes"
        )
    assert corpus.count("How:") == 8
    assert "Why this tool here" not in corpus
    assert "Prompt Coach, at any step" in section
    assert "Prompt Coach, at any step" in corpus


# Every tool the course has introduced and still chains, and the place it
# is used. The assignment says "the tools the course has introduced so
# far", so this is what that sentence has to cash out as. Adobe Acrobat is
# taught in Segment 3 but is not in the chain — see the test above.
SEG5_CHAIN_TOOLS = [
    "Copilot Chat (Basic)", "Prompt Coach", "Study and Learn",
    "Canva", "genspark.ai",
    "Gemini", "Guided Learning", "Deep Research", "a Gem", "file upload",
    "long context", "Canvas", "image generation", "NotebookLM",
]


def test_the_blueprint_chains_every_tool_the_course_has_taught(reopened):
    """"All the tools so far" is a memory test unless the chain names them,
    so every tool the plan's blueprint uses appears on a blueprint slide,
    in the step it belongs to."""
    plan = _plan_section(BLUEPRINT_SECTION)
    text = _blueprint_text(reopened)
    for tool in SEG5_CHAIN_TOOLS:
        assert tool in plan, f"{tool!r} is in no step of the plan's blueprint"
        assert tool in text, f"the blueprint never names {tool!r}"


def test_the_deliverable_is_produced_by_the_chain(reopened):
    """The chain has to end in the thing the assignment asks for, and the
    steps that make it teach have to be in it — otherwise the roadmap is a
    tour of tools rather than a route to a deliverable."""
    corpus = _blueprint_corpus(reopened) + _slide_corpus(
        _seg5_slide(reopened, SEG5_ROADMAP_TITLE))
    for token in ("flashcard", "quiz", "video excerpt", "learning outcome"):
        assert token in corpus.lower(), f"the chain never produces {token!r}"


def test_outcome_slide_states_level_3_achieved_and_level_4_founded(reopened):
    """The plan's outcome sentence, against the map it refers to: Level 3
    achieved, Level 4's foundation established, both under the names the
    7 Levels table in lesson_1.md gives them."""
    plan = _plan_section(ASSIGNMENT_SECTION)
    text = _slide_text(_seg5_slide(reopened, SEG5_OUTCOME_TITLE))
    for token in ("Level 3 achieved", "foundation for Level 4 established"):
        assert token in plan, f"{token!r} is not the plan's outcome"
        assert token in text, f"the outcome slide does not say {token!r}"
    levels = _seven_levels_from_lesson_1()
    assert levels[3] == "Power User" and levels[4] == "Tool Orchestrator"
    assert "7 Levels" in text, "the outcome is not stated against the map"
    for number in (2, 3, 4):
        assert f"Level {number}" in text, f"Level {number} is not on the map"
        assert levels[number] in text, f"Level {number} is unnamed"
    assert "step 5 of the roadmap" in text, "the Gem is credited to the wrong step"


# The plan's effort bullet and its two venues, each asserted against the
# plan's own section before it is looked for on the slide.
SEG5_EFFORT_TOKENS = [
    "3–4 hours", "blueprint first",
    "Practice Week 1", "W3", "Lesson 3 debrief", "W4", "2–3 volunteers",
]


def test_effort_slide_states_the_effort_and_both_venues(reopened):
    plan = _plan_section(EFFORT_SECTION)
    slide = _seg5_slide(reopened, SEG5_EFFORT_TITLE)
    text, corpus = _slide_text(slide), _slide_corpus(slide)
    for token in SEG5_EFFORT_TOKENS:
        assert token in plan, f"{token!r} is not in the plan's effort section"
        assert token in corpus, f"the effort slide drops {token!r}"
    for token in ("3–4 hours", "Practice Week 1", "Lesson 3"):
        assert token in text, f"{token!r} is only in the notes"


def test_effort_slide_says_what_done_looks_like(reopened):
    """Eight steps is a route, not a quota, and the three things that get
    handed in are checkable rather than a feeling — both are the plan's
    words, and both are what stops the room optimising the wrong half."""
    plan = _plan_section(EFFORT_SECTION)
    text = _slide_text(_seg5_slide(reopened, SEG5_EFFORT_TITLE))
    for token in ("not a quota", "only step 2 is required", "AI-use log"):
        assert token in plan, f"{token!r} is not in the plan's effort section"
        assert token in text, f"the effort slide drops {token!r}"


# The close's terms, each asserted to be CONTEXT.md's ubiquitous language
# before it is looked for on the slide.
SEG5_CLOSE_TERMS = ["Practice Week", "Wiki Sandbox", "KM System"]


def test_close_names_practice_week_1_and_lesson_3_in_context_terms(reopened):
    plan = _plan_section(CLOSE_SECTION)
    context = CONTEXT_PATH.read_text(encoding="utf-8")
    text = _slide_text(_seg5_slide(reopened, SEG5_CLOSE_TITLE))
    for term in SEG5_CLOSE_TERMS:
        assert f"**{term}**" in context, f"{term!r} is not a CONTEXT.md term"
        assert term in text, f"the close does not use {term!r}"
    for token in ("Practice Week 1", "W3", "drop-in", "Lesson 3", "W4",
                  "NotebookLM", "Perplexity"):
        assert token in plan, f"{token!r} is not in the plan's close"
        assert token in text, f"the close does not name {token!r}"


def test_close_promises_no_new_material_in_the_practice_week(reopened):
    """CONTEXT.md's Practice Week is a week with no new material; the close
    says so, because that is what makes the drop-in optional."""
    text = _slide_text(_seg5_slide(reopened, SEG5_CLOSE_TITLE)).lower()
    assert "no new material" in text
    assert "optional" in text


def test_segment_5_names_both_lanes(reopened):
    """Segments 3 and 4 each stay in one lane; the chain crosses both, so
    this one names both tools and both lanes."""
    corpus = "\n".join(_segment_corpus(reopened, SEG5_NAME)).lower()
    for token in ("sensitive lane", "daily driver", "copilot", "gemini"):
        assert token in corpus, f"Segment 5 never names {token!r}"


def test_every_seg5_slide_has_speaker_notes(reopened):
    for title in SEG5_TITLES:
        slide = _seg5_slide(reopened, title)
        assert slide.has_notes_slide, f"{title!r} has no notes slide"
        assert slide.notes_slide.notes_text_frame.text.strip(), (
            f"{title!r} has empty speaker notes"
        )


# --- Deck-wide --------------------------------------------------------------


def test_no_youtube_link_anywhere_in_the_deck(reopened):
    """The prompt bans YouTube as a reference. The word itself stays
    allowed — Segment 4 describes turning a YouTube link into a textual
    source — so what is banned is a link to one."""
    for i, slide in enumerate(reopened.slides):
        lowered = _slide_corpus(slide).lower()
        for marker in ("youtube.com", "youtu.be", "m.youtube"):
            assert marker not in lowered, f"{marker!r} on slide {i + 1}"


# --- Master, theme and the Lesson 2 look ------------------------------------


def _theme_xml(reopened):
    from pptx.opc.constants import RELATIONSHIP_TYPE as RT

    return reopened.slide_masters[0].part.part_related_by(RT.THEME).blob.decode("utf-8")


def test_master_defines_the_lesson_2_footer_and_theme_fonts(reopened):
    master = reopened.slide_masters[0]
    master_text = "\n".join(
        s.text_frame.text for s in master.shapes if s.has_text_frame
    )
    assert "Lesson 2 · Week 2" in master_text
    assert "Lesson 1" not in master_text
    theme_xml = _theme_xml(reopened)
    assert 'typeface="Cambria"' in theme_xml
    assert 'typeface="Calibri"' in theme_xml


def test_look_is_not_lesson_1s(reopened):
    """A light canvas, non-Segoe fonts, and none of Lesson 1's theme colours."""
    from deck import template as lesson1

    theme_xml = _theme_xml(reopened)
    assert "Segoe" not in theme_xml
    for color in (lesson1.PRIMARY, lesson1.ACCENT, lesson1.HIGHLIGHT):
        assert str(color) not in theme_xml
    canvas = reopened.slide_masters[0].background.fill.fore_color.rgb
    assert canvas != lesson1.CANVAS
    r, g, b = canvas[0], canvas[1], canvas[2]
    assert min(r, g, b) > 0xE0, "Lesson 2 canvas is light, not navy"


def test_fonts_and_sizes_are_explicit_everywhere(reopened):
    for i, slide in enumerate(reopened.slides):
        for shape in slide.shapes:
            if not shape.has_text_frame:
                continue
            for para in shape.text_frame.paragraphs:
                for run in para.runs:
                    assert run.font.name in ("Cambria", "Calibri"), (
                        f"slide {i + 1}: font {run.font.name!r} in {run.text!r}"
                    )
                    assert run.font.size is not None, (
                        f"slide {i + 1}: inherited size in {run.text!r}"
                    )


# --- Lane colour-coding ------------------------------------------------------


def _fills(slide):
    fills = set()
    for shape in slide.shapes:
        try:
            if shape.fill.type is not None and shape.fill.fore_color.rgb:
                fills.add(str(shape.fill.fore_color.rgb))
        except (AttributeError, TypeError):
            continue
    return fills


def test_lane_segments_are_colour_coded_by_lane(reopened):
    """Segment 3 (Copilot) carries the Sensitive Lane colour and never the
    Daily Driver's; Segment 4 (Gemini) the reverse — divider and every
    slide. Segments about both lanes (2, 5) or neither (1) may use either
    colour where they name a tool, so they are not constrained here."""
    from deck.lesson2 import template

    teal, violet = str(template.SENSITIVE_LANE.color), str(template.DAILY_DRIVER.color)
    all_slides = list(reopened.slides)
    ranges = _segment_ranges(reopened)
    expected = {
        "Copilot Chat (Basic) in practice": (teal, violet),
        "Gemini in depth (+ NotebookLM preview)": (violet, teal),
    }
    for name, (own, other) in expected.items():
        slides = ranges[name]
        divider = all_slides[all_slides.index(slides[0]) - 1]
        for slide in [divider, *slides]:
            fills = _fills(slide)
            assert own in fills, f"{name}: slide without its lane colour"
            assert other not in fills, f"{name}: other lane's colour present"


# --- Terminology ------------------------------------------------------------

# Forbidden synonyms from the CONTEXT.md Avoid-lists, the same list the
# Lesson 1 deck is audited against (see tests/test_deck.py for why the
# generic words are left out).
FORBIDDEN_TERMS = [
    "homework week", "self-study week", "blank week", "blank lesson",
    "breakout team", "domain team",
    "knowledge base", "shared wiki", "team wiki",
    "use case", "case study", "team project",
    "home base", "primary AI environment",
    "workhorse", "main tool",
    "tool ecosystem", "satellite tool",
]


def test_terminology_audit_no_forbidden_synonyms(reopened):
    for i, slide in enumerate(reopened.slides):
        corpus = _slide_corpus(slide).lower()
        for term in FORBIDDEN_TERMS:
            assert term.lower() not in corpus, f"{term!r} on slide {i + 1}"


# --- Integration pass: flow, timing, lane separation (ticket 09) --------------


def test_no_placeholder_slides_remain_anywhere(reopened):
    for i, slide in enumerate(reopened.slides):
        assert "[PLACEHOLDER]" not in _slide_text(slide), f"slide {i + 1}"


# Slides per segment, reviewed against the Timing Overview durations and
# pinned. Segment 4 sits exactly at the density guard's limit on purpose:
# the plan's demo script runs five of its capability slides as one
# 1.5-minute batch, so 15 slides in 15 minutes is the plan's shape, not
# drift. Growing a segment past its pin means changing lesson_2.md first.
EXPECTED_SEGMENT_SLIDE_COUNTS = {
    "Recap + assignment debrief": 3,
    "The sensitive-data rule and data-protection principles": 7,
    "Copilot Chat (Basic) in practice": 7,
    "Gemini in depth (+ NotebookLM preview)": 15,
    "Assignment briefing + close": 8,
}


def test_segment_slide_counts_fit_their_timings(reopened):
    # The density guard fires if a later ticket grows a segment past one
    # slide per minute of its TIMING_OVERVIEW duration (and bumps the pin).
    minutes = {name: int(duration.split()[0]) for name, duration, _ in TIMING_OVERVIEW}
    for name, slides in _segment_ranges(reopened).items():
        expected = EXPECTED_SEGMENT_SLIDE_COUNTS[name]
        assert len(slides) == expected, (
            f"{name}: {len(slides)} slides, expected {expected}"
        )
        assert len(slides) <= minutes[name], f"{name}: over one slide per minute"


def test_the_session_fits_the_hour(reopened):
    """The durations the dividers show the room add up to the length the
    title slide promises. The 60-minute cap is why the Session Template's
    hands-on block is absent from this lesson, so a segment that grows has
    to take its minutes from another one."""
    shown = [
        int(re.search(r"(\d+) min", _slide_text(divider)).group(1))
        for divider in _dividers(reopened)
    ]
    assert len(shown) == len(TIMING_OVERVIEW)
    promised = int(re.search(r"(\d+) minutes", _slide_text(reopened.slides[0])).group(1))
    assert sum(shown) == promised, f"the dividers add up to {sum(shown)} minutes"
    assert promised <= 60


def test_every_content_slide_has_speaker_notes(reopened):
    # Everything except the title slide and the five section dividers.
    dividers = _dividers(reopened)
    for i, slide in enumerate(reopened.slides):
        if i == 0 or slide in dividers:
            continue
        assert slide.has_notes_slide, f"slide {i + 1} has no notes slide"
        assert slide.notes_slide.notes_text_frame.text.strip(), (
            f"slide {i + 1} has empty speaker notes"
        )


def test_every_live_demo_cue_sits_in_a_demo_segment_and_is_mirrored_in_notes(reopened):
    """The lanes are demonstrated one at a time, so every switch to a live
    tool belongs to Segment 3 or Segment 4 and no other segment opens one.
    (How many each of those two runs is pinned in their own sections.) Each
    switch is marked on its slide and repeated in that slide's notes, so a
    facilitator reading only the notes still knows to switch."""
    ranges = _segment_ranges(reopened)
    demo_slides = ranges[SEG3_NAME] + ranges[SEG4_NAME]
    cued = [
        (i, s) for i, s in enumerate(reopened.slides) if "LIVE DEMO" in _slide_text(s)
    ]
    assert cued, "the deck cues no live demo at all"
    for i, slide in cued:
        assert slide in demo_slides, (
            f"slide {i + 1}: a live demo outside the two demo segments"
        )
        assert slide.has_notes_slide, f"slide {i + 1} has no notes slide"
        notes = slide.notes_slide.notes_text_frame.text
        assert re.search(r"\blive\b", notes, re.IGNORECASE), (
            f"slide {i + 1}: the demo is cued on the slide but not in its notes"
        )


# What naming a lane looks like on a slide, beyond the lane's own name and
# the brand of its tool — the words that give a lane away and cannot be read
# off template.Lane. NotebookLM belongs to the Daily Driver's vocabulary: it
# is previewed in Segment 4 and has no business in Segment 3.
LANE_TELLS = {
    "Sensitive Lane": ("Microsoft",),
    "Daily Driver": ("Gemini", "NotebookLM"),
}


def test_every_lane_declared_segment_stays_in_its_lane(reopened):
    """The two lanes are demonstrated separately, so a segment that belongs
    to one stays in it: Segment 3 never names Gemini or the Daily Driver,
    Segment 4 never names Copilot or the Sensitive Lane — on a slide or in a
    note, anywhere in the segment, divider included. Segments that belong to
    neither lane — the recap, the rule, the assignment — may name both.

    Which segments are bound is read off deck/lesson2/build.py rather than
    written out here, so the rule holds for the deck rather than for two
    segments named in a test."""
    from deck.lesson2 import build, template

    declared = [segment for segment in build.SEGMENTS if segment.lane]
    assert len(declared) == 2, "the two demo segments are the lane-bound ones"
    for segment in declared:
        other = next(lane for lane in template.LANES if lane is not segment.lane)
        forbidden = (other.name, other.tool.split()[0], *LANE_TELLS[other.name])
        for text in _segment_corpus(reopened, segment.title):
            lowered = text.lower()
            for term in forbidden:
                assert term.lower() not in lowered, (
                    f"Segment {segment.number} names {term!r}"
                )


def test_one_spelling_of_anonymise_deck_wide(reopened):
    """The deck follows lesson_2.md's British -ise throughout — the
    institutional rule it quotes is spelled that way. A stray "anonymize"
    two lines from "anonymise before upload" reads as a typo."""
    for i, slide in enumerate(reopened.slides):
        assert not re.search(r"anonymiz", _slide_corpus(slide), re.IGNORECASE), (
            f"slide {i + 1}: the other spelling of anonymise"
        )


# --- Builders keep the Lesson 1 call shapes ----------------------------------


def test_builders_accept_the_lesson_1_call_shapes_and_survive_a_roundtrip():
    """A segment author calls the Lesson 2 template exactly as a Lesson 1
    segment module calls deck/template.py; the result opens after a
    save-and-reopen."""
    from pptx.util import Inches

    from deck.lesson2 import template

    prs = template.new_presentation()
    slide = template.add_content_slide(
        prs, "Sample", ["one", (1, "two")], notes="n", lane=template.DAILY_DRIVER,
    )
    template.add_demo_cue(slide, "open Gemini", Inches(1.6), lane=template.DAILY_DRIVER)
    template.add_band(slide, Inches(2.4), Inches(0.8), "Band", sub="sub")
    template.add_card(slide, Inches(0.6), Inches(3.5), Inches(4.0), Inches(1.5),
                      template.DAILY_TINT, "Card", template.INK, "body")
    template.add_card(slide, Inches(4.8), Inches(3.5), Inches(4.0), Inches(1.5),
                      template.SENSITIVE_LANE.color, "Label block", template.WHITE)
    template.add_label(slide, Inches(0.6), Inches(5.3), Inches(4.0), "caption")

    fd, path = tempfile.mkstemp(suffix=".pptx")
    os.close(fd)
    try:
        prs.save(path)
        reopened = Presentation(path)
    finally:
        os.unlink(path)
    text = _slide_text(reopened.slides[0])
    for expected in ("Sample", "one", "two", "LIVE DEMO — open Gemini", "Band",
                     "sub", "Card", "body", "Label block", "caption",
                     "DAILY DRIVER"):
        assert expected in text, expected
    assert reopened.slides[0].notes_slide.notes_text_frame.text == "n"
