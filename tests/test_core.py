"""Tests for the shared deck core (ticket 02 of the Lesson 2 effort).

The core owns the lesson-independent PowerPoint mechanics; a lesson hands
it a Design and gets back a presentation whose master and theme carry
that design. Assertions run against a saved-and-reopened .pptx, as the
Lesson 1 tests do.
"""

import os
import tempfile

import pytest
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.opc.constants import RELATIONSHIP_TYPE as RT

from deck import core

# A deliberately un-Lesson-1 design, so the tests prove the core carries
# whatever design it is handed rather than Lesson 1's constants.
DESIGN = core.Design(
    heading_font="Georgia",
    body_font="Calibri",
    scheme={"dk2": RGBColor(0x11, 0x22, 0x33), "accent1": RGBColor(0x44, 0x55, 0x66)},
    canvas=RGBColor(0xFA, 0xFA, 0xF0),
    footer_text="AI for Research Management · Lesson X · Week X",
    footer_rule=RGBColor(0xCC, 0xCC, 0xCC),
    footer_color=RGBColor(0x66, 0x66, 0x66),
)

_P = "http://schemas.openxmlformats.org/presentationml/2006/main"


def _roundtrip(prs):
    fd, path = tempfile.mkstemp(suffix=".pptx")
    os.close(fd)
    try:
        prs.save(path)
        return Presentation(path)
    finally:
        os.unlink(path)


@pytest.fixture(scope="module")
def reopened():
    return _roundtrip(core.new_presentation(DESIGN))


def test_presentation_is_16_9_for_teams(reopened):
    assert reopened.slide_width == core.SLIDE_WIDTH
    assert reopened.slide_height == core.SLIDE_HEIGHT
    assert round(reopened.slide_width / reopened.slide_height, 3) == round(16 / 9, 3)


def test_theme_carries_the_designs_fonts_and_colours(reopened):
    theme_xml = reopened.slide_masters[0].part.part_related_by(RT.THEME).blob.decode("utf-8")
    assert 'typeface="Georgia"' in theme_xml
    assert 'typeface="Calibri"' in theme_xml
    assert "112233" in theme_xml
    assert "445566" in theme_xml


def test_master_footer_carries_the_designs_text(reopened):
    master = reopened.slide_masters[0]
    master_text = "\n".join(s.text_frame.text for s in master.shapes if s.has_text_frame)
    assert "Lesson X · Week X" in master_text


def test_master_shape_ids_stay_within_signed_32_bit(reopened):
    """The shapes the core adds must stay below 2^31, or PowerPoint
    reports the file as corrupt (see master_next_id)."""
    master = reopened.slide_masters[0]
    added = [s for s in master.shapes if s.name in ("Footer Rule", "Footer Text")]
    assert len(added) == 2
    for shape in added:
        assert 0 < shape.shape_id < 2**31


def test_master_next_id_ignores_the_stock_near_2_31_layout_ids(reopened):
    """python-pptx's own next-id scan picks up the stock master's slide-
    layout ids (numbered from 2^31); the helper must not."""
    master = reopened.slide_masters[0]
    assert master.shapes._next_shape_id >= 2**31, "stock template gotcha no longer present"
    shape_ids = [int(el.get("id")) for el in master.shapes._spTree.iter(f"{{{_P}}}cNvPr")]
    next_id = core.master_next_id(master)
    assert next_id < 2**31
    assert next_id == max(shape_ids) + 1


def test_lesson_2_sub_package_builds_on_the_core_with_its_own_footer():
    from deck.lesson2.build import build_presentation

    reopened = _roundtrip(build_presentation())
    master_text = "\n".join(
        s.text_frame.text for s in reopened.slide_masters[0].shapes if s.has_text_frame
    )
    assert "Lesson 2 · Week 2" in master_text
    assert "Lesson 1" not in master_text
