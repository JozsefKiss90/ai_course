from __future__ import annotations

import subprocess
import sys
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parent.parent
SCRIPT = ROOT / "scripts" / "build_my_budget.py"
DATA = ROOT / "data" / "my_budget.csv"
WORKBOOK = ROOT / "output" / "my_budget.xlsx"


def test_builds_formula_driven_budget_workbook():
    result = subprocess.run(
        [sys.executable, str(SCRIPT)],
        cwd=ROOT,
        capture_output=True,
        text=True,
        check=False,
    )

    assert result.returncode == 0, result.stderr
    assert WORKBOOK.exists()

    budget = pd.read_csv(DATA, sep=";")
    categories = list(dict.fromkeys(budget["category"]))
    workbook = load_workbook(WORKBOOK, data_only=False)
    sheet = workbook["Budget"]
    last_data_row = len(budget) + 1
    last_summary_row = len(categories) + 1
    total_row = last_summary_row + 1

    assert sheet["H1"].value == "Line Cost (EUR)"
    assert sheet["H2"].value == "=D2*F2"
    assert sheet[f"H{last_data_row}"].value == (
        f"=D{last_data_row}*F{last_data_row}"
    )
    assert [sheet.cell(row=row, column=10).value for row in range(2, total_row)] == categories
    assert sheet["K2"].value == (
        f'=SUMIF($B$2:$B${last_data_row},J2,$H$2:$H${last_data_row})'
    )
    assert sheet[f"K{total_row}"].value == f"=SUM(K2:K{last_summary_row})"
    assert sheet["F2"].number_format == '#,##0.00 "EUR"'
    assert sheet["H2"].number_format == '#,##0.00 "EUR"'
    assert sheet["A1"].font.bold
    assert sheet.column_dimensions["C"].width >= 30

    assert len(sheet._charts) == 1
    chart = sheet._charts[0]
    assert chart.title is not None
    assert chart.series[0].cat.numRef.f == f"'Budget'!$J$2:$J${last_summary_row}"
    assert chart.series[0].val.numRef.f == f"'Budget'!$K$2:$K${last_summary_row}"

    for category in categories:
        assert category in result.stdout
