from __future__ import annotations

import sys
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.comments import Comment
from openpyxl.styles import Font, PatternFill


ROOT = Path(__file__).resolve().parent.parent
DATA = ROOT / "data" / "my_budget.csv"
OUTPUT_DIR = ROOT / "output"
WORKBOOK = OUTPUT_DIR / "my_budget.xlsx"

HEADERS = [
    "Category Code",
    "Category",
    "Item",
    "Quantity",
    "Unit",
    "Unit Cost (EUR)",
    "Source Note",
    "Line Cost (EUR)",
]
CURRENCY_FORMAT = '#,##0.00 "EUR"'


def build_workbook(budget: pd.DataFrame) -> None:
    categories = list(dict.fromkeys(budget["category"]))
    last_data_row = len(budget) + 1
    last_summary_row = len(categories) + 1
    total_row = last_summary_row + 1

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Budget"
    sheet.append(HEADERS)

    for row_number, row in enumerate(budget.itertuples(index=False), start=2):
        sheet.append(
            [
                row.category_code,
                row.category,
                row.item,
                row.quantity,
                row.unit,
                row.unit_cost_eur,
                row.source_note,
                f"=D{row_number}*F{row_number}",
            ]
        )

    sheet["J1"] = "Category"
    sheet["K1"] = "Category Total (EUR)"
    for row_number, category in enumerate(categories, start=2):
        sheet.cell(row=row_number, column=10, value=category)
        # SUMIF keeps the summary and chart current when line inputs change in Excel.
        sheet.cell(
            row=row_number,
            column=11,
            value=(
                f'=SUMIF($B$2:$B${last_data_row},J{row_number},'
                f'$H$2:$H${last_data_row})'
            ),
        )

    sheet.cell(row=total_row, column=10, value="Grand Total")
    sheet.cell(
        row=total_row,
        column=11,
        value=f"=SUM(K2:K{last_summary_row})",
    )

    header_fill = PatternFill("solid", fgColor="1F4E78")
    total_fill = PatternFill("solid", fgColor="D9EAF7")
    for cell in (*sheet[1][:8], sheet["J1"], sheet["K1"]):
        cell.font = Font(name="Arial", bold=True, color="FFFFFF")
        cell.fill = header_fill

    for row in sheet.iter_rows(min_row=2, max_row=last_data_row, min_col=1, max_col=7):
        for cell in row:
            cell.font = Font(name="Arial", color="0000FF")
    for row in sheet.iter_rows(min_row=2, max_row=total_row, min_col=8, max_col=11):
        for cell in row:
            cell.font = Font(name="Arial", color="000000")

    for cell in sheet[total_row][9:11]:
        cell.font = Font(name="Arial", bold=True)
        cell.fill = total_fill

    for row_number in range(2, last_data_row + 1):
        sheet.cell(row=row_number, column=6).number_format = CURRENCY_FORMAT
        sheet.cell(row=row_number, column=8).number_format = CURRENCY_FORMAT
    for row_number in range(2, total_row + 1):
        sheet.cell(row=row_number, column=11).number_format = CURRENCY_FORMAT

    sheet["H1"].comment = Comment(
        "Each line uses Quantity x Unit Cost so edits recalculate automatically.",
        "Budget builder",
    )
    sheet["K1"].comment = Comment(
        "SUMIF groups the live line-cost formulas by category; the chart reads this table.",
        "Budget builder",
    )

    widths = {
        "A": 15,
        "B": 18,
        "C": 38,
        "D": 12,
        "E": 16,
        "F": 18,
        "G": 36,
        "H": 18,
        "J": 18,
        "K": 21,
    }
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:H{last_data_row}"

    chart = BarChart()
    chart.type = "col"
    chart.style = 10
    chart.title = "Budget by Category"
    chart.y_axis.title = "Cost (EUR)"
    chart.x_axis.title = "Category"
    # The chart references formula cells rather than copied values so it stays live.
    chart.add_data(
        Reference(sheet, min_col=11, min_row=1, max_row=last_summary_row),
        titles_from_data=True,
    )
    chart.set_categories(
        Reference(sheet, min_col=10, min_row=2, max_row=last_summary_row)
    )
    chart.height = 8
    chart.width = 14
    sheet.add_chart(chart, "M2")

    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    workbook.calculation.calcMode = "auto"
    OUTPUT_DIR.mkdir(exist_ok=True)
    workbook.save(WORKBOOK)


def main() -> int:
    budget = pd.read_csv(DATA, sep=";")
    try:
        build_workbook(budget)
    except PermissionError:
        print(
            f"Cannot write {WORKBOOK.name}: close it in Excel and run again.",
            file=sys.stderr,
        )
        return 1

    totals = (
        budget.assign(cost_eur=budget["quantity"] * budget["unit_cost_eur"])
        .groupby("category", sort=False)["cost_eur"]
        .sum()
    )
    print("Category totals (EUR):")
    for category, amount in totals.items():
        print(f"  {category}: {amount:,.2f} EUR")
    print(f"Workbook written: {WORKBOOK.relative_to(ROOT)}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
